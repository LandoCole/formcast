"""The workbench page's way in to extraction (``demo/web/workbench/``) in Chromium against the real workbench server.

The page opens on "What would you like to do?" (``#choose``): *Fill a PDF* opens the existing start screen as it is,
*Extract from PDFs* switches the page to its extraction mode (``body[data-mode]``), shown only when the server says
it can extract. The Extract modal (``#extract``) takes the batch (files chosen or dropped, or a folder on the
server's machine), the sample, what to extract and how to read it; *Read the sample* builds the template in the
building panel (``#building``) and hands the setup on (``extract:setup`` on the bus) to the setup screen
(``.view[data-only=setup]``: the fields and their status across the samples, a sample page with the readers' boxes, the
field's inspector and its results on each sample), where every change is an edit of the setup on the server. *Run on
all* saves the template and reads the batch in the run panel (``#running``), then opens the review
(``.view[data-only=review]``: the run's numbers, its documents with problems or set aside, a document's page with the
problem's place outlined, a card per problem with its decisions), which exports the run and sends it to fill; the
choice lists the saved runs and extractions to open again. The server runs in-process over a throwaway copy of
``demo/data`` (conftest
``start_demo_server``), with form maps built without detectors, as in ``test_workbench_page``, whose ``Bench`` drives
the page here too; a template is built from the five corpus samples of ``test_server_extract`` by ``CorpusModel``
(gated there as ``GatedModel``), with no detector weights loaded.
"""

from __future__ import annotations

import base64
import dataclasses
import importlib.util
import io
import json
import re
import threading
import zipfile

import pytest
from playwright.sync_api import expect
from test_server_documents import get_json, small_pdf
from test_server_extract import (  # noqa: F401 (corpus_files and problem_files are fixtures)
    SHIFTED, GatedModel, add_files, assert_proven, corpus_files, job_ended, new_batch, problem_files,
    sample_set, start_job, zip_of,
)
from test_server_jobs import wait_for
from test_workbench_page import Bench, bench, no_model_loads, widget_pdf  # noqa: F401 (bench and no_model_loads are fixtures)

from demo.server import extract_api, extraction
from demo.server.api import Api
from formcast.extract.edit import RenameField, SetField
from formcast.extract.setup import STAGES
from tests.extract.corpus import FIELD_RECTS, SHIFT, clients, make_form


def test_the_workbench_opens_on_the_choice_of_fill_or_extract(demo_server, bench):
    b = bench(demo_server, fill=False)
    choose = b.page.locator("#choose")
    expect(choose).to_be_visible()
    expect(b.page.locator("#choose-title")).to_have_text("What would you like to do?")
    expect(choose.locator(".choice .choice-t")).to_have_text(["Fill a PDF", "Extract from PDFs"])
    expect(b.page.locator("#start")).to_be_hidden()
    assert b.page.evaluate("document.body.dataset.mode") == "fill"


def test_fill_a_pdf_opens_the_start_screen_and_the_fill_flow_works_from_it(demo_server, bench):
    b = bench(demo_server, fill=False)
    b.page.click('.choice[data-go="fill"]')
    expect(b.page.locator("#choose")).to_be_hidden()
    expect(b.page.locator("#start")).to_be_visible()
    expect(b.page.locator("#start-title")).to_have_text("Fill a form")
    b.upload(widget_pdf())
    b.build_new()
    expect(b.boxes()).to_have_count(1)

    b.page.click("#doc-name")  # the toolbar's form name goes back to the choice
    expect(b.page.locator("#choose")).to_be_visible()
    b.page.click('.choice[data-go="fill"]')
    expect(b.page.locator("#start")).to_be_visible()
    b.page.click("#start-close")  # a form is open: Back returns to it
    expect(b.page.locator("#start")).to_be_hidden()
    expect(b.page.locator("#choose")).to_be_hidden()
    expect(b.boxes()).to_have_count(1)


def test_back_on_the_start_screen_returns_to_the_choice_when_no_form_is_open(demo_server, bench):
    b = bench(demo_server)
    expect(b.page.locator("#start-close")).to_be_visible()
    b.page.click("#start-close")
    expect(b.page.locator("#start")).to_be_hidden()
    expect(b.page.locator("#choose")).to_be_visible()


def test_extract_from_pdfs_switches_the_page_to_extraction_and_its_toolbar_goes_back(demo_server, bench):
    b = bench(demo_server, fill=False)
    b.page.click('.choice[data-go="extract"]')
    expect(b.page.locator("#choose")).to_be_hidden()
    expect(b.page.locator("#extract")).to_be_visible()
    assert b.page.evaluate("document.body.dataset.mode") == "extract"
    expect(b.page.locator("#bench")).to_be_hidden()  # the fill screen's toolbar and columns go away
    expect(b.page.locator("#doc-name")).to_be_hidden()
    expect(b.page.locator("#doc-pick")).to_be_visible()
    b.page.evaluate("document.getElementById('extract').hidden = true")  # closed, as it is once a setup opens
    b.page.click("#doc-pick")
    expect(b.page.locator("#choose")).to_be_visible()
    b.page.click('.choice[data-go="fill"]')
    assert b.page.evaluate("document.body.dataset.mode") == "fill"
    expect(b.page.locator("#start")).to_be_visible()


def test_show_start_with_a_dataset_offers_it_as_your_data_and_open_uses_it_without_an_upload(demo_server, bench):
    status, dataset = get_json(demo_server, "/api/data", "POST", {"text": "name,tin\nAda Lee,84-2219037\n"})
    assert status == 200
    b = bench(demo_server)
    b.page.evaluate("d => window.__workbench.showStart({dataset: d, form: 'fw9'})", dataset)
    given = b.page.locator("#data-given")
    expect(given).to_be_visible()
    expect(given).to_contain_text(dataset["name"])
    expect(b.page.locator("#data-file-row")).to_be_hidden()
    expect(b.page.locator("#data-text")).to_be_hidden()
    expect(b.page.locator("#form-pick")).to_have_value("saved:fw9")  # the form named: its saved template first
    expect(b.page.locator("#open")).to_have_text("Open and fill")

    b.page.click("#data-given-clear")  # Clear: the file and paste box are back
    expect(given).to_be_hidden()
    expect(b.page.locator("#data-file-row")).to_be_visible()

    b.page.evaluate("d => window.__workbench.showStart({dataset: d})", dataset)
    b.page.select_option("#form-pick", "demo:fw9")
    uploads = []
    b.page.on("request", lambda r: uploads.append(r.url) if r.url.endswith("/api/data") else None)
    b.page.click("#open")
    expect(b.page.locator("#match")).to_be_visible()  # fw9 is stored: the match prompt, with the data kept for it
    assert uploads == []


def test_a_server_that_cannot_extract_offers_only_fill(start_demo_server, page, monkeypatch):
    real = Api.capabilities

    def no_extract(self, req):
        caps = real(self, req)
        caps.pop("extract")
        return caps

    monkeypatch.setattr(Api, "capabilities", no_extract)
    b = Bench(page, start_demo_server())
    b.goto(fill=False)
    expect(page.locator('.choice[data-go="fill"]')).to_be_visible()
    expect(page.locator('.choice[data-go="extract"]')).to_be_hidden()
    assert b.errors == []


# ------------------------------------------------------------------ the Extract modal and building the template

#: The five samples of the server setup tests (``test_server_extract.sample_set``): three fillable, two flattened.
PEOPLE = clients(5, seed=31)
FILES = sample_set(PEOPLE)


@pytest.fixture(autouse=True)
def setup_detectors_off(extract_detectors_off):
    """A setup's samples are read with no detector weights (conftest ``extract_detectors_off``)."""


def corpus_model(**kw) -> GatedModel:
    """The corpus model for the five ``FILES`` (``CorpusModel``, gated as ``test_server_extract.GatedModel`` is)."""
    return GatedModel({name: client for (name, _), client in zip(FILES, PEOPLE)}, **kw)


def extract_server(start_demo_server, model: GatedModel | None = None, **settings) -> str:
    """A server whose AI is the corpus model, reading runs in this process."""
    return start_demo_server(provider_factory=model or corpus_model(), extract_workers=1, **settings)


def pdf_files(files: list[tuple[str, bytes]], mime: str = "application/pdf") -> list[dict]:
    return [{"name": name, "mimeType": mime, "buffer": data} for name, data in files]


def stored_zip(entries: list[tuple[str, bytes]]) -> bytes:
    """A zip of ``entries`` with no compression, so it is as large as they are."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as z:
        for name, data in entries:
            z.writestr(name, data)
    return buf.getvalue()


def open_extract(b: Bench) -> None:
    """From the choice, Extract from PDFs: the Extract modal."""
    b.page.click('.choice[data-go="extract"]')
    expect(b.page.locator("#extract")).to_be_visible()


def add_five(b: Bench) -> None:
    """Choose the five sample PDFs and wait for the batch's scan."""
    b.page.set_input_files("#drop-files", files=pdf_files(FILES))
    expect(b.page.locator("#batch-sum")).to_contain_text("5 fillable or flattened copies of one layout")


def data_transfer(b: Bench, files: list[tuple[str, bytes]], mime: str = "application/pdf"):
    """A DataTransfer holding ``files``, for a drop event."""
    return b.page.evaluate_handle(
        """([files, mime]) => {
            const dt = new DataTransfer();
            for (const [name, b64] of files) {
                const bytes = Uint8Array.from(atob(b64), (c) => c.charCodeAt(0));
                dt.items.add(new File([bytes], name, { type: mime }));
            }
            return dt;
        }""",
        [[[name, base64.b64encode(data).decode()] for name, data in files], mime],
    )


def listen_for_setups(b: Bench) -> None:
    """Keep every setup the bus hears (``extract:setup``) in ``window.__setups``."""
    b.page.evaluate("() => { window.__setups = []; window.__workbench.on('extract:setup', ({ setup }) => window.__setups.push(setup)); }")


def test_extract_from_pdfs_opens_the_extract_modal_without_the_fill_notes_and_back_returns_to_the_choice(demo_server, bench):
    b = bench(demo_server, fill=False)
    expect(b.page.locator("#ai-note")).to_be_visible()  # AI is off: the page says how to turn it on, for filling
    open_extract(b)
    expect(b.page.locator("#extract-title")).to_have_text("Extract from PDFs")
    expect(b.page.locator("#extract .group > .label")).to_have_text(
        ["1 · Your documents", "2 · The sample", "3 · What to extract", "4 · Detection and reading · automatic, prose skipped"]
    )
    expect(b.page.locator("#drop")).to_contain_text("Drop PDFs or a .zip here")
    expect(b.page.locator("#drop-note")).to_have_text(
        "Zips up to 50 MB; for a bigger batch choose the files or a folder, or type a folder on this machine."
    )
    assert b.page.eval_on_selector("#drop-files", "i => i.multiple") is True
    assert b.page.eval_on_selector("#drop-folder", "i => i.webkitdirectory") is True
    expect(b.page.locator("#batch")).to_be_hidden()
    expect(b.page.locator("#n-sample")).to_have_value("5")
    expect(b.page.locator("#sample-how")).to_have_value("spread")
    expect(b.page.locator('#extract .seg [data-mode="ai"]')).to_have_attribute("aria-pressed", "true")
    expect(b.page.locator("#settings-sum")).to_have_text("automatic, prose skipped")
    expect(b.page.locator("#unsure")).to_be_checked()
    expect(b.page.locator("#unsure")).to_be_disabled()
    expect(b.page.locator("#extract-settings")).to_contain_text("Always on in this version")
    expect(b.page.locator("#ai-fallback")).not_to_be_checked()
    expect(b.page.locator("#ai-note")).to_be_hidden()  # fill copy is not shown while extracting

    b.page.click("#extract-back")
    expect(b.page.locator("#extract")).to_be_hidden()
    expect(b.page.locator("#choose")).to_be_visible()


def test_choosing_five_pdfs_shows_the_batch_summary_and_a_strip_cell_for_each_after_the_scan(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    add_five(b)
    expect(b.page.locator("#batch-sum")).to_contain_text("5 PDFs")
    expect(b.page.locator("#batch-strip li")).to_have_count(5)
    expect(b.page.locator('#batch-strip li[data-state="sample"]')).to_have_count(5)  # five of five: all in the sample
    expect(b.page.locator("#batch-note")).to_contain_text("One layout fits 5 of them (3 fillable, 2 flattened)")


def test_dropping_pdfs_and_then_a_zip_puts_them_in_one_batch_and_start_a_new_batch_starts_another(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    expect(b.page.locator("#new-batch")).to_be_hidden()  # no batch yet
    b.page.dispatch_event("#drop", "drop", {"dataTransfer": data_transfer(b, FILES[:3])})
    expect(b.page.locator("#batch-sum")).to_contain_text("3 PDFs")
    archive = zip_of(FILES[3:])
    b.page.dispatch_event("#drop", "drop", {"dataTransfer": data_transfer(b, [("more.zip", archive)], "application/zip")})
    expect(b.page.locator("#batch-sum")).to_contain_text("5 fillable or flattened copies of one layout")
    expect(b.page.locator("#batch-sum")).to_contain_text("5 PDFs")
    batches = start_demo_server.servers[-1].api.batches._batches
    assert len(batches) == 1

    b.page.click("#new-batch")  # the next files start a batch of their own
    expect(b.page.locator("#batch")).to_be_hidden()
    b.page.dispatch_event("#drop", "drop", {"dataTransfer": data_transfer(b, FILES[:2])})
    expect(b.page.locator("#batch-sum")).to_contain_text("2 PDFs")
    expect(b.page.locator("#batch-strip li")).to_have_count(2)
    assert len(batches) == 2 and sorted(len(x.files) for x in batches.values()) == [2, 5]


def test_a_batch_the_server_no_longer_has_is_cleared_and_the_next_files_start_a_new_one(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    add_five(b)
    batches = start_demo_server.servers[-1].api.batches._batches
    batches.clear()  # as after a restart: the batch's id is unknown
    b.page.set_input_files("#drop-files", files=pdf_files(FILES[:1]))
    expect(b.page.locator("#extract-msg")).to_have_text("That batch is no longer on the server; add the documents again.")
    expect(b.page.locator("#batch")).to_be_hidden()
    b.page.set_input_files("#drop-files", files=pdf_files(FILES[:3]))
    expect(b.page.locator("#batch-sum")).to_contain_text("3 PDFs")
    expect(b.page.locator("#extract-msg")).to_be_hidden()
    assert len(batches) == 1


def test_a_file_dropped_beside_the_drop_zone_while_the_modal_is_open_is_ignored(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    url = b.page.url
    drop_on_body = """dt => {
        const events = ["dragover", "drop"].map((type) => new DragEvent(type, { dataTransfer: dt, bubbles: true, cancelable: true }));
        return events.map((e) => { document.querySelector(".sheet h1").dispatchEvent(e); return e.defaultPrevented; });
    }"""
    assert b.page.evaluate(drop_on_body, data_transfer(b, FILES[:1])) == [False, False]  # the modal closed: as ever
    open_extract(b)
    assert b.page.evaluate(drop_on_body, data_transfer(b, FILES[:1])) == [True, True]
    expect(b.page.locator("#batch")).to_be_hidden()  # nothing was added either
    assert b.page.url == url


def test_the_sample_size_must_be_3_to_20(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    add_five(b)
    b.page.fill("#n-sample", "2")
    b.page.dispatch_event("#n-sample", "change")
    expect(b.page.locator("#extract-msg")).to_be_visible()
    expect(b.page.locator("#extract-msg")).to_contain_text("3 to 20")
    b.page.click("#read-sample")
    expect(b.page.locator("#extract-msg")).to_contain_text("3 to 20")
    expect(b.page.locator("#building")).to_be_hidden()
    b.page.fill("#n-sample", "3")
    b.page.dispatch_event("#n-sample", "change")
    expect(b.page.locator("#extract-msg")).to_be_hidden()
    expect(b.page.locator('#batch-strip li[data-state="sample"]')).to_have_count(3)


def test_ill_pick_them_lists_the_batchs_files_with_checkboxes(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    expect(b.page.locator("#picklist")).to_be_hidden()
    add_five(b)
    b.page.select_option("#sample-how", "pick")
    rows = b.page.locator("#picklist li")
    expect(rows).to_have_count(5)
    expect(rows.locator('input[type="checkbox"]')).to_have_count(5)
    expect(rows).to_contain_text([name for name, _ in FILES])
    expect(rows.locator("input:checked")).to_have_count(5)  # the sample picked so far
    expect(b.page.locator("#pick-more")).to_be_hidden()
    rows.nth(1).locator("input").uncheck()
    expect(b.page.locator('#batch-strip li[data-state="sample"]')).to_have_count(4)


def test_the_pick_list_shows_200_files_at_a_time_and_more_shows_the_rest(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    archive = zip_of([(f"f{i:03d}.pdf", small_pdf(text=f"file {i}")) for i in range(201)])
    b.page.set_input_files("#drop-files", files=pdf_files([("many.zip", archive)], "application/zip"))
    expect(b.page.locator("#batch-sum")).to_contain_text("201 PDFs")
    b.page.select_option("#sample-how", "pick")
    expect(b.page.locator("#picklist li")).to_have_count(200)
    expect(b.page.locator("#pick-more")).to_be_visible()
    b.page.click("#pick-more")
    expect(b.page.locator("#picklist li")).to_have_count(201)
    expect(b.page.locator("#picklist li").last).to_contain_text("f200.pdf")
    expect(b.page.locator("#pick-more")).to_be_hidden()


def test_what_to_extract_offers_your_own_lines_and_the_shipped_fill_templates_as_a_destination(demo_server, bench):
    b = bench(demo_server, fill=False)
    open_extract(b)
    expect(b.page.locator("#mode-detail")).to_contain_text("The AI reads the sample and proposes the fields it finds")
    b.page.click('#extract .seg [data-mode="user"]')
    expect(b.page.locator('#extract .seg [data-mode="user"]')).to_have_attribute("aria-pressed", "true")
    expect(b.page.locator('#extract .seg [data-mode="ai"]')).to_have_attribute("aria-pressed", "false")
    expect(b.page.locator("#lines")).to_be_visible()
    b.page.fill("#lines", "investor name\ntax ID")
    b.page.click('#extract .seg [data-mode="dest"]')
    expect(b.page.locator("#lines")).to_be_hidden()
    expect(b.page.locator("#dest-pick")).to_be_visible()
    expect(b.page.locator("#dest-pick option[value='fw9']")).to_have_count(1)
    values = b.page.eval_on_selector_all("#dest-pick option", "os => os.map(o => o.value)")
    assert {"fw9", "menomonie_app", "wadoc_auth"} <= set(values)
    b.page.click('#extract .seg [data-mode="user"]')
    expect(b.page.locator("#lines")).to_have_value("investor name\ntax ID")  # kept while another mode was shown


def test_read_the_sample_with_no_documents_says_to_add_them_first(demo_server, bench):
    b = bench(demo_server, fill=False)
    open_extract(b)
    b.page.click("#read-sample")
    expect(b.page.locator("#extract-msg")).to_be_visible()
    expect(b.page.locator("#extract-msg")).to_contain_text("Add the documents first")
    expect(b.page.locator("#building")).to_be_hidden()


def test_read_the_sample_builds_the_template_in_five_stages_shown_live_and_then_hands_on_the_setup(start_demo_server, bench):
    model = corpus_model(gate="derive")
    b = bench(extract_server(start_demo_server, model), fill=False)
    open_extract(b)
    add_five(b)
    listen_for_setups(b)
    stages = b.page.locator("#build-stages li")
    try:
        b.page.click("#read-sample")
        expect(b.page.locator("#building")).to_be_visible()
        expect(b.page.locator("#extract")).to_be_hidden()
        expect(b.page.locator("#building-title")).to_have_text("Building the template")
        assert model.reached.wait(20)
        expect(stages).to_have_count(5)
        expect(stages.nth(2)).to_have_attribute("data-status", "running")  # designing the fields, held at the model
        expect(stages.nth(0)).to_have_attribute("data-status", "done")
        expect(stages.nth(4)).to_have_attribute("data-status", "pending")
        expect(b.page.locator("#build-readout")).to_contain_text("Design the fields")
    finally:
        model.released.set()
    b.page.wait_for_function("window.__setups.length === 1", timeout=60_000)
    expect(b.page.locator("#building")).to_be_hidden()
    expect(stages).to_contain_text([label for _, label in STAGES])  # each with its detail ("5 samples read", ...)
    assert [stages.nth(i).get_attribute("data-status") for i in range(5)] == ["done"] * 5
    setup = b.page.evaluate("window.__setups[0]")
    assert len(setup["samples"]) == 5
    assert_proven(setup["proof"])  # every field some sample has is ready, and no sample disagrees
    info = setup["template"]["setup"]
    assert (info["start"], info["detect"], info["ocr"], info["sample_how"]) == ("ai", "auto", "auto", "spread")


def test_the_settings_and_your_lines_are_what_the_template_is_built_with(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    add_five(b)
    listen_for_setups(b)
    b.page.select_option("#sample-how", "first")
    b.page.click('#extract .seg [data-mode="user"]')
    b.page.fill("#lines", "investor legal name\n\ntaxpayer ID\n")
    b.page.evaluate("document.getElementById('extract-settings').open = true")
    b.page.select_option("#detect", "off")
    expect(b.page.locator("#settings-sum")).to_have_text("form fields and text only")
    b.page.select_option("#ocr", "never")
    b.page.click("#read-sample")
    b.page.wait_for_function("window.__setups.length === 1", timeout=60_000)
    info = b.page.evaluate("window.__setups[0].template.setup")
    assert info["start"] == "lines" and info["lines"] == ["investor legal name", "taxpayer ID"]
    assert (info["detect"], info["ocr"], info["sample_how"]) == ("off", "never", "first")
    expect(b.page.locator("#build-stages li").nth(1)).to_have_attribute("data-status", "skipped")


def test_cancel_while_building_stops_the_job_and_returns_to_the_extract_modal(start_demo_server, bench):
    model = corpus_model(gate="derive")
    base = extract_server(start_demo_server, model)
    b = bench(base, fill=False)
    open_extract(b)
    add_five(b)
    listen_for_setups(b)
    try:
        with b.page.expect_response(lambda r: r.url.endswith("/api/jobs") and r.request.method == "POST") as started:
            b.page.click("#read-sample")
        job_id = started.value.json()["id"]
        assert model.reached.wait(20)
        b.page.click("#build-cancel")
        expect(b.page.locator("#building")).to_be_hidden()
        expect(b.page.locator("#extract")).to_be_visible()
    finally:
        model.released.set()
    job = wait_for(lambda: (j := get_json(base, f"/api/jobs/{job_id}")[1])["status"] != "running" and j, timeout=30)
    assert job["status"] == "cancelled"
    assert b.page.evaluate("window.__setups.length") == 0


def test_a_build_that_cannot_start_returns_to_the_extract_modal_with_the_message(demo_server, bench):
    b = bench(demo_server, fill=False)  # no model: the server refuses an extract-setup job
    open_extract(b)
    add_five(b)
    b.page.click("#read-sample")
    expect(b.page.locator("#extract-msg")).to_contain_text("AI is not configured")
    expect(b.page.locator("#extract")).to_be_visible()
    expect(b.page.locator("#building")).to_be_hidden()


def test_the_folder_field_scans_a_folder_on_the_server(start_demo_server, bench, tmp_path):
    folder = tmp_path / "incoming"
    folder.mkdir()
    for name, data in FILES:
        (folder / name).write_bytes(data)
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    b.page.fill("#folder", str(folder))
    b.page.click("#scan")
    expect(b.page.locator("#batch-sum")).to_contain_text("5 fillable or flattened copies of one layout")
    expect(b.page.locator("#batch-sum")).to_contain_text("5 PDFs")
    expect(b.page.locator("#batch-strip li")).to_have_count(5)


def test_a_server_that_does_not_open_folders_says_so_in_the_message(start_demo_server, bench, tmp_path):
    b = bench(start_demo_server(folder_scan=False), fill=False)
    open_extract(b)
    b.page.fill("#folder", str(tmp_path))
    b.page.click("#scan")
    expect(b.page.locator("#extract-msg")).to_have_text("This server does not open folders on its machine. Upload the files instead.")
    expect(b.page.locator("#batch")).to_be_hidden()


def test_a_file_over_the_servers_upload_limit_is_skipped_and_named_and_the_rest_are_added(start_demo_server, bench):
    archive = stored_zip(FILES)
    limit = max(len(data) for _, data in FILES) + 4096  # each PDF fits a request of its own; the zip does not
    assert len(archive) > limit
    b = bench(extract_server(start_demo_server, max_upload_bytes=limit), fill=False)
    open_extract(b)
    b.page.set_input_files("#drop-files", files=pdf_files(FILES) + pdf_files([("all.zip", archive)], "application/zip"))
    expect(b.page.locator("#extract-msg")).to_contain_text("all.zip")
    expect(b.page.locator("#batch-sum")).to_contain_text("5 fillable or flattened copies of one layout")
    expect(b.page.locator("#batch-sum")).to_contain_text("5 PDFs")


def test_the_extract_modal_reopens_with_a_setups_choices(demo_server, bench):
    b = bench(demo_server)  # on the start screen: opening the modal closes it, one overlay at a time
    info = {
        "start": "destination", "lines": [], "line_paths": {}, "destination": "fw9", "detect": "manual",
        "pages": {"0": [1]}, "ocr": "never", "vision": True, "sample_how": "first",
    }  # fmt: skip
    b.page.evaluate("info => window.__workbench.openExtract(info)", info)
    expect(b.page.locator("#extract")).to_be_visible()
    expect(b.page.locator("#start")).to_be_hidden()
    assert b.page.evaluate("document.body.dataset.mode") == "extract"
    expect(b.page.locator('#extract .seg [data-mode="dest"]')).to_have_attribute("aria-pressed", "true")
    expect(b.page.locator("#dest-pick")).to_have_value("fw9")
    expect(b.page.locator("#detect")).to_have_value("manual")
    expect(b.page.locator("#ocr")).to_have_value("never")
    expect(b.page.locator("#sample-how")).to_have_value("first")
    expect(b.page.locator("#settings-sum")).to_have_text("you pick the pages")
    assert b.page.eval_on_selector("#pages-row", "row => row.hidden") is False  # shown in the (closed) settings
    expect(b.page.locator("#detect-pages")).to_have_value("2")  # sample 0's page index 1, counted from 1

    lines = {**info, "start": "lines", "lines": ["investor name", "tax ID"], "destination": None, "detect": "auto"}
    b.page.evaluate("info => window.__workbench.openExtract(info)", lines)
    expect(b.page.locator('#extract .seg [data-mode="user"]')).to_have_attribute("aria-pressed", "true")
    expect(b.page.locator("#lines")).to_have_value("investor name\ntax ID")
    expect(b.page.locator("#detect")).to_have_value("auto")
    assert b.page.eval_on_selector("#pages-row", "row => row.hidden") is True


def test_finding_the_blanks_by_hand_searches_the_pages_named_on_every_sample(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    add_five(b)
    listen_for_setups(b)
    b.page.evaluate("document.getElementById('extract-settings').open = true")
    expect(b.page.locator("#pages-row")).to_be_hidden()
    b.page.select_option("#detect", "manual")
    expect(b.page.locator("#pages-row")).to_be_visible()
    expect(b.page.locator("#pages-row")).to_contain_text("Page numbers from 1, with commas and ranges")
    for typed in ["", "0", "4-2", "x"]:  # nothing named, or nothing that names a page: refused before any job
        b.page.fill("#detect-pages", typed)
        b.page.click("#read-sample")
        expect(b.page.locator("#extract-msg")).to_have_text("Name the pages to search, such as 1 or 1, 3-4.")
        expect(b.page.locator("#building")).to_be_hidden()
        expect(b.page.locator("#read-sample")).to_be_enabled()
    b.page.fill("#detect-pages", "1")
    b.page.click("#read-sample")
    b.page.wait_for_function("window.__setups.length === 1", timeout=60_000)
    info = b.page.evaluate("window.__setups[0].template.setup")
    assert info["detect"] == "manual"
    assert info["pages"] == {str(i): [0] for i in range(5)}  # page 1 of each of the five samples, by position
    expect(b.page.locator("#build-stages li").nth(1)).to_have_attribute("data-status", "done")  # searched, not skipped


@pytest.fixture
def held_scan(monkeypatch):
    """Every batch scan waits for ``held_scan.set()`` before it reads its files (``clear()`` holds the next ones)."""
    gate = threading.Event()
    real = extraction.scan

    def scan(*args, **kwargs):
        assert gate.wait(20)
        return real(*args, **kwargs)

    monkeypatch.setattr(extraction, "scan", scan)
    yield gate
    gate.set()


def test_the_strip_reads_not_scanned_yet_until_the_scan_is_done_and_drops_the_sample_marks_when_it_starts_again(
    start_demo_server, bench, held_scan
):
    b = bench(extract_server(start_demo_server), fill=False)
    open_extract(b)
    cells = b.page.locator("#batch-strip li")
    b.page.set_input_files("#drop-files", files=pdf_files(FILES[:3]))
    expect(b.page.locator("#batch-sum")).to_contain_text("scanning")
    expect(cells).to_have_count(3)
    expect(b.page.locator('#batch-strip li[data-state="pending"][title="not scanned yet"]')).to_have_count(3)
    held_scan.set()
    expect(b.page.locator('#batch-strip li[data-state="sample"]')).to_have_count(3)

    held_scan.clear()  # two more files: the batch is scanned again, and its old sample is no longer on show
    b.page.set_input_files("#drop-files", files=pdf_files(FILES[3:]))
    expect(cells).to_have_count(5)
    expect(b.page.locator('#batch-strip li[data-state="pending"]')).to_have_count(5)
    expect(b.page.locator('#batch-strip li[data-state="sample"]')).to_have_count(0)
    held_scan.set()
    expect(b.page.locator('#batch-strip li[data-state="sample"]')).to_have_count(5)

    b.page.fill("#n-sample", "25")  # a size that is refused leaves no sample marked
    b.page.dispatch_event("#n-sample", "change")
    expect(b.page.locator('#batch-strip li[data-state="sample"]')).to_have_count(0)


def test_a_double_click_on_read_the_sample_starts_one_build(start_demo_server, bench):
    model = corpus_model(gate="derive")
    b = bench(extract_server(start_demo_server, model), fill=False)
    open_extract(b)
    add_five(b)
    starts = []
    b.page.on("request", lambda r: starts.append(r.url) if r.method == "POST" and r.url.endswith("/api/jobs") else None)
    try:
        disabled = b.page.evaluate(
            "() => { const go = document.getElementById('read-sample'); go.click(); const d = go.disabled; go.click(); return d; }"
        )
        assert disabled is True  # from the first click until the building panel opens
        assert model.reached.wait(20)
        expect(b.page.locator("#building")).to_be_visible()
    finally:
        model.released.set()
    b.page.wait_for_timeout(300)
    assert len(starts) == 1


# ------------------------------------------------------------------ the setup screen


def setup_on(base: str, files: list[tuple[str, bytes]] = FILES, samples: list[int] | None = None, **spec) -> dict:
    """A setup built through the API from a new batch of ``files`` (``samples`` the sample, else every file), as the
    modal builds it."""
    batch = new_batch(base, name="Client forms")
    status, _ = add_files(base, batch["id"], files)
    assert status == 200
    samples = list(range(len(files))) if samples is None else samples
    status, out = start_job(base, {"kind": "extract-setup", "batch": batch["id"], "samples": samples, "spec": {"name": "Client forms", "names_from_titles": False, **spec}})
    assert status == 200, out
    job = job_ended(base, out["id"])
    assert job["status"] == "done", job
    return job["result"]["setup"]


#: The edits that give the corpus template's tax ID (required) and country (optional) a reason, as the agent's
#: classification leaves them.
REASONS = [
    {"op": "set_field", "path": "investor.tax_id", "required": True, "required_reason": "the form marks it required"},
    {"op": "set_field", "path": "investor.country", "required": False, "required_reason": "a copy may leave it empty"},
]


def edit_setup(base: str, setup: dict, edit: dict) -> dict:
    """The setup with ``edit`` made on it through the API."""
    status, out = get_json(base, f"/api/setups/{setup['id']}/edit", "POST", {"edit": edit})
    assert status == 200, out
    return out


def show_setup(b: Bench, setup: dict) -> None:
    """Open ``setup`` on the setup screen, as the Extract modal does once it is built."""
    b.page.evaluate("s => window.__workbench.openSetup(s)", setup)
    expect(b.page.locator('.view[data-only="setup"]')).to_be_visible()
    expect(b.page.locator("#setup-pv .fc-box").first).to_be_attached()


def row(b: Bench, path: str):
    """The field list's row of ``path`` (a repeated group's field by its first item's path)."""
    return b.page.locator(f'#fields .fc-sl-leaf[data-path="{path}"]')


def picked(b: Bench, kind: str = "box"):
    """The selected field's marks of ``kind`` on the page, drawn heavy."""
    return b.page.locator(f'#setup-pv .fc-box.is-picked[data-kind="{kind}"]')


def edits_sent(b: Bench) -> list[dict]:
    """Every edit the page sends from now on, as sent."""
    sent: list[dict] = []
    b.page.on("request", lambda r: sent.append(r.post_data_json["edit"]) if r.method == "POST" and r.url.endswith("/edit") else None)
    return sent


def setup_idle(b: Bench) -> None:
    """Wait until the setup screen has no edit on its way."""
    b.page.wait_for_function("window.__workbench.setupIdle")


def media_src(b: Bench) -> str:
    return b.page.locator("#setup-pv .fc-pv-media").get_attribute("src") or ""


def keep_the_models_names(b: Bench) -> None:
    """Builds started from the page keep the names the corpus model answers with, whatever its titles would make."""

    def keep(route) -> None:
        body = json.loads(route.request.post_data)
        if body.get("kind") == "extract-setup":
            body["spec"]["names_from_titles"] = False
        route.continue_(post_data=json.dumps(body))

    b.page.route("**/api/jobs", keep)


def on_every_sample(setup: dict) -> list[str]:
    """The plain fields every sample agrees on: ruling P1 asserts ``● 5/5`` only of these."""
    return [
        f["path"] for f in setup["proof"]["fields"]
        if "[" not in f["path"] and all(s["outcome"] == "agree" for s in f["samples"])
    ]  # fmt: skip


def test_the_setup_screen_opens_after_the_build_with_each_fields_status_across_the_samples(start_demo_server, bench):
    b = bench(extract_server(start_demo_server), fill=False)
    keep_the_models_names(b)
    open_extract(b)
    add_five(b)
    listen_for_setups(b)
    b.page.click("#read-sample")
    b.page.wait_for_function("window.__setups.length === 1", timeout=60_000)
    setup = b.page.evaluate("window.__setups[0]")
    expect(b.page.locator('.view[data-only="setup"]')).to_be_visible()
    expect(b.page.locator("#extract-empty")).to_be_hidden()
    expect(b.page.locator("#doc-pick")).to_contain_text("5 PDFs · 5 samples")
    expect(b.page.locator("#strip")).to_contain_text("AI suggests · 5 samples · automatic, prose skipped")
    expect(b.page.locator("#change")).to_be_visible()
    full = on_every_sample(setup)
    assert len(full) >= 5
    for path in full:
        expect(row(b, path).locator(".fc-sl-st")).to_have_text("●")
        expect(row(b, path).locator(".fc-sl-ty")).to_have_text("5/5")
    expect(row(b, "owners[0].name").locator(".fc-sl-name")).to_have_text("Owner name")  # the group's items in one row
    expect(row(b, "owners[1].name")).to_have_count(0)
    expect(b.page.locator("#fields .fc-sl-grp")).to_have_text(["Investor", "Mailing address", "Beneficial owners · up to 4", "Wire instructions"])
    rows = b.page.locator("#fields .fc-sl-leaf").count()
    expect(b.page.locator("#fields-label")).to_have_text(f"Fields · {rows} · 0 need a look")
    expect(b.page.locator("#cap")).to_contain_text(f"{setup['proof']['ready']} of {setup['proof']['total']} fields ready")
    expect(b.page.locator("#sample-no")).to_have_text("Sample 1 / 5")
    expect(b.page.locator("#setup-pv .fc-pv-media")).to_be_attached(timeout=10_000)
    assert f"/api/setups/{setup['id']}/samples/0/pages/0.png" in media_src(b)
    expect(b.page.locator("#setup-pv .fc-box").first).to_be_visible()


def test_selecting_a_field_heavy_draws_its_box_with_its_path_and_the_inspector_says_how_it_is_read(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "investor.tax_id").click()
    expect(row(b, "investor.tax_id")).to_have_class("fc-sl-leaf is-picked")
    expect(picked(b)).to_have_count(1)
    expect(picked(b).locator(".fc-tag")).to_have_text("Taxpayer ID")
    expect(picked(b).locator(".fc-tag")).to_be_visible()
    expect(b.page.locator("#setup-pv .fc-box.is-picked")).to_have_count(2)  # its box and its label, on a fillable copy
    ins = b.page.locator("#ins")
    expect(b.page.locator("#ins-name")).to_have_text("investor.tax_id")
    expect(b.page.locator("#ins-title")).to_have_value("Taxpayer ID")
    expect(b.page.locator("#ins-type")).to_have_value("text")
    expect(ins.locator(".ro")).to_contain_text(["Form field tax_id", "A box on page 1", "Next to the label Taxpayer ID"])
    expect(b.page.locator("#ins-labels")).to_have_value("Taxpayer ID")
    expect(b.page.locator("#ins-required")).to_be_checked()
    expect(b.page.locator("#ins-pattern")).to_have_value(r"^\d{2}-\d{7}$|^\d{3}-\d{2}-\d{4}$")
    expect(ins).to_contain_text("label 'TIN' is on no sample")  # the build's warning about this field
    lines = b.page.locator("#per-sample li")
    expect(lines).to_have_count(5)
    expect(lines.nth(0)).to_contain_text(f"1. {PEOPLE[0].tax_id}")
    expect(lines.nth(3)).to_contain_text("s4.pdf · flattened")
    expect(lines.nth(0)).to_have_class("on")

    b.page.locator('#setup-pv .fc-box[data-field-id="investor.dob#box"]').click()  # a box on the page selects its field
    expect(row(b, "investor.dob")).to_have_class("fc-sl-leaf is-picked")
    expect(b.page.locator("#ins-title")).to_have_value("Date of birth or formation")
    expect(picked(b).locator(".fc-tag")).to_have_text("Date of birth or formation")

    row(b, "investor.type").click()  # a checkbox group: its four boxes, and its answers
    expect(picked(b, "option")).to_have_count(4)
    expect(ins.locator(".answers li")).to_have_count(4)
    expect(ins.locator(".answers")).to_contain_text("trust")

    lines.nth(3).click()  # a sample's line shows that sample
    expect(b.page.locator("#sample-no")).to_have_text("Sample 4 / 5")
    expect(lines.nth(3)).to_have_class("on")


def test_a_box_dragged_on_the_shifted_sample_sends_one_set_box_naming_it_and_lands_where_it_was_dropped(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    for _ in range(SHIFTED):
        b.page.click("#next-s")
    expect(b.page.locator("#sample-no")).to_have_text(f"Sample {SHIFTED + 1} / 5")
    expect(b.page.locator("#sample-kind")).to_have_text(f"s{SHIFTED + 1}.pdf · flattened")
    # the legal name has no check, so its box over the tax ID reads a wrong value there (a value that fails a
    # check, as the date would under the tax ID's box, is passed over for the next way of reading the field)
    row(b, "investor.legal_name").click()
    box = picked(b)
    expect(box).to_have_count(1)
    b.page.wait_for_function("() => document.querySelector('#setup-pv .fc-pv-media')?.src.includes('/samples/4/')")
    sent = edits_sent(b)
    px_per_pt = box.bounding_box()["height"] / 20  # the box is 20 points tall
    before = box.bounding_box()["y"]
    b.drag(box, 0, 30 * px_per_pt)  # down onto the tax ID
    expect(row(b, "investor.legal_name").locator(".fc-sl-ty")).not_to_have_text("5/5")  # proven again
    assert len(sent) == 1, sent
    change = sent[0]
    assert (change["op"], change["path"], change["page"], change["sample"]) == ("set_box", "investor.legal_name", 0, SHIFTED)
    x0, y0, x1, y1 = FIELD_RECTS["investor_name"]
    assert change["rect"] == pytest.approx([x0, y0 + SHIFT + 30, x1, y1 + SHIFT + 30], abs=1)
    expect(b.page.locator("#per-sample li").nth(SHIFTED).locator("span").first).to_have_text("×")
    expect(b.page.locator("#per-sample li").nth(SHIFTED)).to_contain_text(PEOPLE[SHIFTED].tax_id)  # what it read there
    marks = get_json(base, f"/api/setups/{setup['id']}/samples/{SHIFTED}/marks")[1]["marks"]
    mark = next(m for m in marks if (m["path"], m["kind"]) == ("investor.legal_name", "box"))
    assert mark["rect"] == pytest.approx(change["rect"], abs=0.5)  # on this sample, where it was dropped
    strategies = get_json(base, f"/api/setups/{setup['id']}")[1]["template"]["readers"]["investor.legal_name"]["strategies"]
    assert next(s["rect"] for s in strategies if s["kind"] == "box") == pytest.approx([x0, y0 + 30, x1, y1 + 30], abs=1)
    setup_idle(b)
    expect(picked(b)).to_have_count(1)  # drawn again from the new marks, where it was dropped
    assert picked(b).bounding_box()["y"] == pytest.approx(before + 30 * px_per_pt, abs=2)


def test_an_edited_definition_is_kept_after_selecting_another_field_and_back(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "investor.tax_id").click()
    sent = edits_sent(b)
    b.page.fill("#ins-def", "The EIN or SSN, with its dashes.")
    b.page.locator("#ins-def").blur()
    setup_idle(b)
    assert sent == [{"op": "set_field", "path": "investor.tax_id", "description": "The EIN or SSN, with its dashes."}]
    row(b, "investor.dob").click()
    expect(b.page.locator("#ins-def")).to_have_value("The investor's date of birth, or the date the entity was formed.")
    row(b, "investor.tax_id").click()
    expect(b.page.locator("#ins-def")).to_have_value("The EIN or SSN, with its dashes.")

    b.page.fill("#ins-title", "Tax ID")
    b.page.press("#ins-title", "Enter")
    setup_idle(b)
    b.page.uncheck("#ins-required")
    setup_idle(b)
    b.page.fill("#ins-labels", "Taxpayer ID, TIN number")
    b.page.press("#ins-labels", "Enter")
    setup_idle(b)
    assert sent[1:] == [
        {"op": "set_field", "path": "investor.tax_id", "title": "Tax ID"},
        {"op": "set_field", "path": "investor.tax_id", "required": False},
        {"op": "set_labels", "path": "investor.tax_id", "labels": ["Taxpayer ID", "TIN number"]},
    ]
    template = get_json(base, f"/api/setups/{setup['id']}")[1]["template"]
    assert template["schema"]["properties"]["investor"]["properties"]["tax_id"]["title"] == "Tax ID"
    expect(b.page.locator("#ins-title")).to_have_value("Tax ID")
    expect(b.page.locator("#ins-required")).not_to_be_checked()
    expect(b.page.locator("#ins-labels")).to_have_value("Taxpayer ID, TIN number")


def test_not_on_this_form_marks_the_field_absent_bring_back_restores_it_and_remove_takes_it_out(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "wire.bank").click()
    expect(b.page.locator("#ins-absent")).to_have_text("Not on this form")
    b.page.click("#ins-absent")
    expect(row(b, "wire.bank").locator(".fc-sl-st")).to_have_text("○")
    expect(row(b, "wire.bank")).to_have_attribute("data-state", "absent")
    expect(row(b, "wire.bank")).to_have_class("fc-sl-leaf is-picked")  # still the field selected
    expect(b.page.locator("#per-sample li")).to_have_count(5)
    expect(b.page.locator("#per-sample li").first).to_contain_text("not on this form")
    expect(b.page.locator("#ins .ro")).to_contain_text(["Not on this form"])
    expect(b.page.locator('#setup-pv .fc-box[data-field-id^="wire.bank#"]')).to_have_count(0)
    expect(b.page.locator("#ins-absent")).to_have_text("Bring back")
    b.page.click("#ins-absent")
    expect(row(b, "wire.bank")).not_to_have_attribute("data-state", "absent")
    expect(b.page.locator("#ins-absent")).to_have_text("Not on this form")

    b.page.click("#ins-remove")
    expect(row(b, "wire.bank")).to_have_count(0)
    assert "wire.bank" not in get_json(base, f"/api/setups/{setup['id']}")[1]["template"]["readers"]
    expect(b.page.locator("#fields .fc-sl-leaf.is-picked")).to_have_count(1)  # another field is selected


def test_add_a_field_from_a_line_adds_its_fields_and_selects_the_first(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base, start="lines", lines=["tax ID"])
    b = bench(base, fill=False)
    show_setup(b, setup)
    expect(row(b, "wire.bank")).to_have_count(0)
    b.page.fill("#add", "wire details")
    b.page.click("#add-go")
    expect(row(b, "wire.bank")).to_have_class("fc-sl-leaf is-picked", timeout=30_000)  # the CorpusModel's answer
    expect(row(b, "wire.account")).to_be_visible()
    expect(b.page.locator("#add")).to_have_value("")
    expect(b.page.locator("#strip textarea")).to_have_value("tax ID\nwire details")
    expect(b.page.locator("#ins-title")).to_have_value("Sending bank")

    b.page.fill("#add", "the moon's phase")  # nothing on the samples answers it: the job's reason is shown
    b.page.click("#add-go")
    expect(b.page.locator("#add-msg")).to_contain_text("the schema has no fields", timeout=30_000)
    expect(b.page.locator("#add")).to_have_value("the moon's phase")
    expect(b.page.locator("#add-go")).to_be_enabled()


def test_the_arrows_walk_the_samples_and_their_pages_and_the_zoom_fits_the_page(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    for n in range(1, 6):
        expect(b.page.locator("#sample-no")).to_have_text(f"Sample {n} / 5")
        expect(b.page.locator("#sample-kind")).to_have_text(f"s{n}.pdf · {'fillable' if n <= 3 else 'flattened'}")
        b.page.wait_for_function("n => document.querySelector('#setup-pv .fc-pv-media')?.src.includes(`/samples/${n}/pages/0.png`)", arg=n - 1)
        b.page.click("#next-s")
    expect(b.page.locator("#sample-no")).to_have_text("Sample 1 / 5")  # round again
    b.page.click("#prev-s")
    expect(b.page.locator("#sample-no")).to_have_text("Sample 5 / 5")
    expect(b.page.locator("#page-of")).to_have_text("Page 1 / 2")
    expect(b.page.locator("#prev-p")).to_be_disabled()
    b.page.click("#next-p")
    expect(b.page.locator("#page-of")).to_have_text("Page 2 / 2")
    expect(b.page.locator("#next-p")).to_be_disabled()
    b.page.wait_for_function("() => document.querySelector('#setup-pv .fc-pv-media')?.src.includes('/samples/4/pages/1.png')")
    row(b, "investor.tax_id").click()  # its box is on page 1: the page turns to it
    expect(b.page.locator("#page-of")).to_have_text("Page 1 / 2")
    expect(picked(b)).to_have_count(1)

    width = "document.querySelector('#setup-pv .fc-pv-page').getBoundingClientRect().width"
    expect(b.page.locator('#setup-zoom [data-fit="width"]')).to_have_attribute("aria-pressed", "true")
    wide = b.page.evaluate(width)
    b.page.click('#setup-zoom [data-fit="page"]')
    expect(b.page.locator('#setup-zoom [data-fit="page"]')).to_have_attribute("aria-pressed", "true")
    b.page.wait_for_function(f"() => {width} < {wide - 1}")
    b.page.click('#setup-zoom [data-step="1"]')
    expect(b.page.locator('#setup-zoom [data-fit="page"]')).to_have_attribute("aria-pressed", "false")
    b.page.click("#setup-zoom [data-level]")  # back to fit width
    expect(b.page.locator('#setup-zoom [data-fit="width"]')).to_have_attribute("aria-pressed", "true")
    b.page.wait_for_function(f"() => Math.abs({width} - {wide}) < 2")


def test_save_asks_for_a_name_and_stores_the_extraction_template(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    b.page.click("#setup-save")
    expect(b.page.locator("#setup-save-box")).to_be_visible()
    expect(b.page.locator("#setup-save-name")).to_have_value("Client forms")
    b.page.fill("#setup-save-name", "../nope")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-msg")).to_contain_text("letters, digits")
    b.page.fill("#setup-save-name", "client intake")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    expect(b.page.locator("#notice")).to_contain_text("Saved “client intake”")
    listed = get_json(base, "/api/extractions")[1]["extractions"]
    assert [e["name"] for e in listed] == ["client intake"] and listed[0]["samples"] == 5


def test_save_offers_a_name_not_saved_yet_and_asks_before_replacing_a_saved_extraction(start_demo_server, bench):
    base = extract_server(start_demo_server)
    first = setup_on(base)
    assert get_json(base, "/api/extractions/Client%20forms", "PUT", {"setup": first["id"]})[0] == 200
    setup = setup_on(base)  # a second batch of the same name: its setup is named the same
    b = bench(base, fill=False)
    show_setup(b, setup)
    saves = []
    b.page.on("request", lambda r: saves.append(r.url) if r.method == "PUT" else None)
    b.page.click("#setup-save")
    expect(b.page.locator("#setup-save-name")).to_have_value("Client forms 2")  # "Client forms" is saved already

    b.page.fill("#setup-save-name", "Client forms")  # typed over the saved one: asked first
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-replace")).to_be_visible()
    expect(b.page.locator("#setup-replace-text")).to_have_text("Replace the saved extraction “Client forms”?")
    expect(b.page.locator("#setup-save-go")).to_be_hidden()
    b.page.click("#setup-replace-cancel")
    expect(b.page.locator("#setup-replace")).to_be_hidden()
    expect(b.page.locator("#setup-save-box")).to_be_visible()
    assert saves == []  # nothing was sent

    b.page.click("#setup-save-go")
    b.page.click("#setup-replace-go")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    expect(b.page.locator("#notice")).to_contain_text("Saved “Client forms”")
    assert len(saves) == 1
    b.page.click("#setup-save")  # its own name now: saved again without a question
    expect(b.page.locator("#setup-save-name")).to_have_value("Client forms")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    assert len(saves) == 2 and [e["name"] for e in get_json(base, "/api/extractions")[1]["extractions"]] == ["Client forms"]


def test_save_shows_why_the_server_refuses_a_field_whose_column_clashes_with_the_exports_name(start_demo_server, bench):
    from test_server_extract import with_a_field_named_name

    base = extract_server(start_demo_server)
    setup = setup_on(base)
    kept = start_demo_server.servers[-1].api.setups.get(setup["id"])
    kept.template = with_a_field_named_name(kept.template)
    b = bench(base, fill=False)
    show_setup(b, setup)
    b.page.click("#setup-save")
    b.page.fill("#setup-save-name", "Clash")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-msg")).to_have_text(
        "the field 'document_name' would share the document column document_name; change its title"
    )
    expect(b.page.locator("#setup-save-box")).to_be_visible()
    assert get_json(base, "/api/extractions")[1] == {"extractions": []}


def test_an_edit_after_the_server_has_forgotten_the_setup_returns_to_the_choice_with_a_notice(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "investor.tax_id").click()
    start_demo_server.servers[-1].api.setups._setups.clear()  # as after a restart
    b.page.click("#ins-absent")
    expect(b.page.locator("#choose")).to_be_visible()
    expect(b.page.locator("#choose-msg")).to_have_text("This setup is no longer on the server; start again.")
    expect(b.page.locator('.view[data-only="setup"]')).to_be_hidden()
    expect(b.page.locator("#extract-empty")).to_be_visible()
    b.page.click('.choice[data-go="extract"]')  # the message goes with the choice made
    expect(b.page.locator("#extract")).to_be_visible()
    expect(b.page.locator("#choose-msg")).to_be_hidden()


def test_a_repeated_fields_not_on_this_form_after_the_server_forgot_the_setup_leaves_the_screen_idle(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "owners[0].name").click()
    sent = edits_sent(b)
    start_demo_server.servers[-1].api.setups._setups.clear()
    b.page.click("#ins-absent")  # one edit per owner: the first is refused, the other three are skipped
    expect(b.page.locator("#choose")).to_be_visible()
    expect(b.page.locator("#choose-msg")).to_have_text("This setup is no longer on the server; start again.")
    setup_idle(b)
    assert len(sent) == 1, sent


def test_an_edit_made_while_a_line_job_runs_waits_for_it_and_is_then_made(start_demo_server, bench):
    model = corpus_model()
    base = extract_server(start_demo_server, model)
    setup = setup_on(base, start="lines", lines=["tax ID"])
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "investor.tax_id").click()
    sent = edits_sent(b)
    model.gate = "fields"  # the line job holds at the model
    try:
        b.page.fill("#add", "wire details")
        b.page.click("#add-go")
        assert model.reached.wait(20)
        b.page.fill("#ins-def", "The EIN or SSN, with its dashes.")
        b.page.locator("#ins-def").blur()
        b.page.wait_for_timeout(300)
        assert sent == []  # waiting behind the line job, never sent to be refused as busy
    finally:
        model.gate = None
        model.released.set()
    expect(row(b, "wire.bank")).to_be_visible(timeout=30_000)
    setup_idle(b)
    assert sent == [{"op": "set_field", "path": "investor.tax_id", "description": "The EIN or SSN, with its dashes."}]
    schema = get_json(base, f"/api/setups/{setup['id']}")[1]["template"]["schema"]
    assert schema["properties"]["investor"]["properties"]["tax_id"]["description"] == "The EIN or SSN, with its dashes."
    assert "wire" in schema["properties"]
    expect(b.page.locator("#notice")).to_be_hidden()


def names_following(server, setup: dict, base: str) -> dict:
    """``setup`` as a new template's: its names follow its titles."""
    kept = server.servers[-1].api.setups.get(setup["id"])
    info = kept.template.setup.model_copy(update={"name_follows_title": True})
    kept.template = kept.template.model_copy(update={"setup": info})
    return get_json(base, f"/api/setups/{setup['id']}")[1]


def test_a_title_gives_the_field_its_name_and_the_field_stays_selected(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = names_following(start_demo_server, setup_on(base), base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "investor.tax_id").click()
    expect(b.page.locator("#ins-category")).to_have_value("Investor")
    expect(b.page.locator("#ins-name")).to_have_text("investor.tax_id")
    b.page.fill("#ins-title", "TIN number")
    b.page.press("#ins-title", "Enter")
    expect(row(b, "investor.tin_number")).to_have_class(re.compile("is-picked"))
    expect(row(b, "investor.tax_id")).to_have_count(0)
    expect(b.page.locator("#ins-name")).to_have_text("investor.tin_number")


def test_a_repeated_groups_field_moves_in_every_item_and_stays_on_its_item(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = names_following(start_demo_server, setup_on(base), base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "owners[0].name").click()
    expect(b.page.locator("#ins-category")).to_have_js_property("readOnly", True)
    b.page.select_option("#ins-item", "owners[1].name")
    b.page.fill("#ins-title", "Holder")
    b.page.press("#ins-title", "Enter")
    expect(b.page.locator("#ins-item")).to_have_value("owners[1].holder")


def test_an_edit_made_while_a_rename_is_on_its_way_reaches_the_renamed_field(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = names_following(start_demo_server, setup_on(base), base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "investor.country").click()
    held = []
    b.page.route("**/edit", lambda route: held.append(route), times=1)
    b.page.fill("#ins-title", "Nation")
    b.page.press("#ins-title", "Enter")
    b.page.fill("#ins-def", "Where the investor lives.")
    b.page.check("#ins-required")  # asked for on the old name
    row(b, "investor.dob").click()  # and the person has moved on
    for _ in range(200):  # the route handler runs while the page is driven
        if held:
            break
        b.page.wait_for_timeout(50)
    held[0].continue_()
    expect(row(b, "investor.nation").locator(".fc-sl-req")).to_be_visible()
    expect(row(b, "investor.dob")).to_have_class(re.compile("is-picked"))  # the selection stays where the person put it
    expect(b.page.locator("#notice")).to_be_hidden()


def test_a_category_moves_the_field_in_the_list(start_demo_server, bench):
    base = extract_server(start_demo_server)
    b = bench(base, fill=False)
    show_setup(b, names_following(start_demo_server, setup_on(base), base))
    row(b, "investor.country").click()
    b.page.fill("#ins-category", "Mailing")
    b.page.press("#ins-category", "Enter")
    expect(row(b, "mailing.country")).to_have_class(re.compile("is-picked"))
    expect(b.page.locator("#ins-name")).to_have_text("mailing.country")
    b.page.fill("#ins-category", "")
    b.page.press("#ins-category", "Enter")
    expect(row(b, "country")).to_have_class(re.compile("is-picked"))


def test_a_refused_title_says_why_and_puts_the_title_back(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = names_following(start_demo_server, setup_on(base), base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "investor.tax_id").click()
    b.page.fill("#ins-title", "#")
    b.page.press("#ins-title", "Enter")
    expect(b.page.locator("#notice")).to_contain_text("letter or number")
    expect(b.page.locator("#ins-title")).to_have_value("Taxpayer ID")


def test_a_reserved_title_is_refused_and_the_title_put_back(start_demo_server, bench):
    from test_server_extract import with_a_field_named_name

    base = extract_server(start_demo_server)
    setup = setup_on(base)
    kept = start_demo_server.servers[-1].api.setups.get(setup["id"])
    kept.template = with_a_field_named_name(kept.template)
    kept.edit(RenameField(op="rename_field", path="document_name", name="holder"))  # a top-level field, proven again
    setup = names_following(start_demo_server, setup, base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "holder").click()
    b.page.fill("#ins-title", "DOCUMENT NAME")
    b.page.press("#ins-title", "Enter")
    expect(b.page.locator("#notice")).to_contain_text("reserved")
    expect(row(b, "document_name")).to_have_count(0)


def test_once_saved_a_title_keeps_the_name_and_update_name_changes_it(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = names_following(start_demo_server, setup_on(base), base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    b.page.click("#setup-save")
    b.page.fill("#setup-save-name", "Kept")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    row(b, "investor.dob").click()
    expect(b.page.locator("#ins-update-name")).to_be_visible()  # its title says date_of_birth_or_formation
    expect(b.page.locator("#ins-category")).to_have_js_property("readOnly", True)  # part of the export column
    b.page.fill("#ins-title", "Born on")
    with b.page.expect_response("**/edit"):
        b.page.press("#ins-title", "Enter")
    expect(b.page.locator("#ins-name")).to_have_text("investor.dob")
    b.page.click("#ins-update-name")
    expect(row(b, "investor.born_on")).to_have_class(re.compile("is-picked"))
    expect(b.page.locator("#ins-update-name")).to_have_count(0)


def test_the_field_the_save_refuses_can_be_retitled_and_then_saved(start_demo_server, bench):
    from test_server_extract import with_a_field_named_name

    base = extract_server(start_demo_server)
    setup = setup_on(base)
    kept = start_demo_server.servers[-1].api.setups.get(setup["id"])
    kept.template = with_a_field_named_name(kept.template)
    kept.edit(SetField(op="set_field", path="document_name", title="Document name"))  # proven again, so it is on the list
    b = bench(base, fill=False)
    show_setup(b, names_following(start_demo_server, setup, base))
    b.page.click("#setup-save")
    b.page.fill("#setup-save-name", "Clash")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-msg")).to_contain_text("change its title")
    b.page.click("#setup-save-cancel")
    row(b, "document_name").click()
    b.page.fill("#ins-title", "Legal name")
    b.page.press("#ins-title", "Enter")
    expect(row(b, "legal_name")).to_be_visible()
    b.page.click("#setup-save")
    b.page.fill("#setup-save-name", "Clash")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    assert [e["name"] for e in get_json(base, "/api/extractions")[1]["extractions"]] == ["Clash"]


def test_a_refused_edit_puts_back_its_own_box_and_keeps_what_is_typed_in_another(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "investor.tax_id").click()
    held = []
    b.page.route("**/edit", lambda route: held.append(route), times=1)  # its answer waits until another box is typed in
    b.page.fill("#ins-pattern", "(")  # not a regular expression: the server refuses it
    b.page.click("#ins-labels")
    for _ in range(200):  # the route handler runs while the page is driven
        if held:
            break
        b.page.wait_for_timeout(50)
    assert len(held) == 1
    b.page.fill("#ins-labels", "Taxpayer ID, TIN number")
    held[0].continue_()
    expect(b.page.locator("#notice")).to_contain_text("That change wasn't made")
    expect(b.page.locator("#ins-pattern")).to_have_value(r"^\d{2}-\d{7}$|^\d{3}-\d{2}-\d{4}$")
    expect(b.page.locator("#ins-labels")).to_have_value("Taxpayer ID, TIN number")
    expect(b.page.locator("#ins-labels")).to_be_focused()


def test_a_setup_from_your_lines_shows_them_in_the_strip_and_a_field_for_each_of_their_paths(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base, start="lines", lines=["tax ID", "investor name"])
    b = bench(base, fill=False)
    show_setup(b, setup)
    expect(b.page.locator("#strip")).to_contain_text("I'll list them · 5 samples")
    expect(b.page.locator("#strip textarea")).to_have_value("tax ID\ninvestor name")
    wanted = {p.replace("[]", "[0]") for paths in setup["template"]["setup"]["line_paths"].values() for p in paths}
    assert wanted
    shown = set(b.page.eval_on_selector_all("#fields .fc-sl-leaf", "rows => rows.map(r => r.dataset.path)"))
    assert shown == wanted
    row(b, "investor.legal_name").click()
    expect(b.page.locator("#ins")).to_contain_text("From your line “investor name”.")

    b.page.click("#change")  # the Extract modal again, with this setup's choices
    expect(b.page.locator("#extract")).to_be_visible()
    expect(b.page.locator('#extract .seg [data-mode="user"]')).to_have_attribute("aria-pressed", "true")
    expect(b.page.locator("#lines")).to_have_value("tax ID\ninvestor name")


def test_a_setup_from_a_destination_names_it_and_lists_the_fields_this_form_lacks_as_not_on_it(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base, start="destination", destination="fw9")
    absent = [f["path"] for f in setup["proof"]["fields"] if f["absent"]]
    assert len(absent) > 1
    b = bench(base, fill=False)
    show_setup(b, setup)
    expect(b.page.locator("#strip")).to_contain_text("Destination · 5 samples")
    expect(b.page.locator("#strip")).to_contain_text("fw9")
    expect(b.page.locator("#strip")).to_contain_text(f"{len(absent)} of its fields are not on this form")
    for path in absent:
        expect(row(b, path).locator(".fc-sl-st")).to_have_text("○")
        expect(row(b, path)).to_have_attribute("data-state", "absent")
    expect(row(b, "address.street").locator(".fc-sl-st")).to_have_text("●")


def test_a_wording_seen_on_the_samples_is_added_to_the_answers_map_when_asked(start_demo_server, bench):
    people = [dataclasses.replace(PEOPLE[0], country="Estados Unidos"), *PEOPLE[1:]]
    files = sample_set(people)
    model = GatedModel({name: client for (name, _), client in zip(files, people)}, us_files=["s1.pdf"])
    base = extract_server(start_demo_server, model)
    setup = setup_on(base, files)
    b = bench(base, fill=False)
    show_setup(b, setup)
    row(b, "investor.country").click()
    expect(row(b, "investor.country").locator(".fc-sl-st")).to_have_text("◐")
    wordings = b.page.locator("#ins-wordings li")
    expect(wordings).to_have_count(1)
    expect(wordings).to_have_text("Estados Unidos → US")
    calls = dict(model.calls)
    b.page.click("#ins-learn")
    expect(row(b, "investor.country").locator(".fc-sl-st")).to_have_text("●")
    expect(b.page.locator("#ins-wordings")).to_have_count(0)
    assert model.calls == calls  # no model call
    reader = get_json(base, f"/api/setups/{setup['id']}")[1]["template"]["readers"]["investor.country"]
    assert reader["clean"][0]["args"]["values"]["Estados Unidos"] == "US"


def test_the_setup_inspector_shows_the_required_reason_and_a_change_is_set_by_you(start_demo_server, bench):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    for edit in REASONS:
        setup = edit_setup(base, setup, edit)
    b = bench(base, fill=False)
    show_setup(b, setup)
    expect(row(b, "investor.tax_id").locator(".fc-sl-req")).to_have_text("*")  # the list marks the required fields
    expect(row(b, "investor.country").locator(".fc-sl-req")).to_have_count(0)
    row(b, "investor.tax_id").click()
    expect(b.page.locator("#ins-required")).to_be_checked()
    expect(b.page.locator("#ins-required-why")).to_have_text("the form marks it required")
    row(b, "investor.country").click()
    expect(b.page.locator("#ins-required")).not_to_be_checked()
    expect(b.page.locator("#ins-required-why")).to_have_text("a copy may leave it empty")

    sent = edits_sent(b)
    b.page.fill("#ins-title", "Country code")
    b.page.locator("#ins-title").press("Enter")  # another change of the field leaves its level and reason alone
    setup_idle(b)
    expect(b.page.locator("#ins-required-why")).to_have_text("a copy may leave it empty")
    b.page.check("#ins-required")  # the person's choice: sent with no reason, so the server says it is theirs
    setup_idle(b)
    assert sent == [
        {"op": "set_field", "path": "investor.country", "title": "Country code"},
        {"op": "set_field", "path": "investor.country", "required": True},
    ]
    expect(b.page.locator("#ins-required")).to_be_checked()
    expect(b.page.locator("#ins-required-why")).to_have_text("Set by you")
    expect(row(b, "investor.country").locator(".fc-sl-req")).to_have_text("*")
    investor = get_json(base, f"/api/setups/{setup['id']}")[1]["template"]["schema"]["properties"]["investor"]
    assert "country" in investor["required"] and investor["properties"]["country"]["x-required-reason"] == "Set by you"


# ------------------------------------------------------------------ running on the batch, the review, export, fill

#: A run's documents: the five samples, the corpus folder's 22 PDFs (20 client forms and 2 copies of the W-9, which
#: no template fits) and the three problem copies (``PROBLEM_FILES``: a bad tax ID, none, a country not in the map).
RUN_TOTAL = 5 + 22 + 3


@pytest.fixture
def run_files(corpus_files, problem_files) -> list[tuple[str, bytes]]:
    return FILES + corpus_files + problem_files


def saved_run(base: str, files: list[tuple[str, bytes]], edits: list[dict] = (), **spec) -> tuple[dict, str]:
    """(the setup, its run's id): a setup of the batch ``files`` (its first five the sample), with ``edits`` made on
    it, saved as "Client info" and run on the whole batch through the API."""
    setup = setup_on(base, files, samples=list(range(5)), **spec)
    for edit in edits:
        setup = edit_setup(base, setup, edit)
    status, out = get_json(base, "/api/extractions/Client%20info", "PUT", {"setup": setup["id"]})
    assert status == 200, out
    status, out = start_job(base, {"kind": "extract-run", "extraction": "Client info", "batch": setup["batch"]})
    assert status == 200, out
    job = job_ended(base, out["id"])
    assert job["status"] == "done", job
    return setup, job["result"]["run"]


def show_review(b: Bench, run_id: str) -> None:
    """Open the run's review, as Run on all does once the run is done."""
    b.page.evaluate("id => window.__workbench.openRun(id)", run_id)
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible()
    expect(b.page.locator("#docs li").first).to_be_visible()


def stat(b: Bench, key: str):
    """The number of the review's stat ``key`` (total, complete, with_problems, optional_only, set_aside, time)."""
    return b.page.locator(f'#stats .stat[data-stat="{key}"] b')


def section(b: Bench, level: str):
    """The review's switch button of the section ``level`` (required or optional)."""
    return b.page.locator(f'#level-pick [data-level="{level}"]')


def listed(b: Bench, name: str):
    """The review list's document ``name``."""
    return b.page.locator(f'#docs li[data-name="{name}"]')


def card(b: Bench, path: str):
    """The problem card of ``path`` in the document shown."""
    return b.page.locator(f'#problems .problem[data-path="{path}"]')


def problem_doc(base: str, run_id: str, name: str) -> dict:
    """The document ``name`` of the run's problems, as the API lists it."""
    documents = get_json(base, f"/api/runs/{run_id}/problems")[1]["documents"]
    return next(d for d in documents if d["name"] == name)


def page_loaded(b: Bench) -> None:
    b.page.wait_for_function("(() => { const i = document.querySelector('#rpage img'); return i?.complete && i.naturalWidth > 0; })()")


def gate_at_three(monkeypatch) -> tuple[threading.Event, threading.Event]:
    """(reached, go): the run job's reading stops after its third document until ``go`` is set."""
    reached, go = threading.Event(), threading.Event()
    real = extract_api.run_batch

    def gated(*args, on_progress, **kwargs):
        def progress(done, total):
            on_progress(done, total)
            if done == 3:
                reached.set()
                assert go.wait(30)

        return real(*args, on_progress=progress, **kwargs)

    monkeypatch.setattr(extract_api, "run_batch", gated)
    return reached, go


def test_run_on_all_saves_the_setup_first_reads_every_document_with_its_progress_and_opens_the_review(start_demo_server, bench, run_files, monkeypatch):
    reached, go = gate_at_three(monkeypatch)
    base = extract_server(start_demo_server)
    setup = setup_on(base, run_files, samples=list(range(5)))
    b = bench(base, fill=False)
    show_setup(b, setup)
    started = []
    b.page.on("request", lambda r: started.append(r.post_data_json) if r.method == "POST" and r.url.endswith("/api/jobs") else None)
    expect(b.page.locator("#run")).to_be_enabled()
    b.page.click("#run")
    expect(b.page.locator("#setup-save-box")).to_be_visible()  # never saved: the name first
    b.page.fill("#setup-save-name", "Client info")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#running")).to_be_visible()
    assert reached.wait(30)
    expect(b.page.locator("#run-stages li")).to_have_count(1)
    expect(b.page.locator("#run-stages li")).to_contain_text("Read every document")
    expect(b.page.locator("#run-stages li .detail")).to_have_text(f"3 of {RUN_TOTAL}")
    expect(b.page.locator("#run-readout")).to_have_text(f"3 / {RUN_TOTAL} documents")
    go.set()
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible(timeout=60_000)
    expect(b.page.locator("#running")).to_be_hidden()
    assert started == [{"kind": "extract-run", "extraction": "Client info", "batch": setup["batch"], "ai": False}]
    assert [e["name"] for e in get_json(base, "/api/extractions")[1]["extractions"]] == ["Client info"]
    [entry] = get_json(base, "/api/runs")[1]["runs"]
    summary = entry["summary"]
    # bad_tin and no_tin; mars's country is optional, so mars is complete with an optional problem
    assert summary["total"] == RUN_TOTAL and summary["with_problems"] == 2 and summary["no_template"] == 2
    expect(stat(b, "total")).to_have_text(str(RUN_TOTAL))
    expect(stat(b, "complete")).to_have_text(str(summary["complete"]))
    expect(stat(b, "with_problems")).to_have_text("2")
    expect(stat(b, "set_aside")).to_have_text("2")
    expect(stat(b, "time")).to_have_text(re.compile(r"^\d+:\d\d$"))
    shown = [int(stat(b, key).inner_text()) for key in ("complete", "with_problems", "set_aside")]
    assert sum(shown) == RUN_TOTAL
    expect(b.page.locator("#doc-pick")).to_have_text(f"Client forms · {RUN_TOTAL} documents · Client info")

    expect(b.page.locator("#back-setup")).to_have_text("← Go back to setup")
    b.page.click("#back-setup")  # back to the setup it was run from, and to its review again
    expect(b.page.locator('.view[data-only="setup"]')).to_be_visible()
    expect(b.page.locator("#run")).to_be_visible()
    expect(b.page.locator("#doc-pick")).to_contain_text("5 samples")
    b.page.click("#go-review")  # nothing changed: straight back, nothing asked
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible()
    expect(b.page.locator("#reprocess-box")).to_be_hidden()
    expect(b.page.locator("#export-xlsx")).to_be_visible()


def run_from_setup(base: str, b: Bench, files: list[tuple[str, bytes]]) -> dict:
    """Run a setup of the batch ``files`` from its screen, saved as "Client info", and show its review; the setup."""
    setup = setup_on(base, files, samples=list(range(5)))
    show_setup(b, setup)
    b.page.click("#run")
    b.page.fill("#setup-save-name", "Client info")
    b.page.click("#setup-save-go")
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible(timeout=60_000)
    return setup


def test_a_setup_changed_after_its_run_asks_before_reprocessing_and_makes_a_new_run(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    b = bench(base, fill=False)
    run_from_setup(base, b, run_files)
    expect(listed(b, "bad_tin.pdf")).to_be_visible()
    [first] = get_json(base, "/api/runs")[1]["runs"]
    card(b, "investor.tax_id").locator('[data-act="not_in_document"]').click()  # a decision to lose
    expect(card(b, "investor.tax_id").locator(".state")).to_have_text("Marked not in this document")

    b.page.click("#back-setup")
    row(b, "investor.country").click()
    b.page.fill("#ins-title", "Country code")
    b.page.locator("#ins-title").press("Enter")
    setup_idle(b)
    b.page.click("#go-review")
    box = b.page.locator("#reprocess-box")
    expect(box).to_be_visible()
    expect(box).to_contain_text("every document is read again")
    expect(b.page.locator("#reprocess-decisions")).to_have_text("The decision made in this review stays with the earlier run; it isn't carried over.")
    b.page.click("#reprocess-cancel")  # Keep editing: nothing is run
    expect(box).to_be_hidden()
    expect(b.page.locator('.view[data-only="setup"]')).to_be_visible()
    assert len(get_json(base, "/api/runs")[1]["runs"]) == 1

    b.page.click("#setup-save")  # saved: still changed since the run
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    b.page.click("#go-review")
    expect(box).to_be_visible()
    b.page.click("#reprocess-go")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()  # it has a name: no name asked
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible(timeout=60_000)
    expect(listed(b, "bad_tin.pdf")).to_be_visible()  # a new run: its decision stayed with the old one
    runs = get_json(base, "/api/runs")[1]["runs"]
    assert len(runs) == 2 and first["id"] in [r["id"] for r in runs]
    b.page.click("#back-setup")
    b.page.click("#go-review")  # as this run read it: nothing asked
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible()


def test_a_saved_run_goes_back_to_its_templates_setup_and_on_to_its_review(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    b.page.click("#back-setup")
    expect(b.page.locator('.view[data-only="setup"]')).to_be_visible(timeout=60_000)
    expect(b.page.locator("#doc-pick")).to_contain_text("5 samples")
    b.page.click("#go-review")
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible()
    expect(listed(b, "bad_tin.pdf")).to_be_visible()


def test_a_problem_picked_goes_to_its_place_on_the_page_with_what_is_wrong_there(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    page_loaded(b)
    tag = b.page.locator("#rpage .rmark.is-picked .rtag")
    expect(tag).to_have_text("Fails a check")  # bad_tin: read here, in doubt
    listed(b, "no_tin.pdf").click()
    mark = b.page.locator('#rpage .rmark.is-picked.is-missing[data-path="investor.tax_id"]')
    expect(mark).to_be_visible()  # nothing was read: the server says where it should be
    expect(tag).to_have_text("Should be here")
    expect(b.page.locator("#rcap")).to_have_text("The boxes are where each problem field was read, or looked for")

    b.page.click("#rnext-p")  # turned away; picking it again goes back to it
    expect(b.page.locator("#rpage-of")).to_have_text("Page 2")
    expect(mark).to_have_count(0)
    card(b, "investor.tax_id").locator(".head").click()
    expect(b.page.locator("#rpage-of")).to_have_text("Page 1")
    expect(mark).to_be_visible()


def test_the_side_columns_grow_with_the_window_and_never_go_narrower_than_240_and_300(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    left = b.page.locator('.view[data-only="review"] .col-schema')
    right = b.page.locator('.view[data-only="review"] .col-side')

    def widths(width: int) -> tuple[float, float]:
        b.page.set_viewport_size({"width": width, "height": 900})
        return left.bounding_box()["width"], right.bounding_box()["width"]

    assert widths(1100) == (240, 300)
    assert widths(1440) == (pytest.approx(302.4, abs=1), pytest.approx(374.4, abs=1))
    assert widths(1920) == (320, 400)


def test_results_by_field_gives_its_columns_room(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    b.page.set_viewport_size({"width": 1440, "height": 900})
    show_review(b, run_id)
    b.page.click("#see-details")
    sheet = b.page.locator("#field-details .sheet")
    expect(sheet).to_be_visible()
    assert sheet.bounding_box()["width"] > 1300  # it used to stop at 1040
    gaps = b.page.evaluate("""() => {
      const head = document.querySelector('#details-sections .field-table thead tr');
      const range = document.createRange();
      const right = (th) => { range.selectNodeContents(th); return range.getBoundingClientRect().right; };
      const left = (th) => { range.selectNodeContents(th); return range.getBoundingClientRect().left; };
      const cells = [...head.children].slice(3);
      return cells.slice(1).map((th, i) => left(th) - right(cells[i]));
    }""")
    assert min(gaps) >= 16  # no heading runs up against the next
    table = b.page.locator("#details-sections .field-table").first.bounding_box()
    assert table["width"] <= sheet.bounding_box()["width"]  # it fits with no sideways scroll


def test_run_on_all_of_a_saved_setup_asks_no_name_and_a_cancel_opens_what_was_read(start_demo_server, bench, run_files, monkeypatch):
    reached, go = gate_at_three(monkeypatch)
    base = extract_server(start_demo_server)
    setup = setup_on(base, run_files, samples=list(range(5)))
    b = bench(base, fill=False)
    show_setup(b, setup)
    b.page.click("#setup-save")
    b.page.fill("#setup-save-name", "Client info")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    b.page.click("#run")
    expect(b.page.locator("#running")).to_be_visible()
    expect(b.page.locator("#setup-save-box")).to_be_hidden()  # saved and unchanged: no name asked
    assert reached.wait(30)
    b.page.click("#run-cancel")
    expect(b.page.locator("#run-readout")).to_contain_text("Stopping")
    go.set()
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible(timeout=60_000)
    expect(b.page.locator("#notice-text")).to_contain_text(f"stopped after 3 of {RUN_TOTAL} documents")
    expect(stat(b, "total")).to_have_text("3")


def test_run_on_all_right_after_an_edit_asks_for_the_name_again_and_closing_it_starts_nothing(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    setup = setup_on(base, run_files, samples=list(range(5)))
    b = bench(base, fill=False)
    show_setup(b, setup)
    b.page.click("#setup-save")
    b.page.fill("#setup-save-name", "Client info")
    b.page.click("#setup-save-go")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    row(b, "investor.tax_id").click()
    b.page.fill("#ins-title", "Tax ID")
    b.page.click("#run")  # the title's edit is sent as the click takes the focus: the saved template is out of date
    expect(b.page.locator("#setup-save-box")).to_be_visible()
    expect(b.page.locator("#setup-save-name")).to_have_value("Client info")
    b.page.click("#setup-save-cancel")
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    expect(b.page.locator("#running")).to_be_hidden()
    assert get_json(base, "/api/runs")[1] == {"runs": []}


def test_the_review_lists_the_documents_with_problems_and_filters_them_by_field_problem_or_set_aside(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    names = b.page.locator("#docs li .doc-name")
    expect(b.page.locator("#docs-label")).to_have_text("Documents · 2")  # mars's country is optional
    expect(names).to_have_text(["bad_tin.pdf", "no_tin.pdf"])
    expect(listed(b, "no_tin.pdf").locator("small")).to_have_text("1 problem: Taxpayer ID")
    expect(b.page.locator("#docs-legend")).to_have_text("2 of 2 still have problems")
    expect(b.page.locator("#filter-path option").first).to_have_text("Every field")
    expect(b.page.locator("#filter-path option[value='investor.tax_id']")).to_have_count(1)

    b.page.select_option("#filter-status", "blank")
    expect(names).to_have_text(["no_tin.pdf"])
    expect(b.page.locator("#docs-label")).to_have_text("Documents · 1")
    expect(card(b, "investor.tax_id")).to_contain_text("Blank")
    b.page.select_option("#filter-status", "")
    b.page.select_option("#filter-path", "investor.tax_id")
    expect(names).to_have_text(["bad_tin.pdf", "no_tin.pdf"])

    b.page.select_option("#filter-status", "set_aside")  # the documents not read, with why, in place of problems
    expect(b.page.locator("#filter-path")).to_be_disabled()
    expect(b.page.locator("#docs-label")).to_have_text("Documents · 2")
    expect(names).to_have_text(["client_00021.pdf", "client_00022.pdf"])
    expect(listed(b, "client_00021.pdf").locator("small")).to_contain_text("no template fits · closest: Client info, ")
    expect(b.page.locator("#rdoc-name")).to_have_text("client_00021.pdf")
    expect(b.page.locator("#problems")).to_contain_text("no template fits")
    expect(b.page.locator("#problems")).to_contain_text("Extract from PDFs again")
    page_loaded(b)


def test_see_details_opens_each_fields_results_to_read_and_close(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    fields = get_json(base, f"/api/runs/{run_id}")[1]["fields"]
    keys = ("found", "blank", "not_located", "failed_check", "unsure", "not_on_form", "absent", "reviewed")
    folded: dict[str, dict] = {}  # the run's fields folded as the table folds them
    for f in fields:
        row = folded.setdefault(re.sub(r"\[\d+\]", "[]", f["path"]), {k: 0 for k in keys})
        for k in keys:
            row[k] += f[k]
    expect(b.page.locator("#by-field")).to_have_count(0)  # the results are in the modal, not on the page
    expect(b.page.locator("#stats #see-details")).to_have_text("See details")
    details = b.page.locator("#field-details")
    expect(details).to_be_hidden()

    b.page.click("#see-details")
    expect(details).to_be_visible()
    expect(b.page.locator("#details-close")).to_be_focused()
    expect(b.page.locator("#details-sections .field-table")).to_have_count(2)
    expect(b.page.locator("#details-key li")).to_have_text(["Found", "Blank", "Not located", "Doubtful", "Fails a check", "Not applicable"])
    expect(b.page.locator("#details-sections .field-table").first.locator("thead th")).to_have_text(
        ["Field", "Results", "Found %", "Found", "Blank", "Not located", "Doubtful", "Fails a check", "Not applicable", "Reviewed"]
    )
    expect(b.page.locator("#details-sections h3")).to_have_text(["Required fields · 1", "Optional fields · " + str(len(folded) - 1)])
    expect(b.page.locator("#details-sections tbody tr")).to_have_count(len(folded))
    tax = b.page.locator('#details-sections tr[data-path="investor.tax_id"]')
    counts = folded["investor.tax_id"]
    expect(tax.locator("td").nth(2)).to_have_text(str(counts["found"]))
    expect(tax.locator("td").nth(3)).to_have_text(str(counts["blank"]))
    expect(tax.locator("td").nth(4)).to_have_text(str(counts["not_located"]))
    expect(tax.locator("td").nth(6)).to_have_text(str(counts["failed_check"]))
    expect(tax.locator("td").nth(8)).to_have_text(str(counts["reviewed"]))
    expect(tax.locator("th")).to_contain_text("Taxpayer ID")
    expect(tax.locator(".value-bar")).to_have_attribute("aria-label", re.compile(rf"Found {counts['found']}"))
    repeated = next(path for path in folded if "[]" in path)  # the owners: a row for all their items
    items = sum(1 for f in fields if re.sub(r"\[\d+\]", "[]", f["path"]) == repeated)
    assert items > 1
    expect(b.page.locator(f'#details-sections tr[data-path="{repeated}"] th small').first).to_have_text(f"{items} items")
    expect(b.page.locator("#details-tiles .stat[data-stat='fields'] b")).to_have_text(str(len(folded)))

    tax.click()  # only for reading: nothing is filtered, and the modal stays
    expect(details).to_be_visible()
    expect(b.page.locator("#filter-path")).to_have_value("")
    b.page.click("#details-close")
    expect(details).to_be_hidden()
    expect(b.page.locator("#see-details")).to_be_focused()
    b.page.click("#see-details")
    b.page.keyboard.press("Escape")
    expect(details).to_be_hidden()
    b.page.click("#see-details")
    details.click(position={"x": 2, "y": 2})  # outside the sheet
    expect(details).to_be_hidden()

    listed_after = b.page.locator("#docs li .doc-name")
    b.page.select_option("#filter-path", repeated)  # a field picked shows the section it is in
    expect(section(b, "optional")).to_have_attribute("aria-pressed", "true")
    b.page.select_option("#filter-path", "")
    expect(listed_after).to_have_text(["mars.pdf"])
    section(b, "required").click()
    expect(listed_after).to_have_text(["bad_tin.pdf", "no_tin.pdf"])
    listed(b, "no_tin.pdf").click()
    card(b, "investor.tax_id").locator("input").fill("12-3456789")
    card(b, "investor.tax_id").locator('[data-act="value"]').click()  # a decision: the modal follows
    expect(card(b, "investor.tax_id")).to_contain_text("12-3456789")
    b.page.click("#see-details")
    expect(tax.locator("td").nth(8)).to_have_text(str(counts["reviewed"] + 1))
    expect(tax.locator("td").nth(3)).to_have_text(str(counts["blank"] - 1))
    b.page.click("#details-close")
    expect(b.page.locator("#filter-missing")).to_be_hidden()  # only a run made before Blank and Not located has it


def test_a_page_image_that_fails_past_the_first_page_stops_the_pages_only_when_it_is_not_there(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    page_loaded(b)
    failing = re.compile(rf"/api/runs/{run_id}/pages\.png\?.*page=1&")
    b.page.route(failing, lambda route: route.fulfill(status=500, content_type="application/json", body='{"error": "server", "detail": "the server had a problem"}'))
    b.page.click("#rnext-p")  # page 2 fails for another reason than not being there: say so, and keep the pages
    expect(b.page.locator("#rcap")).to_contain_text("This page couldn't be shown: the server had a problem.")
    expect(b.page.locator("#rcap button")).to_have_text("Try again")
    expect(b.page.locator("#rpage-of")).to_have_text("Page 2")
    expect(b.page.locator("#rnext-p")).to_be_enabled()  # the document was not cut short
    b.page.unroute(failing)
    b.page.click("#rcap button")  # again: the page is there now (the form has two)
    expect(b.page.locator("#rpage-of")).to_have_text("Page 2")
    page_loaded(b)
    expect(b.page.locator("#rcap")).to_have_text("The boxes are where each problem field was read, or looked for")
    b.page.click("#rnext-p")  # page 3 is not there (404): the document ends at page 2
    expect(b.page.locator("#rpage-of")).to_have_text("Page 2")
    expect(b.page.locator("#rnext-p")).to_be_disabled()
    page_loaded(b)


def test_the_choices_message_that_the_lists_could_not_be_read_goes_once_they_are_listed(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    saved_run(base, run_files)
    b = bench(base, fill=False)
    expect(b.page.locator("#saved-runs li")).to_have_count(1)
    runs = re.compile(r"/api/runs$")
    b.page.route(runs, lambda route: route.fulfill(status=500, content_type="application/json", body='{"error": "server", "detail": "the server had a problem"}'))
    open_extract(b)
    b.page.click("#extract-back")  # the choice again: its lists can't be read
    expect(b.page.locator("#choose-msg")).to_contain_text("The saved runs and extractions couldn't be listed")
    b.page.unroute(runs)
    open_extract(b)
    b.page.click("#extract-back")  # and now they can: the message goes
    expect(b.page.locator("#choose-msg")).to_be_hidden()
    expect(b.page.locator("#saved-runs li")).to_have_count(1)


def test_choosing_a_document_shows_its_page_with_the_problems_place_outlined_and_its_cards(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    expect(listed(b, "bad_tin.pdf")).to_have_class(re.compile(r"\bon\b"))  # the first is shown
    listed(b, "no_tin.pdf").click()
    expect(listed(b, "no_tin.pdf")).to_have_class(re.compile(r"\bon\b"))
    expect(b.page.locator("#rdoc-name")).to_have_text("no_tin.pdf")
    listed(b, "bad_tin.pdf").click()
    doc = problem_doc(base, run_id, "bad_tin.pdf")
    [problem] = doc["problems"]
    assert problem["path"] == "investor.tax_id" and problem["rect"] and problem["page"] == 0
    img = b.page.locator("#rpage img")
    expect(img).to_have_attribute("src", re.compile(rf"/api/runs/{run_id}/pages\.png\?doc={re.escape(doc['doc_id'])}&page=0&dpi=\d+$"))
    page_loaded(b)
    mark = b.page.locator('#rpage .rmark.is-picked[data-path="investor.tax_id"]')
    expect(mark).to_be_visible()
    width_pt = b.page.evaluate("() => { const i = document.querySelector('#rpage img'); return i.naturalWidth * 72 / Number(new URL(i.src).searchParams.get('dpi')); }")
    box, page = mark.bounding_box(), img.bounding_box()
    scale = page["width"] / width_pt
    x0, y0, x1, _ = problem["rect"]
    assert abs(box["x"] - page["x"] - x0 * scale) < 3 and abs(box["y"] - page["y"] - y0 * scale) < 3
    assert abs(box["width"] - (x1 - x0) * scale) < 3
    expect(b.page.locator("#rcap")).to_have_text("The boxes are where each problem field was read, or looked for")
    c = card(b, "investor.tax_id")
    expect(c).to_contain_text("Taxpayer ID · Fails a check")
    expect(c).to_contain_text("Read: 8422-19037")
    expect(c.locator("input")).to_have_value("8422-19037")
    expect(c.locator('[data-act="accept"]')).to_be_visible()  # a value was read
    listed(b, "no_tin.pdf").click()
    expect(card(b, "investor.tax_id").locator('[data-act="accept"]')).to_have_count(0)  # nothing was read to accept
    expect(card(b, "investor.tax_id")).to_contain_text("Read: nothing")

    b.page.click("#review-zoom [data-fit='page']")  # the page's zoom, as the setup screen's
    expect(b.page.locator("#review-zoom [data-fit='page']")).to_have_attribute("aria-pressed", "true")
    b.page.click("#review-zoom [data-step='1']")
    expect(b.page.locator("#review-zoom [data-fit='page']")).to_have_attribute("aria-pressed", "false")


def test_each_decision_settles_its_card_and_updates_the_stats_and_the_list(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    complete = int(stat(b, "complete").inner_text())
    listed(b, "no_tin.pdf").click()
    c = card(b, "investor.tax_id")
    c.locator('[data-act="value"]').click()  # nothing typed: nothing sent
    expect(c.locator("input")).to_be_focused()
    c.locator("input").fill("12-3456789")
    c.locator('[data-act="value"]').click()
    expect(c.locator(".state")).to_have_text("Value entered: 12-3456789")
    expect(c.locator("input")).to_have_count(0)
    expect(stat(b, "with_problems")).to_have_text("1")  # bad_tin; mars's country is optional
    expect(stat(b, "complete")).to_have_text(str(complete + 1))
    expect(listed(b, "no_tin.pdf").locator("small")).to_have_text("all settled")
    expect(b.page.locator("#docs-legend")).to_have_text("1 of 2 still have problems")
    assert get_json(base, f"/api/runs/{run_id}/problems")[1]["total"] == 2  # bad_tin, and mars of either level

    listed(b, "bad_tin.pdf").click()
    card(b, "investor.tax_id").locator('[data-act="accept"]').click()
    expect(card(b, "investor.tax_id").locator(".state")).to_have_text("Accepted as read: 8422-19037")
    expect(stat(b, "with_problems")).to_have_text("0")
    expect(b.page.locator("#docs-legend")).to_have_text("0 of 2 still have problems")
    section(b, "optional").click()
    listed(b, "mars.pdf").click()
    card(b, "investor.country").locator('[data-act="not_in_document"]').click()
    expect(card(b, "investor.country").locator(".state")).to_have_text("Marked not in this document")
    expect(stat(b, "optional_only")).to_have_text("0")
    expect(b.page.locator("#docs-legend")).to_have_text("0 of 1 still have optional problems")
    run = get_json(base, f"/api/runs/{run_id}")[1]
    assert run["summary"]["with_problems"] == 0 and run["summary"]["with_optional"] == 0
    assert next(f for f in run["fields"] if f["path"] == "investor.tax_id")["reviewed"] == 2

    b.page.select_option("#filter-status", "blank")  # the list asked again: nothing is left to settle
    expect(b.page.locator("#docs li")).to_have_text(["No document has a problem of this kind."])


def test_review_shows_required_problems_first_and_optional_ones_in_their_own_section(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files, edits=REASONS)
    b = bench(base, fill=False)
    show_review(b, run_id)
    names = b.page.locator("#docs li .doc-name")
    expect(b.page.locator("#level-pick [data-level]")).to_have_text(["Required · 2", "Optional · 1"])
    expect(section(b, "required")).to_have_attribute("aria-pressed", "true")  # the Required section first
    expect(names).to_have_text(["bad_tin.pdf", "no_tin.pdf"])  # mars misses only an optional field: not here
    expect(stat(b, "with_problems")).to_have_text("2")
    expect(stat(b, "optional_only")).to_have_text("1")
    expect(b.page.locator("#stats .stat span")).to_have_text(["Documents", "Complete", "Required problems", "Optional only", "Set aside", "Run time"])

    section(b, "optional").click()
    expect(section(b, "optional")).to_have_attribute("aria-pressed", "true")
    expect(names).to_have_text(["mars.pdf"])
    expect(b.page.locator("#docs-label")).to_have_text("Documents · 1")
    expect(b.page.locator("#docs-legend")).to_have_text("1 of 1 still have optional problems")
    expect(b.page.locator("#optional-note")).to_have_count(0)  # no note above the Optional section
    expect(b.page.locator("#rdoc-name")).to_have_text("mars.pdf")
    expect(b.page.locator("#problems .problem")).to_have_count(1)  # only this section's cards
    why = card(b, "investor.country").locator(".why")
    expect(why).to_have_text("Optional: a copy may leave it empty")
    expect(why).to_be_hidden()  # behind the info icon
    expect(card(b, "investor.country").locator(".info")).to_have_attribute("title", re.compile("Optional: a copy may leave it empty"))
    card(b, "investor.country").locator(".info").click()
    expect(why).to_be_visible()
    expect(card(b, "investor.country").locator('[data-act="value"]')).to_have_attribute("title", "Use this value")
    expect(card(b, "investor.country").locator('[data-act="value"]')).to_have_attribute("aria-label", "Use this value")
    card(b, "investor.country").locator(".info").click()
    expect(why).to_be_hidden()

    b.page.select_option("#filter-status", "blank")  # the filters apply inside the section
    expect(b.page.locator("#docs li")).to_have_text(["No document has a problem of this kind."])
    b.page.select_option("#filter-status", "failed_check")
    expect(names).to_have_text(["mars.pdf"])
    b.page.select_option("#filter-status", "")
    b.page.select_option("#filter-path", "investor.country")
    expect(names).to_have_text(["mars.pdf"])
    section(b, "required").click()
    expect(b.page.locator("#docs li")).to_have_text(["No document has a problem of this kind."])
    section(b, "optional").click()
    b.page.select_option("#filter-path", "investor.tax_id")  # a field picked shows the section it is in
    expect(section(b, "required")).to_have_attribute("aria-pressed", "true")
    expect(names).to_have_text(["bad_tin.pdf", "no_tin.pdf"])
    b.page.select_option("#filter-path", "")
    b.page.select_option("#filter-status", "blank")
    expect(names).to_have_text(["no_tin.pdf"])
    expect(card(b, "investor.tax_id").locator(".why")).to_have_count(0)  # a required card gives no reason
    b.page.select_option("#filter-status", "")

    listed(b, "no_tin.pdf").click()  # a required card decided: both counts follow
    card(b, "investor.tax_id").locator("input").fill("12-3456789")
    card(b, "investor.tax_id").locator('[data-act="value"]').click()
    expect(b.page.locator("#level-pick [data-level]")).to_have_text(["Required · 1", "Optional · 1"])
    expect(stat(b, "with_problems")).to_have_text("1")
    section(b, "optional").click()
    card(b, "investor.country").locator('[data-act="not_in_document"]').click()  # and an optional one
    expect(b.page.locator("#level-pick [data-level]")).to_have_text(["Required · 1", "Optional · 0"])
    expect(stat(b, "optional_only")).to_have_text("0")


def test_more_after_a_document_is_settled_lists_every_remaining_problem_document_once(start_demo_server, bench, run_files):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    b.page.evaluate("window.__workbench.reviewPageSize = 1")  # the Required section lists two
    show_review(b, run_id)
    names = b.page.locator("#docs li .doc-name")
    expect(names).to_have_text(["bad_tin.pdf"])
    expect(b.page.locator("#docs-more")).to_be_visible()
    card(b, "investor.tax_id").locator('[data-act="accept"]').click()
    expect(listed(b, "bad_tin.pdf").locator("small")).to_have_text("all settled")
    expect(b.page.locator("#docs-legend")).to_have_text("1 of 2 still have problems")
    b.page.click("#docs-more")  # the server's list no longer holds bad_tin: the next page starts one earlier
    expect(names).to_have_text(["bad_tin.pdf", "no_tin.pdf"])
    expect(b.page.locator("#docs-more")).to_be_hidden()
    expect(b.page.locator("#docs-label")).to_have_text("Documents · 2")
    expect(b.page.locator("#docs-legend")).to_have_text("1 of 2 still have problems")


def test_a_decision_draws_only_its_own_card_again_and_a_card_stays_disabled_while_its_decision_is_on_its_way(start_demo_server, bench):
    both = dataclasses.replace(clients(1, seed=5)[0], tax_id="8422-19037", country="Mars")  # two problems in one
    base = extract_server(start_demo_server)
    country_required = {"op": "set_field", "path": "investor.country", "required": True}  # both cards in one section
    _, run_id = saved_run(base, [*FILES, ("both.pdf", make_form(both, kind="flattened"))], edits=[country_required])
    b = bench(base, fill=False)
    show_review(b, run_id)
    expect(b.page.locator("#rdoc-name")).to_have_text("both.pdf")
    tax, country = card(b, "investor.tax_id"), card(b, "investor.country")
    country.locator("input").fill("US")
    held = []
    b.page.route("**/decisions", lambda route: held.append(route), times=1)
    tax.locator('[data-act="not_in_document"]').click()
    for _ in range(200):  # the route handler runs while the page is driven
        if held:
            break
        b.page.wait_for_timeout(50)
    assert len(held) == 1
    expect(tax.locator('[data-act="value"]')).to_be_disabled()  # on its way
    expect(country.locator('[data-act="value"]')).to_be_enabled()
    held[0].continue_()
    expect(tax.locator(".state")).to_have_text("Marked not in this document")
    expect(country.locator("input")).to_have_value("US")  # the other card was not drawn again
    expect(country.locator("input")).to_be_focused()  # focus moves on to the next open card
    assert b.page.evaluate("document.getElementById('problems').contains(document.activeElement)")

    b.page.route("**/decisions", lambda route: held.append(route), times=1)
    country.locator('[data-act="value"]').click()
    for _ in range(200):
        if len(held) == 2:
            break
        b.page.wait_for_timeout(50)
    assert len(held) == 2
    listed(b, "both.pdf").click()  # the cards drawn again meanwhile: the one deciding stays disabled
    expect(country.locator('[data-act="value"]')).to_be_disabled()
    expect(country.locator("input")).to_be_disabled()
    held[1].continue_()
    expect(country.locator(".state")).to_have_text("Value entered: US")
    expect(stat(b, "with_problems")).to_have_text("0")


def test_export_xlsx_downloads_the_runs_workbook_and_export_parquet_needs_pyarrow(start_demo_server, bench, run_files):
    import openpyxl

    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    with b.page.expect_download() as download:
        b.page.click("#export-xlsx")
    assert download.value.suggested_filename == "Client forms-results.xlsx"
    book = openpyxl.load_workbook(io.BytesIO(download.value.path().read_bytes()), read_only=True)
    assert book.sheetnames[0] == "values"
    assert len(list(book["values"].iter_rows(values_only=True))) == 1 + RUN_TOTAL - 2  # each document read
    if importlib.util.find_spec("pyarrow") is None:
        b.page.click("#export-parquet")
        expect(b.page.locator("#export-parquet")).to_be_hidden()  # the server says it can't: the button goes
        expect(b.page.locator("#notice-text")).to_contain_text("Parquet isn't available on this server")
        expect(b.page.locator("#notice-text")).to_contain_text("formcast[parquet]")
    else:
        with b.page.expect_download() as download:
            b.page.click("#export-parquet")
        assert download.value.suggested_filename == "Client forms-results.zip"


def test_send_to_fill_opens_the_start_screen_with_the_runs_results_as_your_data_and_says_what_is_left_blank(start_demo_server, bench, run_files, monkeypatch):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)
    show_review(b, run_id)
    b.page.click("#send-fill")
    expect(b.page.locator("#start")).to_be_visible()
    assert b.page.evaluate("document.body.dataset.mode") == "fill"
    expect(b.page.locator("#data-given")).to_contain_text("Client forms results")
    note = b.page.locator("#data-given-note")  # on the start sheet itself, not under its overlay
    expect(note).to_be_visible()
    # bad_tin and no_tin; mars's country is optional, so its row is complete
    expect(note).to_have_text(
        "2 documents still have problems; their unsettled fields are left blank, and fill will report them as gaps."
    )
    expect(b.page.locator("#notice")).to_be_hidden()
    expect(b.page.locator("#form-pick")).to_have_value("")  # an AI-suggested template names no destination

    monkeypatch.setattr(extraction, "MAX_ROWS", 5)  # more documents than a dataset holds: it says so
    b.page.click("#start-close")  # no form open: back to the choice, and the run's review from its list
    expect(b.page.locator("#choose")).to_be_visible()
    b.page.locator("#saved-runs li button").first.click()
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible()
    b.page.click("#send-fill")
    expect(b.page.locator("#start")).to_be_visible()
    expect(note).to_be_visible()
    expect(note).to_have_text(
        "Only the first 5 documents are in the data, as many as a dataset holds; export the run for every one."
    )
    b.page.click("#data-given-clear")  # the note goes with the data it is about
    expect(note).to_be_hidden()


def test_send_to_fill_of_a_run_made_for_a_destination_picks_that_form(start_demo_server, bench):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, FILES, start="destination", destination="fw9")
    b = bench(base, fill=False)
    show_review(b, run_id)
    b.page.click("#send-fill")
    expect(b.page.locator("#start")).to_be_visible()
    expect(b.page.locator("#form-pick")).to_have_value("saved:fw9")
    expect(b.page.locator("#data-given")).to_contain_text("Client forms results")


def test_the_choice_lists_the_saved_runs_and_extractions_and_opens_them_or_a_run_folder_on_the_server(start_demo_server, bench, run_files, tmp_path):
    base = extract_server(start_demo_server)
    _, run_id = saved_run(base, run_files)
    b = bench(base, fill=False)  # the page opened again later: the run is on the server
    runs = b.page.locator("#saved-runs li")
    expect(runs).to_have_count(1)
    # mars is complete: its country is optional
    expect(runs).to_contain_text(f"Client forms · Client info · {RUN_TOTAL} documents · 2 with problems")
    expect(b.page.locator("#saved-extractions li")).to_contain_text(["Client info"])
    runs.first.locator("button").click()
    expect(b.page.locator("#choose")).to_be_hidden()
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible()
    assert b.page.evaluate("document.body.dataset.mode") == "extract"
    expect(stat(b, "total")).to_have_text(str(RUN_TOTAL))
    expect(b.page.locator("#docs li .doc-name")).to_have_text(["bad_tin.pdf", "no_tin.pdf"])  # the Required section
    expect(b.page.locator("#back-setup")).to_be_visible()  # no setup of it is open: its template reopens

    b.page.click("#doc-pick")
    expect(b.page.locator("#choose")).to_be_visible()
    b.page.fill("#run-folder", str(tmp_path))  # a folder with no run in it
    b.page.click("#open-run-folder")
    expect(b.page.locator("#choose-msg")).to_contain_text("There is no run in that folder")
    folder = start_demo_server.servers[-1].api.settings.data_dir / "local" / "runs" / run_id  # a run's own folder
    b.page.fill("#run-folder", str(folder))
    b.page.click("#open-run-folder")
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible()
    expect(b.page.locator("#choose-msg")).to_be_hidden()
    expect(stat(b, "total")).to_have_text(str(RUN_TOTAL))

    b.page.click("#doc-pick")
    b.page.locator("#saved-extractions li button").first.click()  # a saved extraction opens on the setup screen
    expect(b.page.locator('.view[data-only="setup"]')).to_be_visible(timeout=30_000)
    expect(b.page.locator("#doc-pick")).to_contain_text("Client info")
    expect(b.page.locator("#run")).to_be_enabled()  # it has no batch: Run on all asks for one


def test_a_saved_extraction_reopened_from_the_choice_runs_on_a_new_batch_and_its_review_follows(start_demo_server, bench, corpus_files, problem_files):
    base = extract_server(start_demo_server)
    setup = setup_on(base)
    assert get_json(base, "/api/extractions/Client%20info", "PUT", {"setup": setup["id"]})[0] == 200
    b = bench(base, fill=False)
    started = []
    b.page.on("request", lambda r: started.append(r.post_data_json) if r.method == "POST" and r.url.endswith("/api/jobs") else None)
    b.page.locator("#saved-extractions li button").first.click()
    expect(b.page.locator('.view[data-only="setup"]')).to_be_visible(timeout=30_000)
    b.page.click("#run")  # saved and unchanged: no name asked, the documents to read with it
    sheet = b.page.locator("#extract")
    expect(sheet).to_be_visible()
    expect(b.page.locator("#setup-save-box")).to_be_hidden()
    expect(b.page.locator("#extract-title")).to_have_text("Run Client info on a batch")
    expect(sheet.locator(".group > .label:visible")).to_have_text(["1 · Your documents"])  # the documents only
    expect(b.page.locator("#read-sample")).to_be_hidden()
    run_these = b.page.locator("#run-these")
    expect(run_these).to_have_text("Run on these →")
    expect(run_these).to_be_disabled()  # no documents yet

    b.page.click("#extract-back")  # back to the setup screen it came from
    expect(sheet).to_be_hidden()
    expect(b.page.locator("#choose")).to_be_hidden()
    expect(b.page.locator('.view[data-only="setup"]')).to_be_visible()
    b.page.click("#run")
    expect(sheet).to_be_visible()

    b.page.set_input_files("#drop-files", files=pdf_files(corpus_files + problem_files))
    expect(b.page.locator("#batch-sum")).to_contain_text("25 PDFs")
    expect(run_these).to_be_enabled(timeout=30_000)  # scanned
    run_these.click()
    expect(b.page.locator('.view[data-only="review"]')).to_be_visible(timeout=60_000)
    expect(sheet).to_be_hidden()
    [entry] = get_json(base, "/api/runs")[1]["runs"]
    batch_id = started[-1]["batch"]
    assert started == [{"kind": "extract-run", "extraction": "Client info", "batch": batch_id, "ai": False}]
    assert batch_id != setup["batch"] and (entry["extraction"], entry["summary"]["total"]) == ("Client info", 25)
    expect(stat(b, "total")).to_have_text("25")
    expect(stat(b, "with_problems")).to_have_text("2")  # mars's country is optional
    expect(b.page.locator("#docs li .doc-name")).to_have_text(["bad_tin.pdf", "no_tin.pdf"])  # the Required section
    b.page.click("#back-setup")  # the setup it was run from is on screen
    expect(b.page.locator('.view[data-only="setup"]')).to_be_visible()

    b.page.click("#doc-pick")  # Extract from PDFs afterwards builds a template again
    open_extract(b)
    expect(b.page.locator("#extract-title")).to_have_text("Extract from PDFs")
    expect(b.page.locator("#read-sample")).to_be_visible()
    expect(run_these).to_be_hidden()
