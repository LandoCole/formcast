"""demo/server extraction: task jobs, the two settings, several files per upload, ``page_png``, batches of PDFs,
setups, saved extraction templates and runs.

Task jobs run a caller's function on a job thread (``Jobs.start_task``); the extraction jobs of the workbench are
built on them. Batches are PDFs uploaded (singly or in zips) or found in a server-side folder, scanned by layout on
a thread, with a paged file list, a sample picked from the main layout and page images. A setup is a template built
from a batch's sample by the ``extract-setup`` job, edited and proven again with no model call, grown by the
``extract-line`` job, and saved as an extraction template (with its sample PDFs) that reopens as a setup. A run
reads a batch with a saved extraction (the ``extract-run`` job, resumed after a cancel), and its problems, set-aside
documents, decisions, page images, export and dataset are served from its folder; a command-line run is opened by
its folder. The PDFs
are the generated client-information corpus (``tests/extract/corpus.py``); the model is ``CorpusModel``, which
answers from the corpus's own truth. Everything runs in process: no network, and no detector weights are loaded.
"""

from __future__ import annotations

import dataclasses
import hashlib
import importlib.util
import io
import json
import os
import subprocess
import sys
import threading
import time
import zipfile
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import quote

import pymupdf
import pytest
from test_server_documents import get_json, http_request, png_size, small_pdf
from test_server_jobs import wait_for

from demo.server import documents as documents_mod
from demo.server import extract_api, extraction
from demo.server.api import Api, ApiError, Request, _uploads
from demo.server.documents import PYMUPDF_LOCK, page_png
from demo.server.jobs import Jobs, JobsBusy
from demo.server.settings import Settings
from demo.server.store import StoredTemplateError
from formcast import FormcastError, SessionCancelled
from formcast.data.ai import DERIVE_SYSTEM
from formcast.extract import samples as samples_mod
from formcast.extract.batch import run_batch
from formcast.extract.model import ExtractionTemplate, SampleRecord, concrete_paths
from formcast.extract.proof import Proof, prove
from formcast.extract.prompts import FIELDS_SYSTEM, VALUES_SYSTEM
from formcast.extract.samples import SamplePdf
from formcast.extract.setup import STAGES
from formcast.pagemap.build import DetectPolicy
from tests.extract.conftest import write_corpus
from tests.extract.corpus import FIELD_RECTS, SHIFT, clients, hand_template, make_form
from tests.extract.fakes import CorpusModel


@pytest.fixture(autouse=True)
def no_model_loads(monkeypatch, extract_detectors_off):
    """No detector weights load: the documents' form maps and a setup's samples are read with no detectors."""
    real = documents_mod.build_form_map
    policy = DetectPolicy(digital=(), scanned=())
    monkeypatch.setattr(documents_mod, "build_form_map", lambda pdf, detectors=None, **kw: real(pdf, detectors={}, policy=policy, ocr="never"))


def wait_ended(jobs: Jobs, job_id: str, timeout: float = 10.0) -> dict:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        snap = jobs.get(job_id).snapshot()
        if snap["status"] != "running":
            return snap
        time.sleep(0.01)
    raise AssertionError("the job never ended")


# ------------------------------------------------------------------ task jobs


def test_a_task_job_starts_with_its_stages_pending_and_ends_done_with_the_stages_and_result_the_runner_set():
    jobs = Jobs()
    release = threading.Event()

    def runner(job):
        assert release.wait(10)
        job.stage("a", "done", "all of it")
        job.set_result(count=3)

    job = jobs.start_task("extract-setup", runner, stages=[("a", "Read the samples"), ("b", "Find the boxes")])
    assert job.kind == "extract-setup"
    assert job.snapshot()["stages"] == [
        {"key": "a", "label": "Read the samples", "status": "pending", "detail": ""},
        {"key": "b", "label": "Find the boxes", "status": "pending", "detail": ""},
    ]
    release.set()
    snap = wait_ended(jobs, job.id)
    assert snap["status"] == "done" and snap["error"] is None and snap["ended"] is not None
    assert snap["stages"][0] == {"key": "a", "label": "Read the samples", "status": "done", "detail": "all of it"}
    assert snap["stages"][1]["status"] == "pending"  # a done job leaves a stage the runner never reached as it was
    assert snap["result"] == {"count": 3}


def test_a_task_job_with_no_stages_has_none():
    jobs = Jobs()
    snap = wait_ended(jobs, jobs.start_task("extract-run", lambda job: None).id)
    assert snap["status"] == "done" and snap["stages"] == []


def test_a_task_job_that_raises_session_cancelled_ends_cancelled_and_skips_its_open_stages():
    jobs = Jobs()

    def runner(job):
        job.stage("a", "running")
        raise SessionCancelled("cancelled")

    snap = wait_ended(jobs, jobs.start_task("extract-run", runner, stages=[("a", "A"), ("b", "B")]).id)
    assert snap["status"] == "cancelled" and snap["error"] is None
    assert [st["status"] for st in snap["stages"]] == ["skipped", "skipped"]


def test_a_task_job_that_raises_a_formcast_error_ends_error_with_its_message():
    jobs = Jobs()

    def runner(job):
        job.stage("a", "running")
        raise FormcastError("no")

    snap = wait_ended(jobs, jobs.start_task("extract-line", runner, stages=[("a", "A")]).id)
    assert snap["status"] == "error" and snap["error"] == "no"
    assert snap["stages"][0]["status"] == "error"


def test_a_task_jobs_error_message_has_the_secrets_removed():
    jobs = Jobs(secrets=["sk-secret-value"])

    def runner(job):
        raise FormcastError("the key sk-secret-value was refused")

    snap = wait_ended(jobs, jobs.start_task("extract-line", runner).id)
    assert snap["error"] == "the key *** was refused"


def test_a_task_job_that_fails_unexpectedly_says_so_generically_and_logs_the_traceback_only(caplog):
    jobs = Jobs(secrets=["sk-secret-value"])

    def runner(job):
        raise RuntimeError("disk melted sk-secret-value")

    with caplog.at_level("ERROR", logger="demo.server"):
        snap = wait_ended(jobs, jobs.start_task("extract-run", runner).id)
    assert snap["status"] == "error"
    assert snap["error"] == "The job stopped because of an unexpected error. The server log has the details."
    assert "melted" not in str(snap) and "sk-secret-value" not in str(snap)
    assert "disk melted" in caplog.text and "sk-secret-value" not in caplog.text  # the log has it, redacted


def test_a_runner_sees_a_cancel_request_through_its_job():
    jobs = Jobs()
    started = threading.Event()

    def runner(job):
        started.set()
        while not job.cancel_requested:
            time.sleep(0.005)
        raise SessionCancelled("cancelled")

    job = jobs.start_task("extract-run", runner)
    assert started.wait(5)
    jobs.cancel(job.id)
    assert wait_ended(jobs, job.id)["status"] == "cancelled"


def test_task_jobs_take_the_same_slots_as_the_other_jobs():
    jobs = Jobs(max_running=2)
    release = threading.Event()
    first = jobs.start_task("extract-run", lambda job: release.wait(10))
    second = jobs.start_task("extract-run", lambda job: release.wait(10))
    with pytest.raises(JobsBusy):
        jobs.start_task("extract-run", lambda job: None)
    release.set()
    wait_ended(jobs, first.id)
    wait_ended(jobs, second.id)
    assert wait_ended(jobs, jobs.start_task("extract-run", lambda job: None).id)["status"] == "done"


def test_a_task_job_that_cannot_start_its_thread_frees_its_slot(monkeypatch):
    jobs = Jobs(max_running=1)

    def refuse(self):
        raise RuntimeError("can't start new thread")

    with monkeypatch.context() as patch:
        patch.setattr(threading.Thread, "start", refuse)
        with pytest.raises(RuntimeError):
            jobs.start_task("extract-run", lambda job: None)
    assert wait_ended(jobs, jobs.start_task("extract-run", lambda job: None).id)["status"] == "done"


# ------------------------------------------------------------------ settings


def test_the_extraction_settings_default_to_a_folder_scan_and_no_workers_of_their_own():
    settings = Settings.from_env({})
    assert settings.folder_scan is True and settings.extract_workers == 0


def test_the_extraction_settings_are_read_from_the_environment():
    settings = Settings.from_env({"FORMCAST_DEMO_FOLDER_SCAN": "off", "FORMCAST_DEMO_EXTRACT_WORKERS": "2"})
    assert settings.folder_scan is False and settings.extract_workers == 2
    assert Settings.from_env({"FORMCAST_DEMO_FOLDER_SCAN": "Yes", "FORMCAST_DEMO_EXTRACT_WORKERS": "0"}).folder_scan is True


@pytest.mark.parametrize(
    "name, value",
    [
        ("FORMCAST_DEMO_EXTRACT_WORKERS", "-1"),
        ("FORMCAST_DEMO_EXTRACT_WORKERS", "two"),
        ("FORMCAST_DEMO_EXTRACT_WORKERS", "1.5"),
        ("FORMCAST_DEMO_FOLDER_SCAN", "maybe"),
    ],
)
def test_a_bad_extraction_setting_is_refused_naming_the_variable_but_not_echoing_its_value(name, value):
    with pytest.raises(ValueError) as caught:
        Settings.from_env({name: value})
    message = str(caught.value)
    assert name in message and value not in message.replace(name, "")


# ------------------------------------------------------------------ the capability and the job kinds


def test_capabilities_say_extract_is_available(demo_server):
    status, caps = get_json(demo_server, "/api/capabilities")
    assert status == 200 and caps["extract"] is True


def test_an_unknown_job_kind_is_refused_naming_every_kind_even_with_no_model(demo_server):
    status, body = get_json(demo_server, "/api/jobs", "POST", {"kind": "nope"})
    assert status == 400 and body["error"] == "bad_kind"
    for kind in ("build", "adapt", "revise", "oneshot", "extract-setup", "extract-line", "extract-run"):
        assert kind in body["detail"]


def test_the_api_holds_a_place_for_each_extraction_manager():
    api = Api(Settings.from_env({}))
    for name in ("batches", "setups", "extractions", "runs"):
        assert hasattr(api, name)


# ------------------------------------------------------------------ several files per upload


def multipart_files(parts: list[tuple[str, str, bytes]]) -> tuple[bytes, dict]:
    boundary = "----formcastmulti"
    body = b""
    for field, filename, data in parts:
        head = f'--{boundary}\r\nContent-Disposition: form-data; name="{field}"; filename="{filename}"\r\n'
        body += head.encode() + b"Content-Type: application/octet-stream\r\n\r\n" + data + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    return body, {"Content-Type": f"multipart/form-data; boundary={boundary}"}


def request(body: bytes, headers: dict) -> Request:
    return Request("POST", "/x", {}, body, {k.lower(): v for k, v in headers.items()}, {})


def test_uploads_returns_every_part_named_file_in_order_and_skips_other_parts():
    body, headers = multipart_files([("file", "a.pdf", b"AAA"), ("note", "n.txt", b"ignored"), ("file", "dir/b.zip", b"BBB")])
    assert _uploads(request(body, headers)) == [("a.pdf", b"AAA"), ("b.zip", b"BBB")]


def test_uploads_with_no_part_named_file_is_a_bad_upload():
    body, headers = multipart_files([("other", "a.pdf", b"AAA")])
    with pytest.raises(ApiError) as caught:
        _uploads(request(body, headers))
    assert caught.value.status == 400 and caught.value.code == "bad_upload"
    with pytest.raises(ApiError) as caught:
        _uploads(request(b"not multipart", {"Content-Type": "text/plain"}))
    assert caught.value.code == "bad_upload"


def test_a_single_file_upload_still_opens_a_document(demo_server):
    body, headers = multipart_files([("file", "one.pdf", small_pdf())])
    status, _, _ = http_request(demo_server, "POST", "/api/documents", body, headers)
    assert status == 200


# ------------------------------------------------------------------ page_png


def test_page_png_renders_a_page_of_the_bytes_with_the_dpi_clamped():
    pdf = small_pdf(pages=2)
    png = page_png(pdf, 1, 72)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    assert page_png(pdf, 1, 144) != png
    assert page_png(pdf, 0, 1) == page_png(pdf, 0, documents_mod.MIN_DPI)
    assert page_png(pdf, 0, 9999) == page_png(pdf, 0, documents_mod.MAX_DPI)


@pytest.mark.parametrize("n", [-1, 2, 99])
def test_page_png_of_a_page_out_of_range_is_an_index_error(n):
    with pytest.raises(IndexError):
        page_png(small_pdf(pages=2), n, 72)


def test_page_png_holds_the_mupdf_lock_while_it_opens_the_pdf(monkeypatch):
    pdf = small_pdf()
    owned = []
    real_open = pymupdf.open

    def watching_open(*args, **kwargs):
        owned.append(PYMUPDF_LOCK._is_owned())
        return real_open(*args, **kwargs)

    monkeypatch.setattr(documents_mod.pymupdf, "open", watching_open)
    page_png(pdf, 0, 72)
    assert owned == [True]


def test_a_documents_page_png_is_the_module_functions_png():
    pdf = small_pdf(pages=2)
    doc = documents_mod.Documents().add(pdf, "x.pdf")
    assert doc.page_png(1, 90) == page_png(pdf, 1, 90)
    with pytest.raises(IndexError):
        doc.page_png(2, 90)


# ------------------------------------------------------------------ batches


@pytest.fixture(scope="module")
def corpus_pdfs() -> list[bytes]:
    """Five filled client-information forms, three fillable and two flattened: one layout."""
    kinds = ["fillable", "flattened", "fillable", "flattened", "fillable"]
    return [make_form(client, kind=kind) for client, kind in zip(clients(5), kinds)]


def zip_of(entries: list[tuple[str, bytes]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            z.writestr(name, data)
    return buf.getvalue()


def new_batch(base: str, **body) -> dict:
    status, batch = get_json(base, "/api/batches", "POST", body)
    assert status == 200, batch
    return batch


def add_files(base: str, batch_id: str, parts: list[tuple[str, bytes]]) -> tuple[int, dict]:
    body, headers = multipart_files([("file", name, data) for name, data in parts])
    status, _, out = http_request(base, "POST", f"/api/batches/{batch_id}/files", body, headers)
    return status, json.loads(out)


def scanned(base: str, batch_id: str) -> dict:
    """The batch once its scan has ended."""

    def ended():
        batch = get_json(base, f"/api/batches/{batch_id}")[1]
        return batch["scan"]["phase"] != "scanning" and batch

    return wait_for(ended, what="the scan to end")


def scan_ended(batch: extraction.Batch) -> dict:
    return wait_for(lambda: (s := batch.summary())["scan"]["phase"] != "scanning" and s, what="the scan to end")


@pytest.fixture
def five(demo_server, corpus_pdfs) -> tuple[str, dict]:
    """(base URL, the batch) of the five corpus PDFs, three uploaded as PDFs and two in a zip, scanned."""
    batch = new_batch(demo_server, name="Client forms")
    status, after = add_files(demo_server, batch["id"], [(f"c{i}.pdf", corpus_pdfs[i]) for i in range(3)])
    assert status == 200 and after["count"] == 3
    status, after = add_files(demo_server, batch["id"], [("more.zip", zip_of([("c3.pdf", corpus_pdfs[3]), ("c4.pdf", corpus_pdfs[4])]))])
    assert status == 200 and after["count"] == 5
    return demo_server, scanned(demo_server, batch["id"])


def batch_folder(start_demo_server, batch_id: str) -> Path:
    return start_demo_server.servers[-1].api.settings.data_dir / "local" / "batches" / batch_id


def test_a_new_batch_is_empty_and_its_scan_is_done_at_once(demo_server):
    batch = new_batch(demo_server, name="Client forms")
    assert len(batch["id"]) == 32
    assert batch == {
        "id": batch["id"], "name": "Client forms", "count": 0, "bytes": 0, "skipped": 0, "folder": False,
        "scan": {"phase": "done", "done": 0, "total": 0, "error": None, "groups": []},
    }  # fmt: skip
    assert get_json(demo_server, f"/api/batches/{batch['id']}") == (200, batch)
    assert new_batch(demo_server)["name"] == "Uploaded files"


def test_pdfs_uploaded_in_one_request_and_a_zip_in_another_are_one_batch_scanned_into_one_layout(five, corpus_pdfs):
    base, batch = five
    assert batch["count"] == 5 and batch["skipped"] == 0 and batch["folder"] is False
    assert batch["bytes"] == sum(len(pdf) for pdf in corpus_pdfs)
    scan = batch["scan"]
    assert scan["phase"] == "done" and scan["done"] == scan["total"] == 5 and scan["error"] is None
    assert scan["groups"] == [{"group": 1, "size": 5, "fillable": 3, "flattened": 2, "closeness": None}]


def test_a_zips_pdfs_are_stored_by_number_under_the_batch_whatever_their_names_and_its_other_entries_are_skipped(start_demo_server, corpus_pdfs, tmp_path):
    base = start_demo_server()
    batch = new_batch(base)
    before = set(tmp_path.rglob("*"))
    archive = zip_of([("../evil.pdf", corpus_pdfs[0]), ("/abs.pdf", corpus_pdfs[1]), ("sub/ok.pdf", corpus_pdfs[2]), ("notes.txt", b"hello")])
    status, after = add_files(base, batch["id"], [("bad names.zip", archive)])
    assert status == 200 and after["count"] == 3 and after["skipped"] == 1
    folder = batch_folder(start_demo_server, batch["id"])
    assert sorted(p.name for p in (folder / "files").iterdir()) == ["0.pdf", "1.pdf", "2.pdf"]
    assert (folder / "files" / "2.pdf").read_bytes() == corpus_pdfs[2]
    written = set(tmp_path.rglob("*")) - before
    assert written and all(p == folder or folder in p.parents or p in folder.parents for p in written)
    assert not any(p.name in ("evil.pdf", "abs.pdf", "ok.pdf") for p in tmp_path.parent.rglob("*"))
    files = get_json(base, f"/api/batches/{batch['id']}/files")[1]["files"]
    assert [f["name"] for f in files] == ["evil.pdf", "abs.pdf", "ok.pdf"]


def test_a_zip_declaring_more_than_the_limit_is_refused_before_anything_of_the_request_is_written(start_demo_server, corpus_pdfs, monkeypatch):
    base = start_demo_server()
    batch = new_batch(base)
    assert add_files(base, batch["id"], [("a.pdf", corpus_pdfs[0])])[0] == 200
    monkeypatch.setattr(extraction, "MAX_ZIP_BYTES", len(corpus_pdfs[1]) + len(corpus_pdfs[2]) - 1)
    archive = zip_of([("b.pdf", corpus_pdfs[1]), ("c.pdf", corpus_pdfs[2])])
    status, body = add_files(base, batch["id"], [("d.pdf", corpus_pdfs[3]), ("big.zip", archive)])
    assert status == 400 and body["error"] == "bad_upload" and "big.zip" in body["detail"]
    after = scanned(base, batch["id"])
    assert after["count"] == 1 and after["skipped"] == 0
    assert [p.name for p in (batch_folder(start_demo_server, batch["id"]) / "files").iterdir()] == ["0.pdf"]


def test_an_upload_that_would_take_the_batch_past_the_most_files_is_refused(demo_server, corpus_pdfs, monkeypatch):
    monkeypatch.setattr(extraction, "MAX_BATCH_FILES", 2)
    batch = new_batch(demo_server)
    status, body = add_files(demo_server, batch["id"], [(f"{i}.pdf", corpus_pdfs[i]) for i in range(3)])
    assert status == 400 and body["error"] == "bad_upload"
    assert get_json(demo_server, f"/api/batches/{batch['id']}")[1]["count"] == 0


def test_a_zip_that_cannot_be_opened_is_refused(demo_server):
    batch = new_batch(demo_server)
    status, body = add_files(demo_server, batch["id"], [("broken.zip", b"PK not a zip")])
    assert status == 400 and body["error"] == "bad_upload" and "broken.zip" in body["detail"]


def test_files_that_are_neither_pdfs_nor_zips_are_skipped_and_counted(demo_server, corpus_pdfs):
    batch = new_batch(demo_server)
    status, after = add_files(demo_server, batch["id"], [("A.PDF", corpus_pdfs[0]), ("notes.txt", b"x"), ("pic.png", b"y")])
    assert status == 200 and after["count"] == 1 and after["skipped"] == 2


def test_a_server_folder_is_scanned_in_place_without_copying_its_files(start_demo_server, corpus_pdfs, tmp_path):
    folder = tmp_path / "incoming"
    (folder / "sub").mkdir(parents=True)
    for i, pdf in enumerate(corpus_pdfs[:3]):
        (folder / f"{i}.pdf").write_bytes(pdf)
    (folder / "sub" / "3.PDF").write_bytes(corpus_pdfs[3])
    (folder / "readme.txt").write_text("not a pdf")
    base = start_demo_server()
    batch = new_batch(base, folder=str(folder))
    assert batch["count"] == 4 and batch["folder"] is True and batch["name"] == "incoming"
    assert str(tmp_path) not in json.dumps(batch)
    done = scanned(base, batch["id"])
    assert done["scan"]["phase"] == "done" and done["scan"]["groups"][0]["size"] == 4
    assert not batch_folder(start_demo_server, batch["id"]).exists()
    files = get_json(base, f"/api/batches/{batch['id']}/files")[1]["files"]
    assert [f["name"] for f in files] == ["0.pdf", "1.pdf", "2.pdf", "3.PDF"]
    assert http_request(base, "GET", f"/api/batches/{batch['id']}/files/3/pages/0.png")[0] == 200
    status, body = add_files(base, batch["id"], [("x.pdf", corpus_pdfs[4])])
    assert status == 400 and body["error"] == "bad_upload"  # a folder's batch is the folder


def test_a_folder_batch_is_refused_when_the_server_does_not_allow_it_or_the_folder_is_no_use(start_demo_server, tmp_path):
    (tmp_path / "empty").mkdir()
    (tmp_path / "a.pdf").write_bytes(small_pdf())
    status, body = get_json(start_demo_server(folder_scan=False), "/api/batches", "POST", {"folder": str(tmp_path)})
    assert status == 403 and body["error"] == "not_allowed"
    base = start_demo_server()
    for folder in (str(tmp_path / "missing"), str(tmp_path / "empty"), str(tmp_path / "a.pdf"), "", 7):
        status, body = get_json(base, "/api/batches", "POST", {"folder": folder})
        assert status == 400 and body["error"] == "bad_folder", folder
        assert str(tmp_path) not in body["detail"]


def test_a_folders_walk_stops_past_the_most_files_and_a_file_that_cannot_be_read_or_a_bad_home_is_never_a_500(start_demo_server, tmp_path, monkeypatch):
    folder = tmp_path / "many"
    folder.mkdir()
    for name in ("a.pdf", "b.pdf", "c.pdf", "d.pdf", "e.pdf"):
        (folder / name).write_bytes(small_pdf())
    base = start_demo_server()
    monkeypatch.setattr(extraction, "MAX_BATCH_FILES", 3)
    walked = []
    rglob = Path.rglob

    def counted(self, pattern):
        for p in rglob(self, pattern):
            walked.append(p)
            yield p

    monkeypatch.setattr(Path, "rglob", counted)
    status, body = get_json(base, "/api/batches", "POST", {"folder": str(folder)})
    assert (status, body["error"]) == (400, "bad_folder") and "more than 3 PDF files" in body["detail"]
    assert len(walked) == 4  # the walk stopped at the fourth PDF
    monkeypatch.setattr(extraction, "MAX_BATCH_FILES", 200_000)

    stat = Path.stat

    def gone_after_the_walk(self, *args, **kw):  # b.pdf disappears between the walk finding it and its stat
        if self.name == "b.pdf":
            raise FileNotFoundError(self)
        return stat(self, *args, **kw)

    monkeypatch.setattr(Path, "stat", gone_after_the_walk)
    status, batch = get_json(base, "/api/batches", "POST", {"folder": str(folder)})
    assert status == 200, batch
    names = [f["name"] for f in get_json(base, f"/api/batches/{batch['id']}/files")[1]["files"]]
    assert "b.pdf" not in names and len(names) == 4
    monkeypatch.setattr(Path, "stat", stat)

    for folder_text in ("~no_such_user_for_formcast/forms", "~no_such_user_for_formcast"):
        status, body = get_json(base, "/api/batches", "POST", {"folder": folder_text})
        assert (status, body["error"]) == (400, "bad_folder"), body
        status, body = get_json(base, "/api/runs", "POST", {"folder": folder_text})
        assert (status, body["error"]) == (400, "bad_run"), body


def test_a_batch_name_must_be_text(demo_server):
    status, body = get_json(demo_server, "/api/batches", "POST", {"name": 5})
    assert status == 400 and body["error"] == "bad_request"


def test_the_files_list_is_paged_and_a_page_is_at_most_the_limit(five):
    base, batch = five
    status, page = get_json(base, f"/api/batches/{batch['id']}/files?offset=1&limit=2")
    assert status == 200 and page["total"] == 5
    assert page["files"] == [
        {"index": 1, "name": "c1.pdf", "size": page["files"][0]["size"], "group": 1},
        {"index": 2, "name": "c2.pdf", "size": page["files"][1]["size"], "group": 1},
    ]
    assert get_json(base, f"/api/batches/{batch['id']}/files?offset=9")[1] == {"total": 5, "files": []}
    for query in ("limit=500", "limit=0", "offset=-1", "limit=two"):
        status, body = get_json(base, f"/api/batches/{batch['id']}/files?{query}")
        assert status == 400 and body["error"] == "bad_request", query


def test_a_sample_is_picked_from_the_main_layout_or_as_the_person_picked_it(five):
    base, batch = five
    url = f"/api/batches/{batch['id']}/sample"
    status, out = get_json(base, url, "POST", {"n": 3, "how": "spread"})
    assert status == 200 and len({s["index"] for s in out["samples"]}) == 3
    assert all(s["group"] == 1 and s["name"] == f"c{s['index']}.pdf" for s in out["samples"])
    assert [s["index"] for s in get_json(base, url, "POST", {"n": 3, "how": "first"})[1]["samples"]] == [0, 1, 2]
    assert [s["index"] for s in get_json(base, url, "POST", {"how": "pick", "picked": [4, 0]})[1]["samples"]] == [4, 0]
    bad = [
        {"n": 2, "how": "spread"}, {"n": 21, "how": "first"}, {"n": "3", "how": "first"}, {"n": 3, "how": "best"},
        {"how": "pick", "picked": [0, 4, 9]}, {"how": "pick", "picked": []}, {"how": "pick", "picked": [1, 1]},
    ]  # fmt: skip
    for body in bad:
        status, err = get_json(base, url, "POST", body)
        assert status == 400 and err["error"] == "bad_sample", body


def test_a_sample_waits_for_the_scan(tmp_path, corpus_pdfs, monkeypatch):
    from formcast.extract import scan as scan_mod

    real_read, release = scan_mod._read, threading.Event()
    monkeypatch.setattr(scan_mod, "_read", lambda path, first: release.wait(10) and real_read(path, first))
    batches = extraction.Batches(tmp_path / "batches")
    batch = batches.create("b")
    batches.add(batch, [("a.pdf", corpus_pdfs[0])])
    with pytest.raises(extraction.ScanNotDone):
        batches.pick(batch, 3, "spread", None)
    assert batches.pick(batch, 3, "pick", [0]) == [0]  # the person's own pick needs no scan
    release.set()
    scan_ended(batch)
    assert batches.pick(batch, 3, "spread", None) == [0]


def test_a_page_of_a_batch_file_is_a_png_and_a_page_or_batch_that_is_not_there_is_404(five):
    base, batch = five
    status, headers, png = http_request(base, "GET", f"/api/batches/{batch['id']}/files/0/pages/0.png?dpi=72")
    assert status == 200 and headers["content-type"] == "image/png"
    assert png_size(png) == (612, 792)
    status, body = get_json(base, f"/api/batches/{batch['id']}/files/0/pages/99.png")
    assert status == 404 and body["error"] == "not_found"
    status, body = get_json(base, f"/api/batches/{batch['id']}/files/5/pages/0.png")
    assert status == 404 and body["error"] == "not_found"
    status, body = get_json(base, f"/api/batches/{batch['id']}/files/0/pages/0.png?dpi=big")
    assert status == 400 and body["error"] == "bad_dpi"
    for path in (f"/api/batches/{'0' * 32}", f"/api/batches/{'0' * 32}/files", f"/api/batches/{'0' * 32}/files/0/pages/0.png"):
        status, body = get_json(base, path)
        assert status == 404 and body == {"error": "not_found", "detail": f"no such batch: {'0' * 32}"}


def test_an_unreadable_file_is_in_no_group_and_its_page_cannot_be_shown(demo_server, corpus_pdfs):
    batch = new_batch(demo_server)
    add_files(demo_server, batch["id"], [("good.pdf", corpus_pdfs[0]), ("broken.pdf", b"%PDF-1.7 not really")])
    done = scanned(demo_server, batch["id"])
    assert done["scan"]["phase"] == "done" and done["scan"]["groups"][0]["size"] == 1
    files = get_json(demo_server, f"/api/batches/{batch['id']}/files")[1]["files"]
    assert [f["group"] for f in files] == [1, None]
    status, body = get_json(demo_server, f"/api/batches/{batch['id']}/files/1/pages/0.png")
    assert status == 422 and body["error"] == "bad_pdf"


def test_an_upload_while_a_scan_runs_starts_it_again_and_the_old_scans_results_are_dropped(tmp_path, corpus_pdfs, monkeypatch):
    from formcast.extract import scan as scan_mod

    real_read = scan_mod._read
    reading, release, reads = threading.Event(), threading.Event(), []

    def slow_first_read(path, first_page_only):
        reads.append(Path(path).name)
        if len(reads) == 1:  # the first scan's first document takes a while
            reading.set()
            assert release.wait(10)
        return real_read(path, first_page_only)

    monkeypatch.setattr(scan_mod, "_read", slow_first_read)
    batches = extraction.Batches(tmp_path / "batches")
    batch = batches.create("b")
    batches.add(batch, [("a.pdf", corpus_pdfs[0]), ("b.pdf", corpus_pdfs[1])])
    assert reading.wait(10)
    batches.add(batch, [("c.pdf", corpus_pdfs[2])])  # starts the scan again, over all three
    assert batch.summary()["scan"] == {"phase": "scanning", "done": 0, "total": 3, "error": None, "groups": []}
    release.set()  # the first scan reads its document, sees it was replaced and stops; the second runs
    scan = scan_ended(batch)["scan"]
    assert scan["phase"] == "done" and scan["done"] == scan["total"] == 3 and scan["groups"][0]["size"] == 3
    assert reads == ["0.pdf", "0.pdf", "1.pdf", "2.pdf"]
    assert [batch.group_of(i) for i in range(3)] == [1, 1, 1]


def test_the_scan_reads_each_document_under_the_mupdf_lock_and_lets_go_at_the_end(tmp_path, corpus_pdfs, monkeypatch):
    from formcast.extract import scan as scan_mod

    real_read, owned = scan_mod._read, []
    monkeypatch.setattr(scan_mod, "_read", lambda path, first: owned.append(PYMUPDF_LOCK._is_owned()) or real_read(path, first))
    batches = extraction.Batches(tmp_path / "batches")
    batch = batches.create("b")
    batches.add(batch, [(f"{i}.pdf", pdf) for i, pdf in enumerate(corpus_pdfs[:3])])
    assert scan_ended(batch)["scan"]["phase"] == "done"
    assert owned == [True, True, True]
    assert PYMUPDF_LOCK.acquire(timeout=5)
    PYMUPDF_LOCK.release()


def test_a_large_batch_is_scanned_on_first_pages_only(tmp_path, corpus_pdfs, monkeypatch):
    real_scan, seen = extraction.scan, []
    monkeypatch.setattr(extraction, "scan", lambda paths, **kw: seen.append(kw["first_page_only"]) or real_scan(paths, **kw))
    monkeypatch.setattr(extraction, "FIRST_PAGE_ONLY_FROM", 2)
    batches = extraction.Batches(tmp_path / "batches")
    batch = batches.create("b")
    batches.add(batch, [("a.pdf", corpus_pdfs[0])])
    scan_ended(batch)
    batches.add(batch, [("b.pdf", corpus_pdfs[1])])
    wait_for(lambda: len(seen) == 2, what="the second scan")
    assert scan_ended(batch)["scan"]["groups"][0]["size"] == 2
    assert seen == [False, True]


def test_a_scan_that_fails_unexpectedly_says_so_generically(tmp_path, corpus_pdfs, monkeypatch, caplog):
    def broken(paths, **kw):
        raise RuntimeError("/secret/path melted")

    monkeypatch.setattr(extraction, "scan", broken)
    batches = extraction.Batches(tmp_path / "batches")
    batch = batches.create("b")
    with caplog.at_level("ERROR", logger="demo.server"):
        batches.add(batch, [("a.pdf", corpus_pdfs[0])])
        scan = scan_ended(batch)["scan"]
    assert scan["phase"] == "error" and scan["error"] and "secret" not in scan["error"]
    assert "melted" in caplog.text
    with pytest.raises(KeyError):
        batches.get("0" * 32)
    assert batches.get(batch.id) is batch


# ------------------------------------------------------------------ setups


#: Five corpus copies, as plan 2's setup tests use: three fillable, two flattened, the last pushed down by SHIFT.
SAMPLE_KINDS = [("fillable", 0.0), ("fillable", 0.0), ("fillable", 0.0), ("flattened", 0.0), ("flattened", SHIFT)]
SHIFTED = 4  # the sample pushed down the page


def sample_set(people) -> list[tuple[str, bytes]]:
    return [
        (f"s{n}.pdf", make_form(client, kind=kind, shift=shift))
        for n, (client, (kind, shift)) in enumerate(zip(people, SAMPLE_KINDS), start=1)
    ]


class GatedModel(CorpusModel):
    """``CorpusModel`` that stops at the first call of kind ``gate`` (``"derive"``, ``"fields"``) until released.

    ``us_files`` are samples whose country the model reads as ``"US"`` whatever the form prints there.
    """

    KINDS = {DERIVE_SYSTEM.split("Reply with one JSON object")[0]: "derive", FIELDS_SYSTEM: "fields"}

    def kind(self, system) -> str | None:
        return next((kind for prompt, kind in self.KINDS.items() if system.startswith(prompt)), None)  # derive may add rules before its reply sentence

    def __init__(self, clients_by_file, *, gate: str | None = None, us_files=()) -> None:
        super().__init__(clients_by_file)
        self.gate = gate
        self.us_files = set(us_files)
        self.reached = threading.Event()
        self.released = threading.Event()

    def complete(self, messages, tools, **kwargs):
        if self.gate is not None and self.kind(messages[0].content) == self.gate:
            self.reached.set()
            assert self.released.wait(10)
        return super().complete(messages, tools, **kwargs)

    def _values(self, text: str) -> dict:
        answer = super()._values(text)
        if any(f"({name}," in text for name in self.us_files) and "investor.country" in answer["values"]:
            answer["values"]["investor.country"] = "US"
        return answer


@pytest.fixture(scope="module")
def setup_people():
    return clients(5, seed=31)


@pytest.fixture(scope="module")
def setup_pdfs(setup_people) -> list[tuple[str, bytes]]:
    return sample_set(setup_people)


@dataclass
class Bench:
    """A server with ``model`` as its AI and a batch of the sample files."""

    base: str
    api: Api
    model: GatedModel
    batch: str
    pdfs: list[bytes]


def bench_of(start_demo_server, files: list[tuple[str, bytes]], model: GatedModel) -> Bench:
    base = start_demo_server(provider_factory=model)
    batch = new_batch(base, name="Client forms")
    status, _ = add_files(base, batch["id"], files)
    assert status == 200
    return Bench(base, start_demo_server.servers[-1].api, model, batch["id"], [data for _, data in files])


def model_for(files: list[tuple[str, bytes]], people, **kw) -> GatedModel:
    return GatedModel({name: client for (name, _), client in zip(files, people)}, **kw)


@pytest.fixture
def bench(start_demo_server, setup_pdfs, setup_people) -> Bench:
    return bench_of(start_demo_server, setup_pdfs, model_for(setup_pdfs, setup_people))


def start_job(base: str, body: dict) -> tuple[int, dict]:
    return get_json(base, "/api/jobs", "POST", body)


def job_of(base: str, job_id: str) -> dict:
    return get_json(base, f"/api/jobs/{job_id}")[1]


def job_ended(base: str, job_id: str) -> dict:
    return wait_for(lambda: (j := job_of(base, job_id))["status"] != "running" and j, timeout=60, what="the job to end")


def setup_job(b: Bench, **spec) -> dict:
    body = {"kind": "extract-setup", "batch": b.batch, "samples": [0, 1, 2, 3, 4], "spec": {"name": "Client forms", "names_from_titles": False, **spec}}
    status, out = start_job(b.base, body)
    assert status == 200, out
    return out


def built(b: Bench, **spec) -> dict:
    """The setup an ``extract-setup`` job built, once it is done."""
    job = job_ended(b.base, setup_job(b, **spec)["id"])
    assert job["status"] == "done", job
    return job["result"]["setup"]


def assert_proven(proof: dict) -> None:
    """Every field some sample has is ready, and no sample disagrees (a field no sample fills cannot be proven)."""
    for field in proof["fields"]:
        assert not any(s["outcome"] == "disagree" for s in field["samples"]), field
        if not field["absent"] and any(s["outcome"] != "missing" for s in field["samples"]):
            assert field["ready"], field


def field_of(setup: dict, path: str) -> dict:
    return next(f for f in setup["proof"]["fields"] if f["path"] == path)


def edit(b: Bench, setup_id: str, change: dict) -> tuple[int, dict]:
    return get_json(b.base, f"/api/setups/{setup_id}/edit", "POST", {"edit": change})


def test_a_setup_job_builds_a_template_from_the_batchs_sample_in_the_five_stages_and_keeps_the_setup(bench):
    job = job_ended(bench.base, setup_job(bench, detect="all", sample_how="the first five")["id"])
    assert job["status"] == "done", job
    assert [s["key"] for s in job["stages"]] == [key for key, _ in STAGES]
    assert all(s["status"] == "done" for s in job["stages"]), job["stages"]
    assert job["usage"]["input_tokens"] == 10 * sum(bench.model.calls.values())
    setup = job["result"]["setup"]
    assert_proven(setup["proof"])
    assert setup["proof"]["ready"] > 0 and setup["batch"] == bench.batch
    template = setup["template"]
    assert template["name"] == "Client forms" and template["template_id"] == "client_forms"
    assert template["setup"]["start"] == "ai" and template["setup"]["detect"] == "all"
    assert template["setup"]["sample_how"] == "the first five"
    assert [s["name"] for s in setup["samples"]] == [f"s{n}.pdf" for n in range(1, 6)]
    assert [s["index"] for s in setup["samples"]] == [0, 1, 2, 3, 4]
    assert [s["kind"] for s in setup["samples"]] == [kind for kind, _ in SAMPLE_KINDS]
    assert setup["samples"][0]["pages"][0] == [612.0, 792.0]
    status, again = get_json(bench.base, f"/api/setups/{setup['id']}")
    assert status == 200 and again == setup


def test_the_fields_stage_is_labelled_by_where_the_fields_start(bench):
    job = job_ended(bench.base, setup_job(bench, start="lines", lines=["tax ID"])["id"])
    assert job["status"] == "done", job
    assert next(s for s in job["stages"] if s["key"] == "fields")["label"] == "Turn the lines into fields"
    assert job["result"]["setup"]["template"]["setup"]["lines"] == ["tax ID"]


def _without_reasons(node):
    """``node`` (a schema) with every ``x-required-reason`` taken out."""
    if isinstance(node, dict):
        return {k: _without_reasons(v) for k, v in node.items() if k != "x-required-reason"}
    if isinstance(node, list):
        return [_without_reasons(v) for v in node]
    return node


def test_a_setup_from_a_destination_marks_the_destinations_fields_the_form_lacks_absent(bench):
    status, listed = get_json(bench.base, "/api/templates")
    assert "fw9" in [t["name"] for t in listed["templates"]]
    schema = bench.api.store.get("fw9").schema_
    setup = built(bench, start="destination", destination="fw9")
    template = setup["template"]
    on_form = set(concrete_paths(hand_template().schema_))  # CorpusModel: the destination's fields the hand template has
    assert _without_reasons(template["schema"]) == schema and template["setup"]["destination"] == "fw9"
    assert template["schema"]["properties"]["legal_name"]["x-required-reason"] == "from the destination's schema"
    assert set(template["absent"]) == {p for p in concrete_paths(schema) if p not in on_form}
    assert "address.street" not in template["absent"] and "address.street" in template["readers"]
    assert_proven(setup["proof"])


def test_a_setup_job_is_refused_before_it_starts_for_a_bad_spec_batch_sample_or_destination(bench):
    def refused(**changes):
        body = {"kind": "extract-setup", "batch": bench.batch, "samples": [0, 1, 2], "spec": {"name": "x", "names_from_titles": False}, **changes}
        status, out = start_job(bench.base, body)
        return status, out["error"], out["detail"]

    assert refused(spec={"name": " "})[:2] == (400, "bad_spec")
    assert refused(spec={"name": "x", "start": "lines"})[:2] == (400, "bad_spec")
    assert refused(spec={"name": "x", "detect": "sometimes"})[:2] == (400, "bad_spec")
    assert refused(spec={"name": "x", "stray": 1})[:2] == (400, "bad_spec")
    assert refused(spec="x")[:2] == (400, "bad_spec")
    assert refused(batch="0" * 32)[:2] == (400, "bad_batch")
    assert refused(samples=[0, 9, 1])[:2] == (400, "bad_sample")
    assert refused(samples=[0, 1])[:2] == (400, "bad_sample")  # fewer than 3 of a batch of 5
    assert refused(samples=[0, 0, 1])[:2] == (400, "bad_sample")
    assert refused(samples="all")[:2] == (400, "bad_sample")
    status, code, detail = refused(spec={"name": "x", "start": "destination", "destination": "no such form"})
    assert (status, code) == (404, "not_found") and "no such template" in detail
    assert bench.model.calls == {}  # no model call for a job that never started


def test_a_batch_of_fewer_than_three_files_builds_from_all_of_them(start_demo_server, setup_pdfs, setup_people):
    b = bench_of(start_demo_server, setup_pdfs[:2], model_for(setup_pdfs, setup_people))
    status, out = start_job(b.base, {"kind": "extract-setup", "batch": b.batch, "samples": [0, 1], "spec": {"name": "x", "names_from_titles": False}})
    assert status == 200, out
    assert job_ended(b.base, out["id"])["status"] == "done"


def test_without_ai_a_setup_job_is_refused_and_an_unknown_setup_is_not_found(demo_server):
    batch = new_batch(demo_server)
    status, out = start_job(demo_server, {"kind": "extract-setup", "batch": batch["id"], "samples": [0], "spec": {"name": "x", "names_from_titles": False}})
    assert (status, out["error"]) == (400, "ai_not_configured")
    status, out = start_job(demo_server, {"kind": "extract-line", "setup": "0" * 32, "line": "wire details"})
    assert (status, out["error"]) == (400, "ai_not_configured")
    for path in ("", "/samples/0/marks", "/samples/0/pages/0.png"):
        status, out = get_json(demo_server, f"/api/setups/{'0' * 32}{path}")
        assert (status, out["error"]) == (404, "not_found") and out["detail"].startswith("no such setup")
    for path, body in (("/edit", {"edit": {"op": "mark_absent", "path": "x"}}), ("/learn", {"path": "x"})):
        status, out = get_json(demo_server, f"/api/setups/{'0' * 32}{path}", "POST", body)
        assert (status, out["error"]) == (404, "not_found")


def test_cancelling_a_setup_job_during_the_fields_stage_ends_it_cancelled_with_no_setup(start_demo_server, setup_pdfs, setup_people):
    model = model_for(setup_pdfs, setup_people, gate="derive")
    b = bench_of(start_demo_server, setup_pdfs, model)
    job_id = setup_job(b)["id"]
    assert model.reached.wait(30)
    running = job_of(b.base, job_id)
    assert next(s for s in running["stages"] if s["key"] == "fields")["status"] == "running"
    assert get_json(b.base, f"/api/jobs/{job_id}/cancel", "POST", {})[0] == 200
    model.released.set()
    job = job_ended(b.base, job_id)
    assert job["status"] == "cancelled" and "setup" not in job["result"]
    assert "readers" not in model.calls  # it stopped before the next stage
    assert not b.api.setups._setups


def test_a_setup_job_cancelled_while_it_reads_the_samples_ends_cancelled_and_leaves_mupdf_free(start_demo_server, setup_pdfs, setup_people, monkeypatch):
    reached, released = threading.Event(), threading.Event()
    real = samples_mod.build_form_map

    def gated(pdf, **kw):  # the first sample's form map waits (holding the lock, as the library does) until released
        if not reached.is_set():
            reached.set()
            assert released.wait(10)
        return real(pdf, **kw)

    monkeypatch.setattr(samples_mod, "build_form_map", gated)
    model = model_for(setup_pdfs, setup_people)
    b = bench_of(start_demo_server, setup_pdfs, model)
    job_id = setup_job(b)["id"]
    assert reached.wait(30)
    assert next(s for s in job_of(b.base, job_id)["stages"] if s["key"] == "read")["status"] == "running"
    assert get_json(b.base, f"/api/jobs/{job_id}/cancel", "POST", {})[0] == 200
    released.set()  # the form map's first page tells on_page, which stops the job
    job = job_ended(b.base, job_id)
    assert job["status"] == "cancelled" and model.calls == {}
    status, headers, png = http_request(b.base, "GET", f"/api/batches/{b.batch}/files/0/pages/0.png?dpi=50")
    assert status == 200 and headers["content-type"] == "image/png"
    assert PYMUPDF_LOCK.acquire(timeout=5)  # from this thread too: nothing kept the lock
    PYMUPDF_LOCK.release()


def test_an_edit_is_applied_and_proven_again_with_no_model_call(bench):
    setup = built(bench)
    calls = dict(bench.model.calls)
    status, after = edit(bench, setup["id"], {"op": "set_labels", "path": "investor.tax_id", "labels": []})
    assert status == 200, after
    kinds = [s["kind"] for s in after["template"]["readers"]["investor.tax_id"]["strategies"]]
    assert "anchor" in [s["kind"] for s in setup["template"]["readers"]["investor.tax_id"]["strategies"]]
    assert "anchor" not in kinds and kinds
    template = ExtractionTemplate.model_validate(after["template"])
    assert Proof.model_validate(after["proof"]) == prove(template, bench.pdfs)
    assert bench.model.calls == calls
    assert get_json(bench.base, f"/api/setups/{setup['id']}")[1] == after


def test_a_bad_edit_is_refused_with_the_librarys_message_and_changes_nothing(bench):
    setup = built(bench)
    status, out = edit(bench, setup["id"], {"op": "set_labels", "path": "investor.nope", "labels": []})
    assert (status, out["error"], out["detail"]) == (400, "bad_edit", "no field investor.nope")
    status, out = edit(bench, setup["id"], {"op": "fly", "path": "investor.tax_id"})
    assert (status, out["error"]) == (400, "bad_edit")
    status, out = get_json(bench.base, f"/api/setups/{setup['id']}/edit", "POST", {"change": {}})
    assert (status, out["error"]) == (400, "bad_edit")
    assert get_json(bench.base, f"/api/setups/{setup['id']}")[1] == setup


def test_a_box_drawn_on_a_shifted_samples_marks_is_stored_in_the_layouts_place(bench):
    setup = built(bench)
    status, out = get_json(bench.base, f"/api/setups/{setup['id']}/samples/{SHIFTED}/marks")
    assert status == 200
    mark = next(m for m in out["marks"] if (m["path"], m["kind"]) == ("investor.dob", "box"))
    x0, y0, x1, y1 = FIELD_RECTS["dob"]
    assert mark["rect"] == pytest.approx([x0, y0 + SHIFT, x1, y1 + SHIFT], abs=0.5)
    drawn = [mark["rect"][0], mark["rect"][1], mark["rect"][2] + 20, mark["rect"][3]]
    change = {"op": "set_box", "path": "investor.dob", "page": mark["page"], "rect": drawn, "sample": SHIFTED}
    status, after = edit(bench, setup["id"], change)
    assert status == 200, after
    box = next(s for s in after["template"]["readers"]["investor.dob"]["strategies"] if s["kind"] == "box")
    assert box["rect"] == pytest.approx([x0, y0, x1 + 20, y1], abs=0.5)
    assert field_of(after, "investor.dob")["ready"]


def test_a_samples_marks_have_a_kind_and_a_place_and_its_pages_are_images(bench):
    setup = built(bench)
    status, out = get_json(bench.base, f"/api/setups/{setup['id']}/samples/0/marks")
    assert status == 200 and out["marks"]
    for mark in out["marks"]:
        assert mark["kind"] in ("read", "box", "label", "option") and len(mark["rect"]) == 4
    status, headers, png = http_request(bench.base, "GET", f"/api/setups/{setup['id']}/samples/0/pages/0.png?dpi=50")
    assert status == 200 and headers["content-type"] == "image/png" and png_size(png)[0] == round(612 * 50 / 72)
    assert get_json(bench.base, f"/api/setups/{setup['id']}/samples/9/marks")[0] == 404
    assert http_request(bench.base, "GET", f"/api/setups/{setup['id']}/samples/0/pages/9.png")[0] == 404
    assert http_request(bench.base, "GET", f"/api/setups/{setup['id']}/samples/0/pages/0.png?dpi=x")[0] == 400


def test_a_line_job_adds_the_fields_of_one_more_line_to_the_setup(bench):
    setup = built(bench, start="lines", lines=["tax ID"])
    assert "wire.bank" not in setup["template"]["readers"]
    status, out = start_job(bench.base, {"kind": "extract-line", "setup": setup["id"], "line": "wire details"})
    assert status == 200, out
    job = job_ended(bench.base, out["id"])
    assert job["status"] == "done", job
    after = job["result"]["setup"]
    assert job["result"]["paths"] == ["wire.bank", "wire.account"]
    assert "wire.bank" in after["template"]["readers"] and field_of(after, "wire.bank")["ready"]
    assert after["template"]["setup"]["line_paths"]["wire details"] == ["wire.bank", "wire.account"]
    assert get_json(bench.base, f"/api/setups/{setup['id']}")[1] == after


def test_a_line_job_is_refused_for_no_line_or_an_unknown_setup(bench):
    setup = built(bench, start="lines", lines=["tax ID"])
    status, out = start_job(bench.base, {"kind": "extract-line", "setup": setup["id"], "line": "  "})
    assert (status, out["error"]) == (400, "bad_request")
    status, out = start_job(bench.base, {"kind": "extract-line", "setup": "0" * 32, "line": "wire details"})
    assert (status, out["error"]) == (404, "not_found") and out["detail"].startswith("no such setup")


def test_while_a_line_job_runs_its_setup_takes_no_edit_and_no_other_line_job(bench):
    setup = built(bench, start="lines", lines=["tax ID"])
    bench.model.gate = "fields"
    status, out = start_job(bench.base, {"kind": "extract-line", "setup": setup["id"], "line": "wire details"})
    assert status == 200, out
    assert bench.model.reached.wait(30)
    tax_id = {"op": "set_labels", "path": "investor.tax_id", "labels": []}
    status, busy = edit(bench, setup["id"], tax_id)
    assert (status, busy["error"]) == (409, "busy_setup")
    status, busy = start_job(bench.base, {"kind": "extract-line", "setup": setup["id"], "line": "date of birth"})
    assert (status, busy["error"]) == (409, "busy_setup")
    status, busy = get_json(bench.base, f"/api/setups/{setup['id']}/learn", "POST", {"path": "investor.country"})
    assert (status, busy["error"]) == (409, "busy_setup")
    bench.model.gate = None
    bench.model.released.set()
    assert job_ended(bench.base, out["id"])["status"] == "done"
    assert edit(bench, setup["id"], tax_id)[0] == 200


def test_a_suggested_wording_is_learned_only_when_asked_for_its_path(start_demo_server, setup_people):
    people = [dataclasses.replace(setup_people[0], country="Estados Unidos"), *setup_people[1:]]
    files = sample_set(people)
    model = model_for(files, people, us_files=["s1.pdf"])
    b = bench_of(start_demo_server, files, model)
    setup = built(b)
    country = field_of(setup, "investor.country")
    assert country["suggested"] == {"Estados Unidos": "US"} and not country["ready"]
    calls = dict(model.calls)
    status, out = get_json(b.base, f"/api/setups/{setup['id']}/learn", "POST", {"path": "investor.tax_id"})
    assert (status, out["error"]) == (400, "bad_request")
    status, after = get_json(b.base, f"/api/setups/{setup['id']}/learn", "POST", {"path": "investor.country"})
    assert status == 200, after
    country = field_of(after, "investor.country")
    assert country["ready"] and country["suggested"] == {}
    values = after["template"]["readers"]["investor.country"]["clean"][0]["args"]["values"]
    assert values["Estados Unidos"] == "US"
    assert model.calls == calls
    status, out = get_json(b.base, f"/api/setups/{setup['id']}/learn", "POST", {"path": 3})
    assert (status, out["error"]) == (400, "bad_request")


# ------------------------------------------------------------------ saved extraction templates


def test_a_setup_saved_as_an_extraction_is_listed_and_reopens_as_a_setup_with_a_fresh_proof(bench):
    setup = built(bench)
    status, saved = get_json(bench.base, "/api/extractions/Client%20info", "PUT", {"setup": setup["id"]})
    assert status == 200 and (saved["name"], saved["template_id"]) == ("Client info", "client_forms")
    assert saved["setup"]["id"] == setup["id"]
    folder = bench.api.settings.data_dir / "local" / "extractions"
    assert (folder / "Client info.json").is_file()
    stored = folder / "assets" / "Client info" / "samples"
    assert [(stored / f"{i}.pdf").read_bytes() for i in range(5)] == bench.pdfs
    status, listed = get_json(bench.base, "/api/extractions")
    assert status == 200
    [entry] = listed["extractions"]
    assert (entry["name"], entry["template_id"], entry["samples"]) == ("Client info", "client_forms", 5)
    assert entry["fields"] == len(concrete_paths(setup["template"]["schema"])) and entry["updated"]
    status, got = get_json(bench.base, "/api/extractions/Client%20info")
    assert status == 200 and got["template"] == {**setup["template"], "name": "Client info"}
    calls = dict(bench.model.calls)
    status, reopened = get_json(bench.base, "/api/setups", "POST", {"extraction": "Client info"})
    assert status == 200, reopened
    assert reopened["id"] != setup["id"] and reopened["batch"] is None
    assert reopened["template"] == got["template"]
    assert reopened["proof"] == setup["proof"]
    assert [s["name"] for s in reopened["samples"]] == [s["name"] for s in setup["samples"]]
    assert [s["index"] for s in reopened["samples"]] == [None] * 5
    assert bench.model.calls == calls
    assert edit(bench, reopened["id"], {"op": "mark_absent", "path": "wire.bank"})[0] == 200


def with_a_field_named_name(template: ExtractionTemplate) -> ExtractionTemplate:
    """``template`` with a top-level field ``document_name``, whose export column is the document's own."""
    schema = json.loads(json.dumps(template.schema_))
    schema["properties"]["document_name"] = {"type": "string", "title": "Document name"}
    return ExtractionTemplate.model_validate({**template.model_dump(by_alias=True), "schema": schema})


def test_a_setup_whose_field_would_share_the_exports_name_column_is_refused_when_saved(bench):
    setup = built(bench)
    kept = bench.api.setups.get(setup["id"])
    kept.template = with_a_field_named_name(kept.template)
    status, out = get_json(bench.base, "/api/extractions/Clash", "PUT", {"setup": setup["id"]})
    assert (status, out["error"]) == (400, "bad_template")
    assert out["detail"] == "the field 'document_name' would share the document column document_name; change its title"
    assert get_json(bench.base, "/api/extractions")[1] == {"extractions": []}  # nothing saved


def test_a_saved_extraction_can_be_deleted_and_names_are_checked(bench):
    setup = built(bench)
    status, out = get_json(bench.base, "/api/extractions/..%2Fx", "PUT", {"setup": setup["id"]})
    assert (status, out["error"]) == (400, "bad_name")
    assert get_json(bench.base, "/api/extractions/..%2Fx")[1]["error"] == "bad_name"
    status, out = get_json(bench.base, "/api/extractions/x", "PUT", {"setup": "0" * 32})
    assert (status, out["error"]) == (404, "not_found")
    assert get_json(bench.base, "/api/extractions/x", "PUT", {"setup": setup["id"]})[0] == 200
    assert get_json(bench.base, "/api/extractions/x", "DELETE") == (200, {"name": "x"})
    for method in ("GET", "DELETE"):
        status, out = get_json(bench.base, "/api/extractions/x", method)
        assert (status, out["error"]) == (404, "not_found") and out["detail"].startswith("no such extraction")
    status, out = get_json(bench.base, "/api/setups", "POST", {"extraction": "x"})
    assert (status, out["error"]) == (404, "not_found")
    assert not (bench.api.settings.data_dir / "local" / "extractions" / "assets" / "x").exists()
    assert get_json(bench.base, "/api/extractions")[1] == {"extractions": []}


def test_the_extraction_store_skips_a_bad_file_in_its_list_and_refuses_to_open_it(tmp_path, setup_pdfs):
    store = extraction.ExtractionStore(tmp_path / "extractions")
    (tmp_path / "extractions").mkdir()
    (tmp_path / "extractions" / "broken.json").write_text("{nope")
    assert store.list() == []
    with pytest.raises(StoredTemplateError):
        store.get("broken")
    with pytest.raises(KeyError):
        store.get("missing")
    with pytest.raises(ValueError):
        store.get("../x")
    pdf = setup_pdfs[0][1]
    template = hand_template().model_copy(update={"samples": [SampleRecord(file="a.pdf", sha256=hashlib.sha256(pdf).hexdigest(), opinion={})]})
    store.put("hand", template, [SamplePdf("a.pdf", pdf)])
    assert store.get("hand") == template and store.samples("hand") == [SamplePdf("a.pdf", pdf)]
    (tmp_path / "extractions" / "assets" / "hand" / "samples" / "0.pdf").unlink()
    with pytest.raises(StoredTemplateError):
        store.samples("hand")


def test_setups_are_kept_in_memory_up_to_the_limit_least_recently_used_dropped_first():
    setups = extraction.Setups(limit=2)
    made = [extraction.Setup(f"{n:032x}", None, None, [], None, None, []) for n in range(3)]
    setups.add(made[0])
    setups.add(made[1])
    assert setups.get(made[0].id) is made[0]  # now the most recently used
    setups.add(made[2])
    assert setups.get(made[0].id) is made[0] and setups.get(made[2].id) is made[2]
    with pytest.raises(KeyError):
        setups.get(made[1].id)


# ------------------------------------------------------------------ runs


PROBLEM_FILES = ("bad_tin.pdf", "no_tin.pdf", "mars.pdf")  # a bad tax ID, no tax ID, a country not in the map
ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="module")
def corpus_files(tmp_path_factory) -> list[tuple[str, bytes]]:
    """The corpus folder's 22 PDFs (``client_00001.pdf`` ...): 20 filled client forms, then 2 copies of the W-9."""
    folder = write_corpus(tmp_path_factory.mktemp("runs") / "corpus")
    return [(p.name, p.read_bytes()) for p in sorted(folder.glob("*.pdf"))]


@pytest.fixture(scope="module")
def problem_files() -> list[tuple[str, bytes]]:
    """Three filled client forms with one problem each (``PROBLEM_FILES``)."""
    base = clients(1, seed=5)[0]
    copies = [
        dataclasses.replace(base, tax_id="8422-19037"),
        dataclasses.replace(base, tax_id=""),
        dataclasses.replace(base, country="Mars"),
    ]
    return [(name, make_form(client, kind="flattened")) for name, client in zip(PROBLEM_FILES, copies)]


@dataclass
class Runner:
    """A server (``extract_workers=1``) with the extraction ``Client info`` saved from a setup, as Task 4 makes it."""

    base: str
    api: Api
    model: GatedModel
    template: dict  # the saved template, as GET /api/extractions/Client%20info gives it


@pytest.fixture
def runner(start_demo_server, setup_pdfs, setup_people) -> Runner:
    model = model_for(setup_pdfs, setup_people)
    base = start_demo_server(provider_factory=model, extract_workers=1)
    api = start_demo_server.servers[-1].api
    samples = new_batch(base, name="Samples")
    assert add_files(base, samples["id"], setup_pdfs)[0] == 200
    setup = built(Bench(base, api, model, samples["id"], [data for _, data in setup_pdfs]))
    status, saved = get_json(base, "/api/extractions/Client%20info", "PUT", {"setup": setup["id"]})
    assert status == 200, saved
    return Runner(base, api, model, get_json(base, "/api/extractions/Client%20info")[1]["template"])


def batch_of(r: Runner, files: list[tuple[str, bytes]], name: str = "Client forms") -> str:
    batch = new_batch(r.base, name=name)
    if files:
        status, out = add_files(r.base, batch["id"], files)
        assert status == 200, out
    return batch["id"]


def start_run(r: Runner, batch: str, **body) -> tuple[int, dict]:
    return start_job(r.base, {"kind": "extract-run", "extraction": "Client info", "batch": batch, **body})


def run_on(r: Runner, batch: str, **body) -> dict:
    """The ``extract-run`` job on ``batch``, once it has ended."""
    status, out = start_run(r, batch, **body)
    assert status == 200, out
    return job_ended(r.base, out["id"])


def read_fields(template: dict) -> list[str]:
    """The template's concrete paths it reads (not marked absent), in schema order."""
    return [p for p in concrete_paths(template["schema"]) if p not in template["absent"]]


def runs_dir(r: Runner) -> Path:
    return r.api.settings.data_dir / "local" / "runs"


@pytest.fixture
def reviewed(runner, corpus_files, problem_files) -> tuple[Runner, str]:
    """(the runner, a finished run's id) of the corpus folder's PDFs and the three problem copies: 25 documents."""
    job = run_on(runner, batch_of(runner, corpus_files + problem_files))
    assert job["status"] == "done", job
    return runner, job["result"]["run"]


def problems_of(r: Runner, run_id: str, query: str = "") -> tuple[int, dict]:
    return get_json(r.base, f"/api/runs/{run_id}/problems{query}")


def test_a_run_reads_every_document_of_the_batch_with_no_model_call_lists_it_and_counts_each_field(runner, corpus_files):
    calls = dict(runner.model.calls)
    batch = batch_of(runner, corpus_files)
    job = run_on(runner, batch)
    assert job["status"] == "done", job
    assert [(s["key"], s["label"], s["status"]) for s in job["stages"]] == [("read", "Read every document", "done")]
    result = job["result"]
    assert result["done"] == result["total"] == 22 and result["summary"]["total"] == 22
    assert result["summary"]["no_template"] == 2 and result["summary"]["complete"] + result["summary"]["with_problems"] == 20
    assert runner.model.calls == calls and job["usage"] == {"input_tokens": 0, "output_tokens": 0}
    run_id = result["run"]
    status, listed = get_json(runner.base, "/api/runs")
    assert status == 200
    [entry] = listed["runs"]
    assert (entry["id"], entry["name"], entry["extraction"]) == (run_id, "Client forms", "Client info")
    assert entry["started"] and entry["summary"] == result["summary"]
    status, run = get_json(runner.base, f"/api/runs/{run_id}")
    assert status == 200, run
    assert (run["id"], run["name"], run["extraction"], run["summary"]) == (run_id, "Client forms", "Client info", result["summary"])
    assert run["seconds"] == run["summary"]["seconds"] > 0
    assert [f["path"] for f in run["fields"]] == read_fields(runner.template)
    tax_id = next(f for f in run["fields"] if f["path"] == "investor.tax_id")
    assert set(tax_id) == {"path", "title", "required", "required_reason", "found", "blank", "not_located", "missing", "failed_check", "unsure", "reviewed", "not_on_form", "absent"}
    assert tax_id["title"] == runner.template["schema"]["properties"]["investor"]["properties"]["tax_id"]["title"]
    assert tax_id["found"] == 20 and tax_id["blank"] == tax_id["failed_check"] == tax_id["unsure"] == tax_id["reviewed"] == 0
    folder = runs_dir(runner) / run_id
    meta = json.loads((folder / "run-meta.json").read_text())
    assert meta == {"name": "Client forms", "extraction": "Client info", "batch": batch, "started": entry["started"]}
    assert (folder / "run.json").is_file() and len((folder / "results.jsonl").read_text().splitlines()) == 22


def test_a_run_is_read_back_from_its_folder_after_a_restart_and_an_unknown_run_is_not_found(runner, corpus_files):
    run_id = run_on(runner, batch_of(runner, corpus_files[:3]))["result"]["run"]
    before = get_json(runner.base, f"/api/runs/{run_id}")[1]
    again = extraction.Runs(runs_dir(runner))  # what a restarted server has
    assert again.get(run_id).summary() == before
    assert [entry["id"] for entry in again.list()] == [run_id]
    with pytest.raises(KeyError):
        again.get("0" * 32)
    status, out = get_json(runner.base, f"/api/runs/{'0' * 32}")
    assert (status, out["error"]) == (404, "not_found") and "no such run" in out["detail"]


def test_a_run_is_refused_before_it_starts_for_an_unknown_extraction_or_batch_or_ai_with_no_model(runner, corpus_files, demo_server):
    batch = batch_of(runner, corpus_files[:2])
    status, out = start_job(runner.base, {"kind": "extract-run", "extraction": "Nope", "batch": batch})
    assert (status, out["error"]) == (404, "not_found") and "no such extraction" in out["detail"]
    status, out = start_job(runner.base, {"kind": "extract-run", "extraction": "../x", "batch": batch})
    assert (status, out["error"]) == (400, "bad_name")
    status, out = start_job(runner.base, {"kind": "extract-run", "batch": batch})
    assert (status, out["error"]) == (400, "bad_request")
    status, out = start_run(runner, "f" * 32)
    assert (status, out["error"]) == (400, "bad_batch")
    status, out = start_run(runner, batch_of(runner, [], name="Empty"))
    assert (status, out["error"]) == (400, "bad_batch")
    status, out = start_run(runner, batch, ai="yes")
    assert (status, out["error"]) == (400, "bad_request")
    status, out = start_job(demo_server, {"kind": "extract-run", "extraction": "Client info", "batch": batch, "ai": True})
    assert (status, out["error"]) == (400, "ai_not_configured")
    # demo_server has no model and shares the data folder (and so the extraction), not the batches, which are in memory
    status, out = start_job(demo_server, {"kind": "extract-run", "extraction": "Client info", "batch": batch})
    assert (status, out["error"]) == (400, "bad_batch")  # a run without the AI needs no model: it got as far as the batch
    assert get_json(runner.base, "/api/runs")[1] == {"runs": []}


def test_a_field_that_would_share_the_exports_name_column_can_be_renamed_and_the_setup_saved(bench):
    setup = built(bench)
    kept = bench.api.setups.get(setup["id"])
    kept.template = with_a_field_named_name(kept.template)
    status, after = edit(bench, setup["id"], {"op": "rename_field", "path": "document_name", "name": "legal_name"})
    assert status == 200, after
    assert "legal_name" in after["template"]["schema"]["properties"]
    assert "legal_name" in [f["path"] for f in after["proof"]["fields"]]
    status, out = get_json(bench.base, "/api/extractions/Clash", "PUT", {"setup": setup["id"]})
    assert status == 200, out


def test_a_saved_templates_title_edit_keeps_the_name_and_offers_the_update(bench):
    setup = built(bench)
    kept = bench.api.setups.get(setup["id"])
    info = kept.template.setup.model_copy(update={"name_follows_title": True})
    kept.template = kept.template.model_copy(update={"setup": info})
    status, after = edit(bench, setup["id"], {"op": "set_field", "path": "investor.tax_id", "title": "TIN number"})
    assert status == 200, after
    assert after["renamed"] == {"investor.tax_id": "investor.tin_number"} and after["name_updates"] == {}
    status, out = get_json(bench.base, "/api/extractions/Kept", "PUT", {"setup": setup["id"]})
    assert status == 200, out
    assert bench.api.extractions.get("Kept").setup.name_follows_title is False
    status, after = edit(bench, setup["id"], {"op": "set_field", "path": "investor.tin_number", "title": "Born on"})
    assert status == 200, after
    assert "investor.tin_number" in [f["path"] for f in after["proof"]["fields"]]
    assert after["name_updates"]["investor.tin_number"] == "born_on"
    status, after = edit(bench, setup["id"], {"op": "rename_field", "path": "investor.tin_number", "name": "born_on"})
    assert status == 200, after
    assert "investor.born_on" not in after["name_updates"]


def test_update_name_is_not_offered_for_a_name_that_is_taken_or_reserved():
    schema = {
        "type": "object",
        "properties": {
            "date": {"type": "string", "title": "Date"},
            "date_2": {"type": "string", "title": "Date"},
            "document_id_2": {"type": "string", "title": "Document ID"},
            "own": {
                "type": "array", "maxItems": 2,
                "items": {"type": "object", "properties": {"x": {"type": "string", "title": "Owner name"}}},
            },
        },
    }  # fmt: skip
    template = hand_template().model_copy(update={"schema_": schema, "readers": {}, "absent": []})
    assert extraction._name_updates(template) == {"own[0].x": "owner_name", "own[1].x": "owner_name"}


def test_a_rename_a_field_cannot_take_is_refused_with_the_librarys_message(bench):
    setup = built(bench)
    status, out = edit(bench, setup["id"], {"op": "rename_field", "path": "investor.tax_id", "name": "tax id"})
    assert (status, out["error"]) == (400, "bad_edit")
    assert "letters, numbers" in out["detail"]


def test_a_run_of_an_extraction_whose_field_would_share_the_exports_name_column_is_refused_before_it_starts(runner, corpus_files):
    template = ExtractionTemplate.model_validate(runner.template)
    samples = runner.api.extractions.samples("Client info")
    runner.api.extractions.put("Clash", with_a_field_named_name(template).model_copy(update={"name": "Clash"}), samples)
    status, out = start_job(runner.base, {"kind": "extract-run", "extraction": "Clash", "batch": batch_of(runner, corpus_files[:2])})
    assert (status, out["error"]) == (400, "bad_template")
    assert out["detail"] == "the field 'document_name' would share the document column document_name; change its title"
    assert get_json(runner.base, "/api/runs")[1] == {"runs": []}


def test_the_problems_are_paged_by_document_each_with_its_problems_titled_and_named(reviewed):
    r, run_id = reviewed
    status, page = problems_of(r, run_id, "?limit=2")
    assert status == 200, page
    assert page["total"] == 3 and len(page["documents"]) == 2
    titles = {f["path"]: f["title"] for f in get_json(r.base, f"/api/runs/{run_id}")[1]["fields"]}
    for doc in page["documents"]:
        assert set(doc) == {"doc_id", "name", "problems"} and doc["name"] in PROBLEM_FILES and doc["problems"]
        for problem in doc["problems"]:
            assert set(problem) == {"path", "title", "required", "required_reason", "status", "raw", "value", "reasons", "page", "rect"}
            assert problem["status"] in ("blank", "failed_check", "unsure") and problem["title"] == titles[problem["path"]]
    rest = problems_of(r, run_id, "?offset=2&limit=2")[1]
    assert rest["total"] == 3 and len(rest["documents"]) == 1
    assert {d["name"] for d in page["documents"] + rest["documents"]} == set(PROBLEM_FILES)
    assert problems_of(r, run_id, "?offset=3")[1] == {"total": 3, "documents": []}
    assert problems_of(r, run_id, "?offset=50&limit=2")[1] == {"total": 3, "documents": []}
    assert problems_of(r, run_id, "?path=investor.tax_id")[1]["total"] == 2
    missing = problems_of(r, run_id, "?status=blank")[1]
    assert missing["total"] == 1 and missing["documents"][0]["name"] == "no_tin.pdf"
    for query in ("?limit=0", "?limit=201", "?offset=-1", "?status=found"):
        status, out = problems_of(r, run_id, query)
        assert (status, out["error"]) == (400, "bad_request"), query


def test_places_say_where_a_field_left_missing_should_be_and_only_the_fields_with_no_place(reviewed):
    r, run_id = reviewed
    documents = problems_of(r, run_id)[1]["documents"]
    missing = next(d for d in documents if d["name"] == "no_tin.pdf")
    read = next(d for d in documents if d["name"] == "bad_tin.pdf")
    assert [(p["path"], p["page"], p["rect"]) for p in missing["problems"]] == [("investor.tax_id", None, None)]
    status, out = get_json(r.base, f"/api/runs/{run_id}/places?doc={quote(missing['doc_id'])}")
    assert status == 200, out
    place = out["places"]["investor.tax_id"]
    x0, y0, x1, y1 = place["rect"]
    assert set(out["places"]) == {"investor.tax_id"} and place["page"] == 0 and x0 < x1 and y0 < y1
    assert get_json(r.base, f"/api/runs/{run_id}/places?doc={quote(read['doc_id'])}") == (200, {"places": {}})  # read: it has one
    assert get_json(r.base, f"/api/runs/{run_id}/places?doc=nope")[0] == 404
    assert get_json(r.base, f"/api/runs/{run_id}/places")[0] == 400


def _reasoned(r: Runner) -> None:
    """The saved extraction with a reason under its tax ID (required) and its country (optional), as the agent's
    classification leaves them; a run made after this reads them."""
    template = r.api.extractions.get("Client info")
    investor = template.schema_["properties"]["investor"]["properties"]
    investor["tax_id"]["x-required-reason"] = "the form marks it required"
    investor["country"]["x-required-reason"] = "a copy may leave it empty"
    r.api.extractions.put("Client info", template, r.api.extractions.samples("Client info"))


@pytest.fixture
def leveled(runner, corpus_files, problem_files) -> tuple[Runner, str]:
    """Like ``reviewed``, with the extraction's reasons set: (the runner, a finished run's id)."""
    _reasoned(runner)
    job = run_on(runner, batch_of(runner, corpus_files + problem_files))
    assert job["status"] == "done", job
    return runner, job["result"]["run"]


def test_problems_by_level_list_only_that_levels_fields(leveled):
    r, run_id = leveled
    both = problems_of(r, run_id)[1]
    assert both["total"] == 3
    for doc in both["documents"]:
        for problem in doc["problems"]:
            assert problem["required"] == (problem["path"] == "investor.tax_id")
    required = problems_of(r, run_id, "?level=required")[1]
    assert required["total"] == 2 and {d["name"] for d in required["documents"]} == {"bad_tin.pdf", "no_tin.pdf"}
    assert {(p["path"], p["required"]) for d in required["documents"] for p in d["problems"]} == {("investor.tax_id", True)}
    optional = problems_of(r, run_id, "?level=optional")[1]
    [doc] = optional["documents"]
    assert optional["total"] == 1 and doc["name"] == "mars.pdf"
    assert [(p["path"], p["required"], p["status"]) for p in doc["problems"]] == [("investor.country", False, "failed_check")]
    assert problems_of(r, run_id, "?level=optional&path=investor.tax_id")[1] == {"total": 0, "documents": []}
    assert problems_of(r, run_id, "?level=required&status=blank")[1]["total"] == 1


def test_a_bad_level_is_a_bad_request(leveled):
    r, run_id = leveled
    for query in ("?level=maybe", "?level=Required", "?level=both"):
        status, out = problems_of(r, run_id, query)
        assert (status, out["error"]) == (400, "bad_request"), query


def test_the_runs_fields_say_required_and_why(leveled):
    r, run_id = leveled
    fields = {f["path"]: f for f in get_json(r.base, f"/api/runs/{run_id}")[1]["fields"]}
    assert (fields["investor.tax_id"]["required"], fields["investor.tax_id"]["required_reason"]) == (True, "the form marks it required")
    assert (fields["investor.country"]["required"], fields["investor.country"]["required_reason"]) == (False, "a copy may leave it empty")
    assert fields["investor.legal_name"]["required_reason"] == ""  # a field the agent gave no reason for says none
    [mars] = [d for d in problems_of(r, run_id, "?level=optional")[1]["documents"]]
    assert mars["problems"][0]["required_reason"] == "a copy may leave it empty"
    [tax_id] = [p for d in problems_of(r, run_id, "?status=blank")[1]["documents"] for p in d["problems"]]
    assert tax_id["required_reason"] == "the form marks it required"


def test_the_summary_counts_optional_only_documents(leveled):
    r, run_id = leveled
    summary = get_json(r.base, f"/api/runs/{run_id}")[1]["summary"]
    assert summary["with_problems"] == 2  # bad_tin and no_tin; mars is complete
    assert (summary["optional_only"], summary["with_optional"]) == (1, 1)  # mars's country
    [entry] = [e for e in get_json(r.base, "/api/runs")[1]["runs"] if e["id"] == run_id]
    assert entry["summary"]["optional_only"] == 1
    [doc] = problems_of(r, run_id, "?level=optional")[1]["documents"]
    decision = {"doc_id": doc["doc_id"], "path": "investor.country", "decision": "not_in_document"}
    after = get_json(r.base, f"/api/runs/{run_id}/decisions", "POST", decision)[1]["summary"]
    assert (after["optional_only"], after["with_optional"], after["with_problems"]) == (0, 0, 2)


def test_to_data_counts_only_required_problems(leveled):
    r, run_id = leveled
    out = get_json(r.base, f"/api/runs/{run_id}/to-data", "POST", {})[1]
    assert out["with_problems"] == 2  # mars's country is optional: its row is complete
    [doc] = problems_of(r, run_id, "?level=required&status=blank")[1]["documents"]
    decision = {"doc_id": doc["doc_id"], "path": "investor.tax_id", "decision": "value", "value": "12-3456789"}
    assert get_json(r.base, f"/api/runs/{run_id}/decisions", "POST", decision)[0] == 200
    assert get_json(r.base, f"/api/runs/{run_id}/to-data", "POST", {})[1]["with_problems"] == 1


def test_a_run_made_before_the_levels_were_kept_counts_every_field_as_required(leveled):
    r, run_id = leveled
    run_json = runs_dir(r) / run_id / "run.json"
    kept = json.loads(run_json.read_text())
    del kept["required"]
    run_json.write_text(json.dumps(kept))
    r.api.runs._runs.pop(run_id)  # as after a restart: the run is read from its folder again
    out = get_json(r.base, f"/api/runs/{run_id}")[1]
    assert out["summary"]["with_problems"] == 3 and out["summary"]["optional_only"] == 0
    assert {f["required"] for f in out["fields"]} == {True}
    assert problems_of(r, run_id, "?level=optional")[1]["total"] == 0


def test_the_set_aside_documents_are_the_corpus_folders_w9_copies_with_the_closest_template(reviewed):
    r, run_id = reviewed
    status, out = get_json(r.base, f"/api/runs/{run_id}/set-aside")
    assert status == 200, out
    assert out["total"] == 2
    assert [d["name"] for d in out["documents"]] == ["client_00021.pdf", "client_00022.pdf"]
    for doc in out["documents"]:
        assert set(doc) == {"doc_id", "name", "reason", "closest_template", "closest_score"}
        assert (doc["reason"], doc["closest_template"]) == ("no template fits", "Client info")
        assert isinstance(doc["closest_score"], float) and 0 <= doc["closest_score"] < 1
    assert get_json(r.base, f"/api/runs/{run_id}/set-aside?offset=1&limit=1")[1]["documents"] == out["documents"][1:]
    assert get_json(r.base, f"/api/runs/{run_id}/set-aside?limit=201")[0] == 400


def test_a_decision_settles_its_problem_and_the_document_when_it_was_its_only_one(reviewed):
    r, run_id = reviewed
    before = get_json(r.base, f"/api/runs/{run_id}")[1]
    [doc] = problems_of(r, run_id, "?status=blank")[1]["documents"]
    [problem] = doc["problems"]
    decision = {"doc_id": doc["doc_id"], "path": problem["path"], "decision": "value", "value": "12-3456789"}
    status, after = get_json(r.base, f"/api/runs/{run_id}/decisions", "POST", decision)
    assert status == 200, after
    assert after["summary"]["with_problems"] == before["summary"]["with_problems"] - 1
    assert after["summary"]["complete"] == before["summary"]["complete"] + 1
    assert next(f for f in after["fields"] if f["path"] == problem["path"])["reviewed"] == 1
    left = problems_of(r, run_id)[1]
    assert left["total"] == 2 and doc["doc_id"] not in [d["doc_id"] for d in left["documents"]]
    assert get_json(r.base, f"/api/runs/{run_id}")[1]["summary"] == after["summary"]


def test_a_runs_results_are_read_once_and_a_decision_a_page_image_and_a_problems_page_after_it_read_nothing_again(reviewed, monkeypatch):
    from formcast.extract.batch import Results

    r, run_id = reviewed
    reads = {"loads": 0, "lines": 0}  # every file read from the start; one document's line read back
    read_all, line = Results._read_all, Results._line

    def counted_read_all(self):
        reads["loads"] += 1
        read_all(self)

    def counted_line(self, doc_id):
        reads["lines"] += 1
        return line(self, doc_id)

    monkeypatch.setattr(Results, "_read_all", counted_read_all)
    monkeypatch.setattr(Results, "_line", counted_line)
    r.api.runs._runs.clear()  # as a restarted server: the run is read from its folder when first asked for
    page = problems_of(r, run_id)[1]
    assert reads["loads"] == 1 and page["total"] == 3
    get_json(r.base, f"/api/runs/{run_id}")
    get_json(r.base, "/api/runs")
    assert reads["loads"] == 1  # the listing and the run use the kept counts

    [doc] = problems_of(r, run_id, "?status=blank")[1]["documents"]
    reads["lines"] = 0
    decision = {"doc_id": doc["doc_id"], "path": "investor.tax_id", "decision": "value", "value": "12-3456789"}
    status, after = get_json(r.base, f"/api/runs/{run_id}/decisions", "POST", decision)
    assert status == 200 and after["summary"]["with_problems"] == 1  # bad_tin; mars's country is optional
    status, _, _ = http_request(r.base, "GET", f"/api/runs/{run_id}/pages.png?doc={doc['doc_id']}&page=0&dpi=40")
    assert status == 200
    page = problems_of(r, run_id, "?limit=1")[1]
    assert page["total"] == 2 and len(page["documents"]) == 1
    assert get_json(r.base, f"/api/runs/{run_id}/set-aside?limit=1")[1]["total"] == 2
    assert reads["loads"] == 1  # nothing was read again from the start
    assert reads["lines"] <= 6  # the decided document (checked, taken back, counted), the page image's, the pages'

    other = problems_of(r, run_id)[1]["documents"][0]
    with (runs_dir(r) / run_id / "reviews.jsonl").open("a") as sink:  # another process settles that document
        for problem in other["problems"]:
            sink.write(json.dumps({"doc_id": other["doc_id"], "path": problem["path"], "decision": "not_in_document"}) + "\n")
    assert get_json(r.base, f"/api/runs/{run_id}")[1]["summary"]["with_problems"] == 0  # mars: optional only
    assert problems_of(r, run_id)[1]["total"] == 1
    assert reads["loads"] == 1  # the appended lines are read on their own


def test_a_decision_the_library_refuses_is_a_bad_decision_with_its_message(reviewed):
    r, run_id = reviewed
    [doc] = problems_of(r, run_id, "?status=blank")[1]["documents"]
    for body, says in [
        ({"doc_id": "nope.pdf", "path": "investor.tax_id", "decision": "value", "value": "1"}, "no document"),
        ({"doc_id": doc["doc_id"], "path": "nope", "decision": "value", "value": "1"}, "no field"),
        ({"doc_id": doc["doc_id"], "path": "investor.tax_id", "decision": "accept"}, "nothing to accept"),
        ({"doc_id": doc["doc_id"], "path": "investor.tax_id", "decision": "value"}, "needs a value"),
        ({"doc_id": doc["doc_id"], "path": "investor.tax_id", "decision": "maybe"}, "decision"),
    ]:
        status, out = get_json(r.base, f"/api/runs/{run_id}/decisions", "POST", body)
        assert (status, out["error"]) == (400, "bad_decision"), body
        assert says in out["detail"], out
    assert problems_of(r, run_id)[1]["total"] == 3


def test_a_page_of_a_document_of_the_run_is_a_png_and_anything_else_is_not_found(reviewed):
    r, run_id = reviewed
    [doc] = problems_of(r, run_id, "?status=blank")[1]["documents"]
    status, headers, body = http_request(r.base, "GET", f"/api/runs/{run_id}/pages.png?doc={doc['doc_id']}&page=0&dpi=40")
    assert status == 200 and headers["content-type"] == "image/png"
    assert png_size(body)[0] == round(612 * 40 / 72)
    for query in ("?doc=../etc/passwd&page=0", "?doc=..%2F..%2Fetc%2Fpasswd", f"?doc={doc['doc_id']}&page=9", "?doc=nope.pdf"):
        status, _, out = http_request(r.base, "GET", f"/api/runs/{run_id}/pages.png{query}")
        assert status == 404, query
        assert str(r.api.settings.data_dir) not in json.loads(out)["detail"]
    assert http_request(r.base, "GET", f"/api/runs/{run_id}/pages.png?page=0")[0] == 400


def test_a_run_exports_as_a_workbook_or_a_zip_of_csv_files(reviewed):
    import openpyxl

    r, run_id = reviewed
    status, headers, body = http_request(r.base, "GET", f"/api/runs/{run_id}/export?format=xlsx")
    assert status == 200, body[:200]
    assert headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert headers["content-disposition"] == 'attachment; filename="Client forms-results.xlsx"'
    book = openpyxl.load_workbook(io.BytesIO(body), read_only=True)
    assert "values" in book.sheetnames
    values = list(book["values"].iter_rows(values_only=True))
    assert len(values) == 1 + 23  # the header, then each document read
    assert list(values[0][:3]) == ["document_id", "document_name", "document_file"]
    shown = {row[1] for row in values[1:]}
    assert set(PROBLEM_FILES) <= shown and "client_00001.pdf" in shown  # the uploaded files' own names
    assert all(row[2] == row[1] for row in values[1:])  # an uploaded batch's file is its name, not where it is stored
    aside = list(book["set_aside"].iter_rows(values_only=True))
    assert list(aside[0][:3]) == ["document_id", "document_name", "document_file"] and [row[1] for row in aside[1:]] == ["client_00021.pdf", "client_00022.pdf"]
    cells = [str(c) for sheet in book.sheetnames for row in book[sheet].iter_rows(values_only=True) for c in row if c is not None]
    assert not [c for c in cells if "local/batches" in c or str(r.api.settings.data_dir) in c]  # no data-dir path
    status, headers, body = http_request(r.base, "GET", f"/api/runs/{run_id}/export?format=csv")
    assert status == 200 and headers["content-type"] == "application/zip"
    assert headers["content-disposition"] == 'attachment; filename="Client forms-results.zip"'
    with zipfile.ZipFile(io.BytesIO(body)) as z:
        assert sorted(z.namelist()) == ["fields.csv", "set_aside.csv", "status.csv", "values.csv"]
        assert len(z.read("values.csv").decode("utf-8").splitlines()) == 1 + 23
    status, _, body = http_request(r.base, "GET", f"/api/runs/{run_id}/export?format=pdf")
    assert (status, json.loads(body)["error"]) == (400, "bad_request")
    assert [p.name for p in runs_dir(r).iterdir()] == [run_id]  # the export's files were not left behind


def test_a_parquet_export_is_a_zip_or_not_supported_without_pyarrow(reviewed):
    r, run_id = reviewed
    status, headers, body = http_request(r.base, "GET", f"/api/runs/{run_id}/export?format=parquet")
    if importlib.util.find_spec("pyarrow") is None:
        out = json.loads(body)
        assert (status, out["error"]) == (400, "not_supported") and "formcast[parquet]" in out["detail"]
    else:
        assert status == 200 and headers["content-type"] == "application/zip"
        assert "values.parquet" in zipfile.ZipFile(io.BytesIO(body)).namelist()


def test_a_run_sent_to_fill_is_a_dataset_of_the_found_values_of_the_read_documents_by_schema_path(reviewed):
    r, run_id = reviewed
    status, out = get_json(r.base, f"/api/runs/{run_id}/to-data", "POST", {})
    assert status == 200, out
    dataset = out["dataset"]
    assert out["truncated"] is False and dataset["truncated"] is False
    assert (dataset["name"], dataset["kind"], dataset["format"]) == ("Client forms results", "table", "json")
    assert [c["name"] for c in dataset["columns"]] == read_fields(r.template)
    assert dataset["rows"] == 23  # every document read; the two set aside are not rows
    assert out["with_problems"] == 2  # bad_tin and no_tin; mars's country is optional, so its row is complete
    rows = r.api.datasets.get(dataset["id"]).table.rows
    tax_ids = [row["investor.tax_id"] for row in rows]
    assert "8422-19037" not in tax_ids and tax_ids.count(None) == 2  # a value that failed its check is a gap, not a cell
    assert [row["investor.country"] for row in rows].count(None) == 1  # nor is mars's country, which is not in the map
    assert get_json(r.base, f"/api/data/{dataset['id']}")[1]["rows"] == 23
    [doc] = problems_of(r, run_id, "?status=blank")[1]["documents"]
    decision = {"doc_id": doc["doc_id"], "path": "investor.tax_id", "decision": "value", "value": "12-3456789"}
    assert get_json(r.base, f"/api/runs/{run_id}/decisions", "POST", decision)[0] == 200
    out = get_json(r.base, f"/api/runs/{run_id}/to-data", "POST", {})[1]
    assert out["with_problems"] == 1
    tax_ids = [row["investor.tax_id"] for row in r.api.datasets.get(out["dataset"]["id"]).table.rows]
    assert "12-3456789" in tax_ids and tax_ids.count(None) == 1  # the typed value is found


def test_a_run_sent_to_fill_keeps_the_first_rows_when_it_has_more_than_a_dataset_holds(reviewed, monkeypatch):
    r, run_id = reviewed
    monkeypatch.setattr(extraction, "MAX_ROWS", 5)
    status, out = get_json(r.base, f"/api/runs/{run_id}/to-data", "POST", {})
    assert status == 200, out
    assert out["truncated"] is True and out["dataset"]["truncated"] is True and out["dataset"]["rows"] == 5


def test_a_run_with_the_ai_asks_it_only_about_the_documents_whose_required_fields_failed(runner, problem_files):
    calls = runner.model.calls.get("values", 0)
    job = run_on(runner, batch_of(runner, problem_files), ai=True)
    assert job["status"] == "done", job
    assert [(s["key"], s["label"], s["status"]) for s in job["stages"]] == [
        ("read", "Read every document", "done"), ("ai", "Ask the AI about the failing documents", "done"),
    ]  # fmt: skip
    assert runner.model.calls["values"] - calls == 2  # bad_tin and no_tin: the tax ID is the one required field
    assert job["usage"]["input_tokens"] == 2 * 10 and job["result"]["summary"]["total"] == 3
    assert (runs_dir(runner) / job["result"]["run"] / "ai.jsonl").is_file()


def test_a_run_cancelled_while_it_asks_the_ai_ends_cancelled_with_its_summary(runner, problem_files):
    runner.model.KINDS = {**GatedModel.KINDS, VALUES_SYSTEM: "values"}
    runner.model.gate = "values"
    status, out = start_run(runner, batch_of(runner, problem_files), ai=True)
    assert status == 200, out
    assert runner.model.reached.wait(20)
    assert get_json(runner.base, f"/api/jobs/{out['id']}/cancel", "POST", {})[0] == 200
    runner.model.released.set()
    job = job_ended(runner.base, out["id"])
    assert job["status"] == "cancelled" and job["result"]["summary"]["total"] == 3
    assert [s["status"] for s in job["stages"]] == ["done", "skipped"]


def test_a_second_run_while_the_jobs_slots_are_taken_is_busy_and_starts_nothing(runner, corpus_files):
    batch = batch_of(runner, corpus_files[:2])
    release = threading.Event()
    held = [runner.api.jobs.start_task("extract-run", lambda job: release.wait(10)) for _ in range(2)]
    try:
        status, out = start_run(runner, batch)
        assert (status, out["error"]) == (429, "busy")
        assert out["detail"] == "2 jobs are already running; wait for one to finish or cancel it."
        assert get_json(runner.base, "/api/runs")[1] == {"runs": []}
        assert not runs_dir(runner).exists() or not list(runs_dir(runner).iterdir())
    finally:
        release.set()
    for job in held:
        wait_ended(runner.api.jobs, job.id)
    assert run_on(runner, batch)["status"] == "done"


def test_a_run_cancelled_after_three_documents_keeps_them_and_a_new_run_of_the_batch_resumes_it(runner, corpus_files, monkeypatch):
    batch = batch_of(runner, corpus_files)
    reached, go = threading.Event(), threading.Event()
    real = extract_api.run_batch

    def gated(*args, on_progress, **kwargs):
        def progress(done, total):
            on_progress(done, total)
            if done == 3:
                reached.set()
                assert go.wait(10)

        return real(*args, on_progress=progress, **kwargs)

    monkeypatch.setattr(extract_api, "run_batch", gated)
    status, out = start_run(runner, batch)
    assert status == 200, out
    assert reached.wait(20)
    job = job_of(runner.base, out["id"])
    assert job["status"] == "running" and (job["result"]["done"], job["result"]["total"]) == (3, 22)
    run_id = job["result"]["run"]
    status, busy = start_run(runner, batch)  # one job per run at a time
    assert (status, busy["error"]) == (409, "busy_run")
    assert get_json(runner.base, f"/api/jobs/{out['id']}/cancel", "POST", {})[0] == 200
    go.set()
    job = job_ended(runner.base, out["id"])
    assert job["status"] == "cancelled" and job["result"]["run"] == run_id and job["result"]["done"] == 3
    assert get_json(runner.base, f"/api/runs/{run_id}")[1]["summary"]["total"] == 3
    entered, resume = threading.Event(), threading.Event()

    def held(*args, **kwargs):
        entered.set()
        assert resume.wait(10)
        return real(*args, **kwargs)

    monkeypatch.setattr(extract_api, "run_batch", held)
    status, out = start_run(runner, batch)
    assert status == 200, out
    assert entered.wait(20)
    job = job_of(runner.base, out["id"])
    assert (job["result"]["run"], job["result"]["done"], job["result"]["total"]) == (run_id, 3, 22)  # from where it was
    resume.set()
    again = job_ended(runner.base, out["id"])
    assert again["status"] == "done" and again["result"]["run"] == run_id
    assert again["result"]["summary"]["total"] == 22 and again["result"]["done"] == 22
    lines = (runs_dir(runner) / run_id / "results.jsonl").read_text().splitlines()
    ids = [json.loads(line)["doc_id"] for line in lines]
    assert len(ids) == 22 and len(set(ids)) == 22
    assert [entry["id"] for entry in get_json(runner.base, "/api/runs")[1]["runs"]] == [run_id]


@pytest.mark.parametrize(("setting", "cpus", "workers"), [(0, 8, 8), (3, 8, 3), (16, 8, 8), (0, None, 1)])
def test_a_runs_workers_are_one_per_cpu_by_default_and_never_more_than_the_cpus(monkeypatch, setting, cpus, workers):
    monkeypatch.setattr(extract_api.os, "cpu_count", lambda: cpus)
    settings = Settings.from_env({})
    settings.extract_workers = setting
    assert extract_api._workers(settings) == workers


def test_a_document_set_aside_by_an_internal_error_shows_its_kind_but_not_the_exceptions_words():
    assert extract_api._reason("internal error: KeyError: '/srv/secret/a.pdf'") == "internal error: KeyError"
    assert extract_api._reason("cannot read PDF: the reader crashed") == "cannot read PDF: the reader crashed"
    assert extract_api._reason("no template fits") == "no template fits"


def cli_run(template: dict, folder: Path, out: Path, tmp_path: Path) -> None:
    """``python -m formcast.extract run`` with ``template`` (written as JSON) over ``folder`` into ``out``."""
    path = tmp_path / f"{template['template_id']}.json"
    path.write_text(json.dumps(template))
    done = subprocess.run(
        [sys.executable, "-m", "formcast.extract", "run", str(path), str(folder), "--out", str(out), "--workers", "1"],
        capture_output=True, text=True, timeout=120, cwd=ROOT,
        env={**os.environ, "PYTHONPATH": os.pathsep.join([str(ROOT / "src"), str(ROOT)])},
    )  # fmt: skip
    assert done.returncode == 0, done.stderr


def folder_of(files: list[tuple[str, bytes]], folder: Path) -> Path:
    folder.mkdir()
    for name, data in files:
        (folder / name).write_bytes(data)
    return folder


def test_a_command_line_run_is_opened_by_its_folder_when_its_template_is_a_saved_extraction(runner, corpus_files, tmp_path):
    cli_run(runner.template, folder_of(corpus_files, tmp_path / "corpus"), tmp_path / "out", tmp_path)
    status, run = get_json(runner.base, "/api/runs", "POST", {"folder": str(tmp_path / "out")})
    assert status == 200, run
    assert (run["name"], run["extraction"], run["summary"]["total"]) == ("out", "Client info", 22)
    assert [entry["id"] for entry in get_json(runner.base, "/api/runs")[1]["runs"]] == [run["id"]]
    assert get_json(runner.base, "/api/runs", "POST", {"folder": str(tmp_path / "out")})[1]["id"] == run["id"]  # opened once
    aside = get_json(runner.base, f"/api/runs/{run['id']}/set-aside")[1]
    assert [d["name"] for d in aside["documents"]] == ["client_00021.pdf", "client_00022.pdf"]
    status, headers, _ = http_request(runner.base, "GET", f"/api/runs/{run['id']}/pages.png?doc=client_00001.pdf&page=0&dpi=40")
    assert status == 200 and headers["content-type"] == "image/png"
    assert not runs_dir(runner).exists() or not list(runs_dir(runner).iterdir())  # registered, not copied


def test_a_command_line_run_of_another_template_or_no_run_at_all_is_a_bad_run(runner, corpus_files, tmp_path, start_demo_server):
    folder = folder_of(corpus_files[:2], tmp_path / "corpus")
    edited = ExtractionTemplate.model_validate({**runner.template, "name": "Edited"})
    run_batch([edited], sorted(folder.glob("*.pdf")), tmp_path / "edited", workers=1)
    unknown = edited.model_copy(update={"template_id": "someone_elses"})
    run_batch([unknown], sorted(folder.glob("*.pdf")), tmp_path / "unknown", workers=1)
    for out, says in [(tmp_path / "edited", "client_forms"), (tmp_path / "unknown", "someone_elses")]:
        status, body = get_json(runner.base, "/api/runs", "POST", {"folder": str(out)})
        assert (status, body["error"]) == (400, "bad_run") and says in body["detail"], body
        assert str(tmp_path) not in body["detail"]
    for value in (str(folder), "relative/path", 7):
        status, body = get_json(runner.base, "/api/runs", "POST", {"folder": value})
        assert (status, body["error"]) == (400, "bad_run"), value
        assert str(tmp_path) not in body["detail"]
    assert get_json(runner.base, "/api/runs")[1] == {"runs": []}
    closed = start_demo_server(folder_scan=False)
    status, body = get_json(closed, "/api/runs", "POST", {"folder": str(tmp_path / "edited")})
    assert (status, body["error"]) == (403, "not_allowed")
